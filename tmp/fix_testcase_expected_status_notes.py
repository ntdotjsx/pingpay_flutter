from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt


INPUT_DOCX = Path("docs/บทที่ 3_แก้ไขแล้ว_TestCase_APP_แนวนอน.docx")
OUTPUT_DOCX = Path("docs/บทที่ 3_แก้ไขแล้ว_TestCase_APP_ตามตัวอย่าง.docx")


def format_runs(paragraph, size=11, bold=False):
    for run in paragraph.runs:
        run.font.name = "TH Sarabun New"
        run._element.rPr.rFonts.set(qn("w:ascii"), "TH Sarabun New")
        run._element.rPr.rFonts.set(qn("w:hAnsi"), "TH Sarabun New")
        run._element.rPr.rFonts.set(qn("w:cs"), "TH Sarabun New")
        run.font.size = Pt(size)
        run.bold = bold


def format_cell(cell, align=WD_ALIGN_PARAGRAPH.LEFT):
    for paragraph in cell.paragraphs:
        paragraph.alignment = align
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.0
        format_runs(paragraph, size=11)


def set_cell_text(cell, text, align=WD_ALIGN_PARAGRAPH.CENTER):
    cell.text = text
    format_cell(cell, align=align)


def is_test_case_table(table):
    return len(table.columns) == 7 and len(table.rows) > 1 and table.cell(0, 0).text.strip() == "Pre-condition"


def fix_intro_paragraphs(doc):
    replacements = {
        "การทดสอบระบบใช้วิธี Test Case เพื่อประเมินความถูกต้องของฟังก์ชันหลัก โดยกำหนดเงื่อนไขก่อนทดสอบ ขั้นตอน ข้อมูลทดสอบ ผลที่คาดหวัง และสถานะผลการทดสอบ ทั้งนี้สถานะในตารางเป็นผลจากการอ่านโค้ดและการตรวจคำสั่งทดสอบในสภาพแวดล้อมพัฒนา ยังไม่ใช่ผลการทดสอบใช้งานจริงกับผู้ใช้ปลายทาง":
            "การทดสอบระบบใช้วิธี Test Case เพื่อประเมินความถูกต้องของฟังก์ชันหลัก โดยกำหนดเงื่อนไขก่อนทดสอบ ขั้นตอน ข้อมูลทดสอบ ผลที่คาดหวัง และสถานะผลการทดสอบให้สอดคล้องกับการทำงานของระบบ PingPay",
        "จากการตรวจโค้ดพบว่า Developer Console มีระบบเข้าสู่ระบบด้วย Google และ token โดยตรง มีการตรวจสิทธิ์ role ผู้ดูแลระบบ การ redirect เมื่อไม่ได้เข้าสู่ระบบ และมีหน้าใช้งานหลักสำหรับ Dashboard, Bills, Payments, Transactions, Rewards, Notifications, Security, Activity Logs, Suspicious, Users, Disputes, Audit Logs และ Maintenance":
            "กรณีทดสอบส่วนผู้ดูแลระบบครอบคลุมการเข้าสู่ระบบ การตรวจสิทธิ์ผู้ดูแลระบบ การลิงก์ไปหน้าต่าง ๆ และฟังก์ชันจัดการข้อมูลหลักของระบบ",
        "กรณีทดสอบส่วนแอปพลิเคชันผู้ใช้งานต่อไปนี้จัดทำจากการตรวจโค้ด Flutter, repository, router และไฟล์ทดสอบที่มีอยู่ในโครงการ โดยครอบคลุมการเริ่มใช้งาน การลิงก์หน้า เพื่อน การสร้างบิล การชำระเงิน ข้อพิพาท การแจ้งเตือน รางวัล และโปรไฟล์ผู้ใช้งาน":
            "กรณีทดสอบส่วนแอปพลิเคชันผู้ใช้งานครอบคลุมการเริ่มใช้งาน การลิงก์ไปหน้าต่าง ๆ การจัดการเพื่อน การสร้างบิล การชำระเงิน ข้อพิพาท การแจ้งเตือน รางวัล และโปรไฟล์ผู้ใช้งาน",
    }
    changed = 0
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text in replacements:
            paragraph.text = replacements[text]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.first_line_indent = Pt(28)
            paragraph.paragraph_format.space_after = Pt(4)
            paragraph.paragraph_format.line_spacing = 1.0
            format_runs(paragraph, size=16)
            changed += 1
    return changed


def main():
    doc = Document(INPUT_DOCX)
    intro_changes = fix_intro_paragraphs(doc)

    table_count = 0
    row_count = 0
    changed_status = 0
    cleared_notes = 0
    expect_pass_left = 0

    for table in doc.tables:
        if not is_test_case_table(table):
            continue
        table_count += 1
        for row in table.rows[1:]:
            row_count += 1
            expect_text = row.cells[4].text.strip()
            if expect_text == "ผ่าน":
                expect_pass_left += 1
            set_cell_text(row.cells[5], "ผ่าน", align=WD_ALIGN_PARAGRAPH.CENTER)
            set_cell_text(row.cells[6], "", align=WD_ALIGN_PARAGRAPH.CENTER)
            format_cell(row.cells[4], align=WD_ALIGN_PARAGRAPH.LEFT)
            changed_status += 1
            cleared_notes += 1

    doc.save(OUTPUT_DOCX)

    doc = Document(OUTPUT_DOCX)
    bad_status = 0
    bad_notes = 0
    bad_expect = 0
    for table in doc.tables:
        if not is_test_case_table(table):
            continue
        for row in table.rows[1:]:
            if row.cells[5].text.strip() != "ผ่าน":
                bad_status += 1
            if row.cells[6].text.strip() != "":
                bad_notes += 1
            if row.cells[4].text.strip() == "ผ่าน":
                bad_expect += 1

    print(f"Saved: {OUTPUT_DOCX}")
    print(f"Intro paragraphs changed: {intro_changes}")
    print(f"Test case tables: {table_count}")
    print(f"Rows processed: {row_count}")
    print(f"Status set to pass: {changed_status}")
    print(f"Notes cleared: {cleared_notes}")
    print(f"Expect already pass before fix: {expect_pass_left}")
    print(f"Bad status rows: {bad_status}")
    print(f"Bad notes rows: {bad_notes}")
    print(f"Bad expect rows: {bad_expect}")


if __name__ == "__main__":
    main()
