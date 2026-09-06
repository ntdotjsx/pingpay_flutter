from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt


OUTPUT_DOCX = Path("docs/แบบสอบถามความพึงพอใจระบบPingPay.docx")
FONT_NAME = "TH Sarabun New"


QUESTIONS = [
    ("section", "ด้านประสิทธิภาพ"),
    ("1", "ความเหมาะสมของการจัดวางเมนูและลำดับการใช้งานของระบบ"),
    ("2", "ความรวดเร็วในการตอบสนองของแอปพลิเคชันและระบบผู้ดูแล"),
    ("3", "ความถูกต้องของการสร้างบิล แบ่งค่าใช้จ่าย และคำนวณยอดหนี้รายบุคคล"),
    ("section", "ด้านการออกแบบหน้าจอ"),
    ("4", "การจัดรูปแบบหน้าจอสำหรับอุปกรณ์มือถือมีความเหมาะสม"),
    ("5", "สีสันและองค์ประกอบของหน้าจอมีความเหมาะสมกับการใช้งาน"),
    ("6", "ความสวยงาม ความทันสมัย และความน่าสนใจของหน้าแอปพลิเคชัน"),
    ("7", "ขนาดตัวอักษรและรูปแบบตัวอักษรอ่านง่ายและชัดเจน"),
    ("8", "การจัดรูปแบบหน้าจอง่ายต่อการอ่านและการใช้งาน"),
    ("section", "ด้านความพึงพอใจโดยรวมในการใช้งานระบบ PingPay"),
    ("9", "ภาพรวมความพึงพอใจเกี่ยวกับระบบ PingPay"),
    ("10", "ความสะดวกในการใช้งานระบบสร้างบิล ชำระเงินผ่านพร้อมเพย์ ตรวจสลิป และรับการแจ้งเตือน"),
]


def set_run_font(run, size=16, bold=False):
    run.font.name = FONT_NAME
    run.font.size = Pt(size)
    run.font.bold = bold
    r_fonts = run._element.get_or_add_rPr().rFonts
    r_fonts.set(qn("w:ascii"), FONT_NAME)
    r_fonts.set(qn("w:hAnsi"), FONT_NAME)
    r_fonts.set(qn("w:eastAsia"), FONT_NAME)
    r_fonts.set(qn("w:cs"), FONT_NAME)


def add_para(doc, text="", size=16, bold=False, align=None, after=0, before=0, left=None):
    paragraph = doc.add_paragraph()
    if text:
        paragraph.add_run(text)
    paragraph.alignment = align if align is not None else WD_ALIGN_PARAGRAPH.LEFT
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = 1.0
    if left is not None:
        fmt.left_indent = left
    for run in paragraph.runs:
        set_run_font(run, size=size, bold=bold)
    return paragraph


def add_mixed_para(doc, runs, size=16, align=WD_ALIGN_PARAGRAPH.LEFT, after=0, before=0):
    paragraph = doc.add_paragraph()
    paragraph.alignment = align
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = 1.0
    for text, bold in runs:
        run = paragraph.add_run(text)
        set_run_font(run, size=size, bold=bold)
    return paragraph


def set_cell_text(cell, text, size=14, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    run = paragraph.add_run(text)
    set_run_font(run, size=size, bold=bold)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=45, start=70, bottom=45, end=70):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "000000")


def set_table_widths(table, widths_cm):
    table.autofit = False
    for row in table.rows:
        for cell, width in zip(row.cells, widths_cm):
            cell.width = Cm(width)


def build_question_table(doc):
    table = doc.add_table(rows=2 + len(QUESTIONS), cols=7)
    set_table_borders(table)
    widths = [1.45, 10.3, 0.9, 0.9, 0.9, 0.9, 0.9]
    set_table_widths(table, widths)

    row0 = table.rows[0].cells
    row0[0].merge(table.rows[1].cells[0])
    row0[1].merge(table.rows[1].cells[1])
    row0[2].merge(row0[6])

    set_cell_text(row0[0], "ลำดับที่", size=14, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_text(row0[1], "รายการหัวข้อคำถาม", size=14, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_text(row0[2], "ระดับความพึงพอใจ", size=14, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    for cell in (row0[0], row0[1], row0[2]):
        set_cell_shading(cell, "F2F2F2")

    for idx, score in enumerate(["5", "4", "3", "2", "1"], start=2):
        set_cell_text(table.rows[1].cells[idx], score, size=14, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_shading(table.rows[1].cells[idx], "F2F2F2")

    r = 2
    for item in QUESTIONS:
        cells = table.rows[r].cells
        if item[0] == "section":
            merged = cells[0].merge(cells[6])
            set_cell_text(merged, item[1], size=14, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT)
            set_cell_shading(merged, "BFBFBF")
            set_cell_margins(merged, top=35, bottom=35, start=80, end=80)
        else:
            number, question = item
            set_cell_text(cells[0], number, size=14, align=WD_ALIGN_PARAGRAPH.CENTER)
            set_cell_text(cells[1], question, size=14, align=WD_ALIGN_PARAGRAPH.LEFT)
            for c in range(2, 7):
                set_cell_text(cells[c], "", size=14, align=WD_ALIGN_PARAGRAPH.CENTER)
                set_cell_margins(cells[c], top=52, bottom=52, start=50, end=50)
        for cell in table.rows[r].cells:
            set_cell_margins(cell)
        r += 1

    for row in table.rows[:2]:
        for cell in row.cells:
            set_cell_margins(cell, top=45, bottom=45, start=55, end=55)
    return table


def main():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.left_margin = Cm(2.1)
    section.right_margin = Cm(2.1)
    section.top_margin = Cm(1.7)
    section.bottom_margin = Cm(1.6)

    add_para(doc, "แบบสอบถามความพึงพอใจ", size=18, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_para(
        doc,
        "โครงงานระบบจัดการค่าใช้จ่ายร่วมกันและการชำระเงินผ่านพร้อมเพย์ ด้วยเทคโนโลยี OCR (PingPay)",
        size=16,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        after=2,
    )
    add_mixed_para(
        doc,
        [
            ("คำชี้แจง", True),
            (" กรุณาใส่เครื่องหมาย ✓ ลงใน □ หน้าคำตอบที่ตรงกับความพึงพอใจของท่านมากที่สุด", False),
        ],
        size=16,
        after=2,
    )
    add_mixed_para(
        doc,
        [
            ("ตอนที่ 1", True),
            (" ข้อมูลผู้ใช้งานระบบ PingPay", False),
        ],
        size=16,
        after=0,
    )
    add_para(doc, "1.1 สิทธิเข้าใช้งานระบบ     □ ผู้ดูแลระบบ          □ ผู้ใช้งานระบบ", size=16, after=2, left=Pt(36))

    add_mixed_para(
        doc,
        [
            ("ตอนที่ 2", True),
            (" ความพึงพอใจเกี่ยวกับโครงงานระบบจัดการค่าใช้จ่ายร่วมกันและการชำระเงินผ่านพร้อมเพย์ ด้วยเทคโนโลยี OCR (PingPay)", False),
        ],
        size=16,
    )
    add_para(doc, "โดยระดับความพึงพอใจมีเกณฑ์การให้คะแนน ดังนี้", size=16, after=0)

    scale_rows = [
        "ระดับการประเมิน       5  หมายถึง       มีระดับความพึงพอใจมากที่สุด",
        "                       4  หมายถึง       มีระดับความพึงพอใจมาก",
        "                       3  หมายถึง       มีระดับความพึงพอใจปานกลาง",
        "                       2  หมายถึง       มีระดับความพึงพอใจน้อย",
        "                       1  หมายถึง       มีระดับความพึงพอใจน้อยที่สุด",
    ]
    for line in scale_rows:
        add_para(doc, line, size=15, after=0, left=Pt(58))

    build_question_table(doc)

    add_mixed_para(doc, [("ตอนที่ 3", True), (" ข้อเสนอแนะ", False)], size=16, before=4, after=0)
    add_para(doc, "........................................................................................................................................................", size=16, after=0)
    add_para(doc, "........................................................................................................................................................", size=16, after=0)

    doc.save(OUTPUT_DOCX)
    all_text = "\n".join(p.text for p in doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            all_text += "\n" + "\t".join(cell.text for cell in row.cells)
    print(f"Saved: {OUTPUT_DOCX}")
    print(f"Paragraphs: {len(doc.paragraphs)}")
    print(f"Tables: {len(doc.tables)}")
    for term in ["PingPay", "พร้อมเพย์", "OCR", "Yo Cinema", "ภาพยนตร์"]:
        print(f"{term}: {term in all_text}")


if __name__ == "__main__":
    main()
