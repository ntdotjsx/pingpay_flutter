from pathlib import Path
import re

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


DOCX = Path("docs/บทที่ 3_แก้ไขแล้ว_ตารางครบ.docx")


ROWS = [
    ("id", "รหัสผู้ใช้งาน", "uuid", "PK, Default: gen_random_uuid()", "", "550e8400-e29b-41d4-a716-446655440001"),
    ("user_code", "รหัสสาธารณะสำหรับค้นหาผู้ใช้ในระบบ", "varchar(32)", "Not Null, Unique", "", "USR-8F3K2A"),
    ("email", "อีเมลของผู้ใช้งาน", "varchar(255)", "", "", "somchai@example.com"),
    ("display_name", "ชื่อที่แสดงในแอปพลิเคชัน", "varchar(128)", "", "", "สมชาย"),
    ("full_name", "ชื่อและนามสกุลจริงของผู้ใช้งาน", "varchar(128)", "", "", "สมชาย ใจดี"),
    ("address", "ที่อยู่ปัจจุบันของผู้ใช้งาน", "text", "", "", "99/1 ถนนสุขุมวิท แขวงพระโขนง กรุงเทพฯ 10110"),
    ("phone_number", "หมายเลขโทรศัพท์มือถือของผู้ใช้งาน", "varchar(32)", "", "", "0812345678"),
    ("bank_account_number", "เลขที่บัญชีธนาคารสำหรับรับเงิน", "varchar(32)", "", "", "1234567890"),
    ("bank_name", "ชื่อธนาคารของบัญชีรับเงิน", "varchar(64)", "", "", "กสิกรไทย"),
    ("bank_code", "รหัสธนาคารของบัญชีรับเงิน", "varchar(16)", "", "", "004"),
    ("truemoney_phone", "หมายเลขโทรศัพท์สำหรับกระเป๋าเงินทรูมันนี่", "varchar(32)", "", "", "0812345678"),
    ("prompt_pay_id", "รหัสพร้อมเพย์สำหรับรับเงิน", "varchar(32)", "", "", "0812345678"),
    ("prompt_pay_id_type", "ประเภทรหัสพร้อมเพย์ที่ใช้รับเงิน", "promptpay_id_type", "", "", "mobile_number"),
    ("prompt_pay_verified_at", "วันและเวลาที่พร้อมเพย์ผ่านการตรวจสอบ", "timestamp", "", "", "2026-09-01 08:00:00"),
    ("avatar_url", "ที่อยู่ไฟล์รูปภาพโปรไฟล์", "text", "", "", "https://cdn.pingpay.app/avatars/u001.png"),
    ("reward_points", "จำนวนคะแนนสะสมของผู้ใช้งาน", "integer", "Not Null, Default: 0", "", "120"),
    ("shipping_address", "ที่อยู่สำหรับจัดส่งของรางวัล", "text", "", "", "99/1 ถนนสุขุมวิท แขวงพระโขนง กรุงเทพฯ 10110"),
    ("shipping_phone", "หมายเลขโทรศัพท์สำหรับจัดส่งของรางวัล", "varchar(32)", "", "", "0812345678"),
    ("shipping_recipient_name", "ชื่อผู้รับของรางวัล", "varchar(128)", "", "", "สมชาย ใจดี"),
    ("profile_completed_at", "วันและเวลาที่กรอกข้อมูลโปรไฟล์ครบถ้วน", "timestamp", "", "", "2026-09-01 08:03:00"),
    ("role", "บทบาทการใช้งานในระบบ", "user_role", "Not Null, Default: user", "", "user"),
    ("account_status", "สถานะบัญชีผู้ใช้งาน", "account_status", "Not Null, Default: active", "", "active"),
    ("suspended_until", "วันและเวลาสิ้นสุดการระงับบัญชี", "timestamp", "", "", "2026-09-15 08:00:00"),
    ("created_at", "วันและเวลาที่สร้างบัญชี", "timestamp", "Not Null, Default: now()", "", "2026-09-01 08:00:00"),
    ("updated_at", "วันและเวลาที่แก้ไขบัญชีล่าสุด", "timestamp", "Not Null, Default: now()", "", "2026-09-01 08:03:00"),
]


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, width):
    tc_pr = cell._tc.get_or_add_tcPr()
    tcw = tc_pr.find(qn("w:tcW"))
    if tcw is None:
        tcw = OxmlElement("w:tcW")
        tc_pr.append(tcw)
    tcw.set(qn("w:w"), str(width))
    tcw.set(qn("w:type"), "dxa")


def set_cell_margins(cell, top=60, start=80, bottom=60, end=80):
    tc_pr = cell._tc.get_or_add_tcPr()
    mar = tc_pr.find(qn("w:tcMar"))
    if mar is None:
        mar = OxmlElement("w:tcMar")
        tc_pr.append(mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def repeat_header_row(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = tr_pr.find(qn("w:tblHeader"))
    if tbl_header is None:
        tbl_header = OxmlElement("w:tblHeader")
        tr_pr.append(tbl_header)
    tbl_header.set(qn("w:val"), "true")


def keep_row_together(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = tr_pr.find(qn("w:cantSplit"))
    if cant_split is None:
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)
    cant_split.set(qn("w:val"), "true")


def set_table_geometry(table, widths):
    table.autofit = False
    table.allow_autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_layout = tbl_pr.find(qn("w:tblLayout"))
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    if grid is None:
        grid = OxmlElement("w:tblGrid")
        table._tbl.insert(0, grid)
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            set_cell_width(cell, width)


def format_cell(cell, *, header=False, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    set_cell_margins(cell)
    if header:
        set_cell_shading(cell, "D9EAF7")
    for paragraph in cell.paragraphs:
        paragraph.alignment = align
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.0
        for run in paragraph.runs:
            run.font.name = "TH Sarabun New"
            run._element.rPr.rFonts.set(qn("w:ascii"), "TH Sarabun New")
            run._element.rPr.rFonts.set(qn("w:hAnsi"), "TH Sarabun New")
            run._element.rPr.rFonts.set(qn("w:cs"), "TH Sarabun New")
            run.font.size = Pt(12)
            run.bold = bold


def remove_row(table, row):
    table._tbl.remove(row._tr)


def main():
    doc = Document(DOCX)
    table = doc.tables[0]
    headers = [cell.text for cell in table.rows[0].cells]

    while len(table.rows) - 1 > len(ROWS):
        remove_row(table, table.rows[-1])
    while len(table.rows) - 1 < len(ROWS):
        table.add_row()

    for cell, header in zip(table.rows[0].cells, headers):
        cell.text = header
        format_cell(cell, header=True, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    repeat_header_row(table.rows[0])
    keep_row_together(table.rows[0])

    for row, values in zip(table.rows[1:], ROWS):
        keep_row_together(row)
        for idx, (cell, value) in enumerate(zip(row.cells, values)):
            cell.text = value
            align = WD_ALIGN_PARAGRAPH.LEFT
            if idx in (2, 3, 4):
                align = WD_ALIGN_PARAGRAPH.CENTER
            format_cell(cell, align=align)

    set_table_geometry(table, [1350, 2500, 1250, 1700, 1300, 1500])
    doc.save(DOCX)

    doc = Document(DOCX)
    table = doc.tables[0]
    violations = []
    for row in table.rows[1:]:
        desc = row.cells[1].text.strip()
        if re.search(r"[A-Za-z]", desc):
            violations.append(desc)
    empty_required_context = []
    for row in table.rows[1:]:
        cells = [cell.text.strip() for cell in row.cells]
        if not cells[5]:
            empty_required_context.append(cells[0])
    print(f"Saved: {DOCX}")
    print(f"Rows including header: {len(table.rows)}")
    print(f"Description ASCII-letter violations: {len(violations)}")
    print(f"Rows without Simple Data: {empty_required_context}")
    if violations:
        raise SystemExit(2)


if __name__ == "__main__":
    main()
