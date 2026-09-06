from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt


INPUT_DOCX = Path("docs/แบบทดสอบภาคปฏิบัติวิชาโครงงาน.docx")
OUTPUT_DOCX = Path("docs/แบบทดสอบภาคปฏิบัติวิชาโครงงาน_ระบบPingPay.docx")

FONT_NAME = "TH Sarabun New"


PART1_ROWS = [
    ("ลำดับ", "รายการประเมิน", "คะแนนเต็ม", "คะแนนที่ได้", "หมายเหตุ"),
    ("section", "1. เอกสารรูปเล่มโครงงานระบบ PingPay (15 คะแนน)", "", "", ""),
    ("1.1", "ความถูกต้องตามรูปแบบที่กำหนด ครบถ้วนตั้งแต่บทที่ 1-5 สารบัญ ตาราง แผนภาพ และภาคผนวก", "5", "", ""),
    ("1.2", "ความสมบูรณ์ของเนื้อหาเกี่ยวกับระบบจัดการค่าใช้จ่ายร่วมกัน พร้อมเพย์ OCR การชำระเงิน และผู้ดูแลระบบ", "5", "", ""),
    ("1.3", "การใช้ภาษาไทย การอ้างอิง การจัดรูปแบบ แผนภาพฐานข้อมูล และคำอธิบายระบบถูกต้องชัดเจน", "5", "", ""),
    ("section", "2. การนำเสนอผลงาน (20 คะแนน)", "", "", ""),
    ("2.1", "สื่อนำเสนอมีความชัดเจน แสดงแนวคิด ปัญหา วัตถุประสงค์ และภาพรวมระบบ PingPay ได้ครบถ้วน", "5", "", ""),
    ("2.2", "ลำดับการนำเสนอเข้าใจง่าย ครอบคลุมส่วนผู้ใช้งาน แอปพลิเคชัน ผู้ดูแลระบบ ฐานข้อมูล และผลการทดสอบ", "5", "", ""),
    ("2.3", "อธิบายการทำงานของฟังก์ชันสำคัญ เช่น การสร้างบิล การแบ่งยอด การชำระเงิน การตรวจสลิป และการแจ้งเตือนได้ชัดเจน", "5", "", ""),
    ("2.4", "บุคลิกภาพ น้ำเสียง การใช้ภาษา และการบริหารเวลาในการนำเสนอเหมาะสม", "5", "", ""),
    ("section", "3. การตอบข้อซักถาม (15 คะแนน)", "", "", ""),
    ("3.1", "เข้าใจโครงสร้างระบบของตนเอง ทั้ง Flutter, API, ฐานข้อมูล, ระบบแจ้งเตือน และระบบผู้ดูแล", "5", "", ""),
    ("3.2", "ตอบคำถามเรื่องการออกแบบฐานข้อมูล ความปลอดภัย การตรวจสอบสิทธิ์ และการจัดการข้อมูลทางการเงินได้ตรงประเด็น", "5", "", ""),
    ("3.3", "สามารถอธิบายปัญหา อุปสรรค วิธีแก้ไข และเหตุผลในการเลือกเทคโนโลยีของโครงงานได้ชัดเจน", "5", "", ""),
    ("section", "4. ความคิดริเริ่มสร้างสรรค์และประโยชน์ของระบบ (10 คะแนน)", "", "", ""),
    ("4.1", "มีการประยุกต์ใช้ OCR, PromptPay QR, การตรวจสอบสลิป และการแจ้งเตือนเพื่อแก้ปัญหาการหารค่าใช้จ่ายร่วมกัน", "5", "", ""),
    ("4.2", "ระบบมีประโยชน์ต่อผู้ใช้งาน ช่วยลดความผิดพลาดในการคำนวณ ติดตามยอดหนี้ และเพิ่มความสะดวกในการชำระเงิน", "5", "", ""),
    ("", "รวมคะแนนส่วนที่ 1", "60", "", ""),
]


PART2_ROWS = [
    ("ลำดับ", "รายการประเมิน", "คะแนนเต็ม", "คะแนนที่ได้", "หมายเหตุ"),
    ("section", "1. การทำงานตามฟังก์ชันหลักของระบบ PingPay (20 คะแนน)", "", "", ""),
    ("1.1", "เข้าสู่ระบบด้วย Google ยอมรับนโยบาย PDPA ตั้งรหัส PIN และเข้าสู่แอปพลิเคชันได้ถูกต้อง", "5", "", "ทดสอบส่วนผู้ใช้งาน"),
    ("1.2", "จัดการโปรไฟล์ ช่องทางรับเงิน ข้อมูล PromptPay และรายชื่อเพื่อนได้ครบถ้วน", "5", "", "เพิ่ม/ยอมรับ/ลบเพื่อน"),
    ("1.3", "สร้างบิลจากการกรอกข้อมูลหรือสแกนใบเสร็จด้วย OCR แบ่งยอดค่าใช้จ่าย และบันทึกยอดหนี้รายบุคคลได้ถูกต้อง", "5", "", "ตรวจยอดรวมและยอดแบ่ง"),
    ("1.4", "ชำระเงินผ่าน PromptPay แนบสลิป ตรวจสอบสลิป ยืนยันหรือปฏิเสธการชำระเงิน และอัปเดตสถานะหนี้ได้ถูกต้อง", "5", "", "ครอบคลุมการแจ้งเตือน"),
    ("section", "2. ความถูกต้อง ความปลอดภัย และความเสถียร (10 คะแนน)", "", "", ""),
    ("2.1", "ตรวจสอบข้อมูลที่จำเป็น เช่น จำนวนเงิน สมาชิกบิล สถานะการชำระเงิน และสิทธิ์การเข้าถึงข้อมูลได้ถูกต้อง", "5", "", "ป้องกันข้อมูลผิดพลาด"),
    ("2.2", "ระบบ API ฐานข้อมูล session token PIN และการบันทึกเหตุการณ์ความปลอดภัยทำงานได้เสถียร", "5", "", "ไม่เกิดข้อผิดพลาดระหว่างสาธิต"),
    ("section", "3. ส่วนติดต่อผู้ใช้และระบบผู้ดูแล (10 คะแนน)", "", "", ""),
    ("3.1", "หน้าจอแอปพลิเคชันใช้งานง่าย เมนูหลักครบถ้วน ได้แก่ หน้าหลัก ชำระเงิน บิลของฉัน ของรางวัล และโปรไฟล์", "5", "", "เหมาะกับผู้ใช้งานทั่วไป"),
    ("3.2", "ระบบผู้ดูแลสามารถเข้าสู่ระบบ ตรวจสอบ Dashboard ผู้ใช้งาน บิล การชำระเงิน ธุรกรรม ข้อพิพาท การแจ้งเตือน และ Security Logs ได้", "5", "", "เหมาะกับผู้ดูแลระบบ"),
    ("", "รวมคะแนนส่วนที่ 2", "40", "", ""),
]


TEST_CASES = [
    "1. ผู้ใช้งานเข้าสู่ระบบด้วย Google ยอมรับ PDPA ตั้งค่า PIN และเข้าสู่หน้าหลักสำเร็จ",
    "2. ผู้ใช้งานแก้ไขโปรไฟล์ เพิ่มข้อมูล PromptPay และกำหนดช่องทางรับเงินได้",
    "3. ผู้ใช้งานค้นหาเพื่อน ส่งคำขอ ยอมรับคำขอ และแสดงรายชื่อเพื่อนได้ถูกต้อง",
    "4. ผู้ใช้งานสร้างบิลใหม่จากการกรอกข้อมูลหรือสแกนใบเสร็จด้วย OCR พร้อมตรวจยอดรวมได้",
    "5. ระบบแบ่งค่าใช้จ่ายให้สมาชิกในบิล บันทึกยอดหนี้รายบุคคล และแสดงสถานะค้างชำระได้",
    "6. ผู้ใช้งานชำระเงินผ่าน PromptPay QR แนบสลิป และระบบบันทึกประวัติการชำระเงินได้",
    "7. เจ้าของบิลยืนยันหรือปฏิเสธสลิปการชำระเงิน และระบบปรับสถานะหนี้พร้อมแจ้งเตือนได้",
    "8. ผู้ใช้งานดูการแจ้งเตือน ประวัติการชำระเงิน รายการบิล ของรางวัล และส่งข้อพิพาทได้",
    "9. ผู้ดูแลระบบเข้าสู่ Developer Console และตรวจสอบ Dashboard ผู้ใช้งาน บิล การชำระเงิน ธุรกรรม ข้อพิพาท การแจ้งเตือน และบันทึกความปลอดภัยได้",
    "10. ระบบรองรับการตรวจสอบข้อมูลผิดพลาด เช่น ข้อมูลไม่ครบ จำนวนเงินไม่ถูกต้อง สิทธิ์ไม่ถูกต้อง หรือ session หมดอายุ",
]


SUMMARY_ROWS = [
    ("ส่วนที่", "คะแนนเต็ม", "คะแนนที่ได้"),
    ("ส่วนที่ 1: การนำเสนอและสอบปากเปล่าโครงงาน PingPay", "60", ""),
    ("ส่วนที่ 2: การทดสอบประสิทธิภาพระบบ PingPay", "40", ""),
    ("รวมทั้งสิ้น", "100", ""),
]


def set_run_font(run, size=16, bold=False):
    run.font.name = FONT_NAME
    run.font.size = Pt(size)
    run.font.bold = bold
    r_fonts = run._element.get_or_add_rPr().rFonts
    r_fonts.set(qn("w:ascii"), FONT_NAME)
    r_fonts.set(qn("w:hAnsi"), FONT_NAME)
    r_fonts.set(qn("w:eastAsia"), FONT_NAME)
    r_fonts.set(qn("w:cs"), FONT_NAME)


def style_paragraph(paragraph, size=16, bold=False, align=None, before=0, after=0, first_line=None, left=None):
    if align is not None:
        paragraph.alignment = align
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = 1.0
    if first_line is not None:
        fmt.first_line_indent = first_line
    if left is not None:
        fmt.left_indent = left
    for run in paragraph.runs:
        set_run_font(run, size=size, bold=bold)


def clear_body_keep_section(doc):
    body = doc._body._element
    sect_pr = body.sectPr
    sect_copy = deepcopy(sect_pr) if sect_pr is not None else None
    for child in list(body):
        body.remove(child)
    if sect_copy is not None:
        body.append(sect_copy)


def set_cell_text(cell, text, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT, size=13):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    run = paragraph.add_run(text)
    set_run_font(run, size=size, bold=bold)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=80, bottom=80, end=80):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "000000")


def set_table_width(table, widths_cm):
    table.autofit = False
    for row in table.rows:
        for cell, width in zip(row.cells, widths_cm):
            cell.width = Cm(width)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = tr_pr.find(qn("w:tblHeader"))
    if tbl_header is None:
        tbl_header = OxmlElement("w:tblHeader")
        tr_pr.append(tbl_header)
    tbl_header.set(qn("w:val"), "true")


def add_table(doc, rows, widths_cm):
    table = doc.add_table(rows=len(rows), cols=len(widths_cm))
    try:
        table.style = "Table Grid"
    except KeyError:
        pass
    set_table_borders(table)
    set_table_width(table, widths_cm)

    for r_idx, row_data in enumerate(rows):
        cells = table.rows[r_idx].cells
        if r_idx == 0:
            set_repeat_table_header(table.rows[r_idx])
        if row_data[0] == "section":
            merged = cells[0].merge(cells[-1])
            set_cell_text(merged, row_data[1], bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, size=13)
            set_cell_shading(merged, "D9EAF7")
            set_cell_margins(merged, top=45, bottom=45, start=90, end=90)
            continue

        is_header = r_idx == 0
        is_total = row_data[1].startswith("รวมคะแนน")
        for c_idx, text in enumerate(row_data):
            align = WD_ALIGN_PARAGRAPH.LEFT
            if c_idx in (0, 2, 3):
                align = WD_ALIGN_PARAGRAPH.CENTER
            if is_header:
                align = WD_ALIGN_PARAGRAPH.CENTER
            set_cell_text(cells[c_idx], text, bold=is_header or is_total, align=align, size=13)
            set_cell_margins(cells[c_idx], top=42, bottom=42, start=70, end=70)
            if is_header or is_total:
                set_cell_shading(cells[c_idx], "EDEDED")
    return table


def add_para(doc, text="", size=16, bold=False, align=None, before=0, after=0, first_line=None, left=None):
    paragraph = doc.add_paragraph()
    if text:
        paragraph.add_run(text)
    style_paragraph(paragraph, size=size, bold=bold, align=align, before=before, after=after, first_line=first_line, left=left)
    return paragraph


def add_blank(doc, after=2):
    paragraph = doc.add_paragraph()
    style_paragraph(paragraph, size=8, after=after)
    return paragraph


def main():
    doc = Document(INPUT_DOCX)
    clear_body_keep_section(doc)

    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.7)

    add_para(doc, "แบบทดสอบภาคปฏิบัติ: วิชาโครงงานด้านคอมพิวเตอร์โปรแกรมเมอร์ (31903-2016)", 16, True, WD_ALIGN_PARAGRAPH.CENTER, after=0)
    add_para(doc, "วิทยาลัยเทคนิคนวมินทราชินีมุกดาหาร  ภาคเรียนที่: 1 ปีการศึกษา: 2568", 16, False, WD_ALIGN_PARAGRAPH.CENTER, after=6)
    add_para(doc, "ชื่อ-สกุล: ..................................................... รหัสนักศึกษา: ........................................................................", 16, False, after=0)
    add_para(doc, "ชื่อโครงงาน: ระบบจัดการค่าใช้จ่ายร่วมกันและการชำระเงินผ่านพร้อมเพย์ ด้วยเทคโนโลยี OCR (PingPay)", 16, False, after=8)

    add_para(doc, "ส่วนที่ 1: การนำเสนอและสอบปากเปล่าโครงงาน (60 คะแนน)", 16, True, after=2)
    add_para(doc, "คำชี้แจง: ให้นักศึกษานำเสนอความคืบหน้าและผลการดำเนินงานโครงงานระบบ PingPay ต่อคณะกรรมการสอบ โดยใช้เวลาในการนำเสนอ 15 นาที และตอบข้อซักถาม 10-15 นาที", 16, False, after=2)
    add_para(doc, "เกณฑ์การให้คะแนน:", 16, True, after=2)
    add_table(doc, PART1_ROWS, [1.25, 10.7, 1.9, 2.1, 2.8])

    add_para(doc, "ข้อเสนอแนะจากกรรมการ:", 16, True, before=8, after=2)
    add_para(doc, "............................................................................................................................................................", 16, False, after=0)
    add_para(doc, "............................................................................................................................................................", 16, False, after=4)

    add_para(doc, "ส่วนที่ 2: การทดสอบประสิทธิภาพระบบ PingPay (40 คะแนน)", 16, True, before=4, after=2)
    add_para(doc, "คำชี้แจง: ให้นักศึกษาสาธิตการทำงานของระบบ PingPay ตามกรณีทดสอบ (Test Case) ที่กำหนด โดยครอบคลุมทั้งส่วนแอปพลิเคชันผู้ใช้งานและส่วนผู้ดูแลระบบ", 16, False, after=2)
    add_para(doc, "เกณฑ์การให้คะแนน:", 16, True, after=2)
    add_table(doc, PART2_ROWS, [1.25, 10.3, 1.9, 2.1, 3.2])

    add_para(doc, "สถานการณ์/กรณีทดสอบ (Test Case) ที่ใช้ในการประเมิน:", 16, True, before=8, after=2)
    for item in TEST_CASES:
        add_para(doc, item, 13, False, after=0, left=Pt(18))

    add_blank(doc)
    add_para(doc, "สรุปผลคะแนน", 16, True, after=2)
    add_table(doc, SUMMARY_ROWS, [10.8, 3.2, 4.0])

    add_para(doc, "ลงชื่อ ................................................... กรรมการสอบ", 15, False, WD_ALIGN_PARAGRAPH.RIGHT, before=4, after=0)
    add_para(doc, "(...................................................)", 15, False, WD_ALIGN_PARAGRAPH.RIGHT, after=0)
    add_para(doc, "วันที่ ........ เดือน ........................ พ.ศ. ........", 15, False, WD_ALIGN_PARAGRAPH.RIGHT, after=0)

    doc.save(OUTPUT_DOCX)
    all_text = "\n".join(p.text for p in doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            all_text += "\n" + "\t".join(cell.text for cell in row.cells)
    print(f"Saved: {OUTPUT_DOCX}")
    print(f"Paragraphs: {len(doc.paragraphs)}")
    print(f"Tables: {len(doc.tables)}")
    for term in ["PingPay", "พร้อมเพย์", "OCR", "โดเนท", "สตรีมเมอร์", "ภาพยนตร์", "Yo Cinema", "TSS"]:
        print(f"{term}: {term in all_text}")


if __name__ == "__main__":
    main()
