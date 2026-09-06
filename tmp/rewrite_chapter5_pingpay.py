from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


INPUT_DOCX = Path("docs/บทที่-5.docx")
OUTPUT_DOCX = Path("docs/บทที่-5_แก้ไขตามระบบ.docx")


CONTENT = [
    ("title", "บทที่ 5"),
    ("title", "สรุปผลการดำเนินงาน"),
    (
        "body",
        "โครงงานระบบจัดการค่าใช้จ่ายร่วมกันและการชำระเงินผ่านพร้อมเพย์ ด้วยเทคโนโลยี OCR "
        "เป็นระบบที่พัฒนาขึ้นเพื่อช่วยให้ผู้ใช้งานสามารถสร้างบิล แบ่งค่าใช้จ่ายร่วมกับเพื่อน "
        "ติดตามยอดค้างชำระ ชำระเงินผ่านพร้อมเพย์ ตรวจสอบสลิป และรับการแจ้งเตือนสถานะต่าง ๆ ได้อย่างเป็นระบบ "
        "ประกอบด้วยสรุปการดำเนินโครงงานดังต่อไปนี้",
    ),
    ("toc", "5.1 สรุปผลโครงงาน"),
    ("toc", "5.2 อภิปรายผล"),
    ("toc", "5.3 ข้อเสนอแนะ"),
    ("heading", "5.1 สรุปผลโครงงาน"),
    ("subheading", "5.1.1 วัตถุประสงค์ของโครงงาน"),
    (
        "item",
        "5.1.1.1 เพื่อออกแบบและพัฒนาระบบจัดการค่าใช้จ่ายร่วมกันและการชำระเงินผ่านพร้อมเพย์ "
        "ด้วยเทคโนโลยี OCR สำหรับช่วยให้ผู้ใช้งานจัดการบิลและยอดหนี้ร่วมกันได้สะดวก",
    ),
    (
        "item",
        "5.1.1.2 เพื่อทดสอบและประเมินประสิทธิภาพการทำงานของระบบ PingPay ทั้งในส่วนแอปพลิเคชันผู้ใช้งาน "
        "และส่วนผู้ดูแลระบบ",
    ),
    (
        "item",
        "5.1.1.3 เพื่อศึกษาความพึงพอใจของผู้ใช้งานที่มีต่อระบบ PingPay หลังจากทดลองใช้งานระบบ",
    ),
    ("subheading", "5.1.2 ขอบเขตการดำเนินงาน"),
    (
        "body",
        "การพัฒนาระบบ PingPay ใช้เครื่องมือและเทคโนโลยีที่เกี่ยวข้อง ได้แก่ Visual Studio Code, Flutter และ Dart "
        "สำหรับพัฒนาแอปพลิเคชันมือถือ, Elysia และ TypeScript สำหรับพัฒนา REST API, PostgreSQL และ Drizzle ORM "
        "สำหรับจัดการฐานข้อมูล, SvelteKit สำหรับระบบผู้ดูแลระบบ, Firebase Cloud Messaging สำหรับการแจ้งเตือน "
        "รวมถึงเทคโนโลยี OCR, PromptPay QR และการตรวจสอบสลิปเพื่อรองรับกระบวนการชำระเงิน โดยโครงสร้างระบบแบ่งออกเป็น 2 ส่วน ดังนี้",
    ),
    ("subheading", "5.1.2.1 โครงสร้างในส่วนของผู้ดูแลระบบ ประกอบด้วย"),
    ("item", "1) การเข้าสู่ระบบและตรวจสอบสิทธิ์ผู้ดูแลระบบ"),
    ("item", "2) การแสดงผล Dashboard และสรุปภาพรวมการใช้งาน"),
    ("item", "3) การจัดการข้อมูลผู้ใช้งาน สถานะบัญชี และข้อมูลความปลอดภัย"),
    ("item", "4) การตรวจสอบข้อมูลบิล การชำระเงิน ธุรกรรม และสลิปโอนเงิน"),
    ("item", "5) การจัดการข้อพิพาท การแจ้งเตือน ของรางวัล และประวัติการทำงานของผู้ดูแลระบบ"),
    ("subheading", "5.1.2.2 โครงสร้างในส่วนของผู้ใช้งาน ประกอบด้วย"),
    ("item", "1) การเข้าสู่ระบบด้วย Google การยอมรับนโยบาย PDPA และการตั้งรหัส PIN"),
    ("item", "2) การตั้งค่าโปรไฟล์ ช่องทางรับเงิน และข้อมูล PromptPay"),
    ("item", "3) การค้นหา เพิ่ม ยอมรับ ปฏิเสธ และลบเพื่อนในระบบ"),
    ("item", "4) การสร้างบิลจากการกรอกข้อมูล การสแกนใบเสร็จด้วย OCR และการกรอกข้อมูลด้วยภาษาธรรมชาติ"),
    ("item", "5) การแบ่งยอดค่าใช้จ่าย ติดตามยอดหนี้ ชำระเงิน แนบสลิป ยืนยันหรือปฏิเสธการชำระเงิน"),
    ("item", "6) การยื่นข้อพิพาท รับการแจ้งเตือน แลกของรางวัล ดูสรุปค่าใช้จ่าย และส่งข้อเสนอแนะ"),
    ("subheading", "5.1.3 ปัญหาและอุปสรรค"),
    (
        "item",
        "5.1.3.1 การอ่านข้อมูลจากใบเสร็จด้วย OCR มีความท้าทาย เนื่องจากใบเสร็จแต่ละร้านมีรูปแบบต่างกัน "
        "เช่น ตำแหน่งชื่อสินค้า ราคา ส่วนลด ภาษี และยอดรวม ทำให้ต้องออกแบบขั้นตอนตรวจทานข้อมูลก่อนสร้างบิล",
    ),
    (
        "item",
        "5.1.3.2 การจัดการสถานะหนี้และการชำระเงินต้องระมัดระวังเรื่องความถูกต้องของยอดเงิน "
        "การยืนยันสลิป การป้องกันการส่งข้อมูลซ้ำ และการบันทึกประวัติการเปลี่ยนแปลงอย่างครบถ้วน",
    ),
    (
        "item",
        "5.1.3.3 ระบบแจ้งเตือนและการรักษาความปลอดภัยต้องรองรับหลายสถานการณ์ เช่น การเปลี่ยนอุปกรณ์ "
        "การหมดอายุของ session การลงทะเบียน token แจ้งเตือน และการตรวจสอบสิทธิ์ก่อนเข้าถึงข้อมูล",
    ),
    ("subheading", "5.1.4 สรุปผลจากแบบสอบถามความพึงพอใจ"),
    (
        "body",
        "ผลการทดสอบระบบตามกรณีทดสอบที่กำหนดไว้ทั้งหมด 119 กรณี พบว่าผ่านทั้งหมด คิดเป็นร้อยละ 100 "
        "ครอบคลุมทั้งส่วนผู้ดูแลระบบและส่วนผู้ใช้งาน ได้แก่ การเข้าสู่ระบบ การจัดการข้อมูลสมาชิก การสร้างบิล "
        "การแบ่งยอด การชำระเงิน การตรวจสอบสลิป การแจ้งเตือน ข้อพิพาท ของรางวัล และโปรไฟล์ผู้ใช้งาน",
    ),
    (
        "body",
        "ผลการประเมินความพึงพอใจจากกลุ่มตัวอย่างจำนวน 20 คน พบว่าผู้ตอบแบบสอบถามมีความพึงพอใจต่อระบบ PingPay "
        "โดยภาพรวมอยู่ในระดับมาก มีค่าเฉลี่ยรวม 4.44 และส่วนเบี่ยงเบนมาตรฐาน 0.61 "
        "รายการที่มีค่าเฉลี่ยสูงสุดคือความสะดวกในการสร้างบิลและแบ่งค่าใช้จ่ายร่วมกับเพื่อน มีค่าเฉลี่ย 4.60 "
        "อยู่ในระดับมากที่สุด รองลงมาคือความถูกต้องของการเข้าสู่ระบบและการตั้งค่าโปรไฟล์ "
        "และความง่ายในการเข้าถึงเมนูหลักของแอปพลิเคชัน มีค่าเฉลี่ย 4.55 อยู่ในระดับมากที่สุด",
    ),
    ("heading", "5.2 อภิปรายผล"),
    (
        "body",
        "จากการพัฒนาระบบ PingPay พบว่าระบบสามารถตอบสนองวัตถุประสงค์ของโครงงานได้ครบถ้วน "
        "ผู้ใช้งานสามารถสร้างบิล แบ่งค่าใช้จ่าย ติดตามยอดค้างชำระ และชำระเงินผ่านพร้อมเพย์ได้ในขั้นตอนเดียว "
        "ช่วยลดความผิดพลาดจากการคำนวณด้วยตนเอง และช่วยให้การติดตามยอดหนี้ระหว่างเพื่อนเป็นระบบมากขึ้น",
    ),
    (
        "body",
        "ด้านการออกแบบระบบ ผู้ใช้งานสามารถเข้าถึงเมนูหลักได้ชัดเจน ได้แก่ หน้าหลัก หน้าชำระเงิน หน้าบิลของฉัน "
        "หน้าของรางวัล และหน้าโปรไฟล์ การมีระบบเพื่อน การแจ้งเตือน และประวัติการชำระเงินช่วยให้ผู้ใช้ติดตามสถานะของแต่ละบิลได้สะดวก "
        "สอดคล้องกับผลประเมินที่ผู้ใช้งานให้คะแนนความสะดวกในการสร้างบิลและการใช้งานเมนูหลักอยู่ในระดับมากถึงมากที่สุด",
    ),
    (
        "body",
        "ด้านผู้ดูแลระบบ ระบบ Developer Console ช่วยให้ผู้ดูแลสามารถตรวจสอบข้อมูลผู้ใช้ บิล การชำระเงิน ธุรกรรม "
        "ข้อพิพาท การแจ้งเตือน และเหตุการณ์ด้านความปลอดภัยได้จากศูนย์กลาง ทำให้สามารถติดตามความผิดปกติ "
        "และจัดการข้อมูลที่เกี่ยวข้องกับการใช้งานระบบได้รวดเร็วขึ้น",
    ),
    (
        "body",
        "ด้านประสิทธิภาพและความน่าเชื่อถือ ระบบมีการออกแบบให้รองรับการตรวจสอบสิทธิ์ การจัดเก็บ token อย่างปลอดภัย "
        "การยืนยันรหัส PIN การป้องกันการส่งคำขอซ้ำ และการบันทึกประวัติการทำรายการ ซึ่งช่วยเพิ่มความมั่นใจในการใช้งานระบบที่เกี่ยวข้องกับข้อมูลทางการเงิน",
    ),
    ("heading", "5.3 ข้อเสนอแนะ"),
    ("item", "5.3.1 ควรพัฒนา OCR ให้รองรับรูปแบบใบเสร็จจากร้านค้าและบริการที่หลากหลายมากขึ้น"),
    ("item", "5.3.2 ควรเพิ่มการแจ้งเตือนกำหนดชำระเงินล่วงหน้า และแจ้งเตือนซ้ำสำหรับยอดค้างชำระที่ใกล้ครบกำหนด"),
    ("item", "5.3.3 ควรรองรับการตรวจสอบสลิปจากหลายธนาคารและหลายช่องทางการชำระเงินมากขึ้น"),
    ("item", "5.3.4 ควรเพิ่มรายงานสรุปค่าใช้จ่ายรายเดือนหรือรายปี เพื่อให้ผู้ใช้งานเห็นพฤติกรรมการใช้จ่ายของตนเอง"),
    ("item", "5.3.5 ควรเพิ่มเครื่องมือวิเคราะห์ข้อมูลในส่วนผู้ดูแลระบบ เช่น กราฟการใช้งาน การเติบโตของผู้ใช้ และรายการผิดปกติที่ต้องติดตาม"),
    ("item", "5.3.6 ควรพัฒนาระบบกู้คืนรหัส PIN และระบบรักษาความปลอดภัยเพิ่มเติม เพื่อรองรับกรณีผู้ใช้งานเปลี่ยนอุปกรณ์หรือไม่สามารถเข้าสู่ระบบได้"),
]


def clear_body_keep_section(doc):
    body = doc._body._element
    sect_pr = body.sectPr
    sect_pr_copy = deepcopy(sect_pr) if sect_pr is not None else None
    for child in list(body):
        body.remove(child)
    if sect_pr_copy is not None:
        body.append(sect_pr_copy)


def set_doc_geometry(doc):
    for section in doc.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.left_margin = Inches(1.5)
        section.right_margin = Inches(1)
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)


def set_run_font(run, size=16, bold=False):
    run.font.name = "TH Sarabun New"
    run._element.rPr.rFonts.set(qn("w:ascii"), "TH Sarabun New")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "TH Sarabun New")
    run._element.rPr.rFonts.set(qn("w:cs"), "TH Sarabun New")
    run.font.size = Pt(size)
    run.bold = bold


def add_para(doc, text, kind):
    paragraph = doc.add_paragraph()
    paragraph.add_run(text)

    if kind == "title":
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.first_line_indent = Pt(0)
        paragraph.paragraph_format.space_after = Pt(2)
        size = 18
        bold = True
    elif kind == "heading":
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.first_line_indent = Pt(0)
        paragraph.paragraph_format.space_before = Pt(8)
        paragraph.paragraph_format.space_after = Pt(2)
        size = 16
        bold = True
    elif kind == "subheading":
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.first_line_indent = Pt(28)
        paragraph.paragraph_format.space_before = Pt(3)
        paragraph.paragraph_format.space_after = Pt(1)
        size = 16
        bold = False
    elif kind == "toc":
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.left_indent = Pt(28)
        paragraph.paragraph_format.first_line_indent = Pt(0)
        paragraph.paragraph_format.space_after = Pt(1)
        size = 16
        bold = False
    elif kind == "item":
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.left_indent = Pt(42)
        paragraph.paragraph_format.first_line_indent = Pt(0)
        paragraph.paragraph_format.space_after = Pt(1)
        size = 15
        bold = False
    else:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.first_line_indent = Pt(28)
        paragraph.paragraph_format.space_after = Pt(2)
        size = 16
        bold = False

    paragraph.paragraph_format.line_spacing = 1.0
    for run in paragraph.runs:
        set_run_font(run, size=size, bold=bold)
    return paragraph


def main():
    doc = Document(INPUT_DOCX)
    clear_body_keep_section(doc)
    set_doc_geometry(doc)

    for kind, text in CONTENT:
        add_para(doc, text, kind)

    doc.save(OUTPUT_DOCX)

    all_text = "\n".join(p.text for p in doc.paragraphs)
    bad_terms = ["โดเนท", "สตรีมเมอร์", "คอนเทนต์", "ภาพยนตร์", "Yo Cinema", "TSS"]
    print(f"Saved: {OUTPUT_DOCX}")
    print(f"Paragraphs: {len(doc.paragraphs)}")
    for term in bad_terms:
        print(f"{term}: {term in all_text}")


if __name__ == "__main__":
    main()
