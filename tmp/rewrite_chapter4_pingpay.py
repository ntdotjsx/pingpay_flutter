from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


INPUT_DOCX = Path("docs/บทที่-4.docx")
OUTPUT_DOCX = Path("docs/บทที่-4_แก้ไขตามระบบ.docx")


def clear_body_keep_section(doc):
    body = doc._body._element
    sect_pr = body.sectPr
    sect_pr_copy = deepcopy(sect_pr) if sect_pr is not None else None
    for child in list(body):
        body.remove(child)
    if sect_pr_copy is not None:
        body.append(sect_pr_copy)


def set_doc_geometry(doc):
    for section in doc.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.left_margin = Inches(1.5)
        section.right_margin = Inches(1)
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)


def set_run_font(run, size=16, bold=False):
    run.font.name = "TH Sarabun New"
    run._element.rPr.rFonts.set(qn("w:ascii"), "TH Sarabun New")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "TH Sarabun New")
    run._element.rPr.rFonts.set(qn("w:cs"), "TH Sarabun New")
    run.font.size = Pt(size)
    run.bold = bold


def format_paragraph(paragraph, size=16, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT, first_indent=True, before=0, after=2):
    paragraph.alignment = align
    paragraph.paragraph_format.first_line_indent = Pt(28) if first_indent else Pt(0)
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = 1.0
    for run in paragraph.runs:
        set_run_font(run, size=size, bold=bold)


def add_paragraph(doc, text, *, size=16, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT, first_indent=True, before=0, after=2):
    paragraph = doc.add_paragraph()
    paragraph.add_run(text)
    format_paragraph(paragraph, size=size, bold=bold, align=align, first_indent=first_indent, before=before, after=after)
    return paragraph


def add_heading(doc, text):
    return add_paragraph(doc, text, size=16, bold=True, first_indent=False, before=6, after=2)


def add_title(doc, text, size=18):
    return add_paragraph(doc, text, size=size, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, first_indent=False, before=0, after=2)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=80, bottom=80, end=80):
    tc_pr = cell._tc.get_or_add_tcPr()
    mar = tc_pr.find(qn("w:tcMar"))
    if mar is None:
        mar = OxmlElement("w:tcMar")
        tc_pr.append(mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width):
    tc_pr = cell._tc.get_or_add_tcPr()
    tcw = tc_pr.find(qn("w:tcW"))
    if tcw is None:
        tcw = OxmlElement("w:tcW")
        tc_pr.append(tcw)
    tcw.set(qn("w:w"), str(width))
    tcw.set(qn("w:type"), "dxa")


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = tr_pr.find(qn("w:tblHeader"))
    if tbl_header is None:
        tbl_header = OxmlElement("w:tblHeader")
        tr_pr.append(tbl_header)
    tbl_header.set(qn("w:val"), "true")


def set_table_geometry(table, widths):
    table.autofit = False
    table.allow_autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_layout = tbl_pr.find(qn("w:tblLayout"))
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")
    grid = table._tbl.tblGrid
    if grid is None:
        grid = OxmlElement("w:tblGrid")
        table._tbl.insert(0, grid)
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            set_cell_width(cell, width)


def format_cell(cell, header=False, align=WD_ALIGN_PARAGRAPH.CENTER):
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    set_cell_margins(cell)
    if header:
        set_cell_shading(cell, "D9EAF7")
    for paragraph in cell.paragraphs:
        paragraph.alignment = align
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.0
        for run in paragraph.runs:
            set_run_font(run, size=14, bold=header)


def add_table(doc, caption, headers, rows, widths, aligns=None):
    add_paragraph(doc, caption, size=16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, first_indent=False, before=6, after=2)
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, headers):
        cell.text = text
        format_cell(cell, header=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    repeat_header(table.rows[0])
    aligns = aligns or [WD_ALIGN_PARAGRAPH.CENTER] * len(headers)
    for values in rows:
        cells = table.add_row().cells
        for idx, (cell, value) in enumerate(zip(cells, values)):
            cell.text = str(value)
            format_cell(cell, header=False, align=aligns[idx])
    set_table_geometry(table, widths)
    add_paragraph(doc, "", size=8, first_indent=False, after=0)
    return table


def main():
    doc = Document(INPUT_DOCX)
    clear_body_keep_section(doc)
    set_doc_geometry(doc)

    add_title(doc, "บทที่ 4")
    add_title(doc, "ผลการดำเนินงาน")
    add_paragraph(
        doc,
        "การดำเนินงานโครงงานพัฒนาระบบจัดการค่าใช้จ่ายร่วมกันและการชำระเงินผ่านพร้อมเพย์ ด้วยเทคโนโลยี OCR "
        "ผู้จัดทำได้พัฒนาระบบ PingPay ตามวัตถุประสงค์ของโครงงาน โดยระบบประกอบด้วยแอปพลิเคชันสำหรับผู้ใช้งาน "
        "และระบบผู้ดูแลระบบสำหรับตรวจสอบข้อมูลและบริหารจัดการการใช้งาน",
    )
    add_paragraph(
        doc,
        "ผลการดำเนินงานแบ่งออกเป็น 2 ส่วน ได้แก่ ผลการทดสอบระบบตามกรณีทดสอบ และผลการศึกษาความพึงพอใจของผู้ใช้งานที่มีต่อระบบ",
    )

    add_heading(doc, "4.1 ผลการทดสอบระบบจัดการค่าใช้จ่ายร่วมกันและการชำระเงินผ่านพร้อมเพย์ ด้วยเทคโนโลยี OCR")
    add_paragraph(
        doc,
        "ผู้จัดทำได้ทดสอบระบบตามกรณีทดสอบที่กำหนดไว้ในบทที่ 3 โดยครอบคลุมการทำงานของผู้ดูแลระบบและแอปพลิเคชันผู้ใช้งาน "
        "ผลการทดสอบพบว่าฟังก์ชันหลักของระบบสามารถทำงานได้ตามผลลัพธ์ที่คาดหวัง ดังแสดงในตารางที่ 4.1",
    )

    test_rows = [
        ("1", "การเข้าสู่ระบบของผู้ดูแลระบบ", "8", "8", "0", "ผ่าน"),
        ("2", "การลิงก์ไปหน้าต่าง ๆ ของผู้ดูแลระบบ", "13", "13", "0", "ผ่าน"),
        ("3", "ฟังก์ชันจัดการข้อมูลของผู้ดูแลระบบ", "15", "15", "0", "ผ่าน"),
        ("4", "ข้อผิดพลาดและสถานะระบบของผู้ดูแลระบบ", "10", "10", "0", "ผ่าน"),
        ("5", "การเริ่มใช้งานแอปพลิเคชันของผู้ใช้งาน", "9", "9", "0", "ผ่าน"),
        ("6", "การลิงก์ไปหน้าต่าง ๆ ของแอปพลิเคชันผู้ใช้งาน", "13", "13", "0", "ผ่าน"),
        ("7", "ฟังก์ชันเพื่อนของผู้ใช้งาน", "11", "11", "0", "ผ่าน"),
        ("8", "การสร้างบิลและการแบ่งยอดของผู้ใช้งาน", "13", "13", "0", "ผ่าน"),
        ("9", "การชำระเงินและข้อพิพาทของผู้ใช้งาน", "13", "13", "0", "ผ่าน"),
        ("10", "การแจ้งเตือน รางวัล และโปรไฟล์ของผู้ใช้งาน", "14", "14", "0", "ผ่าน"),
        ("รวม", "รวมผลการทดสอบทั้งหมด", "119", "119", "0", "ผ่าน"),
    ]
    add_table(
        doc,
        "ตารางที่ 4.1 ผลการทดสอบระบบ PingPay ตามกรณีทดสอบ",
        ["ลำดับ", "กรณีทดสอบระบบ", "จำนวนกรณี", "ผ่าน", "ไม่ผ่าน", "ผลการทดสอบ"],
        test_rows,
        [720, 4080, 1050, 950, 950, 1610],
        [
            WD_ALIGN_PARAGRAPH.CENTER,
            WD_ALIGN_PARAGRAPH.LEFT,
            WD_ALIGN_PARAGRAPH.CENTER,
            WD_ALIGN_PARAGRAPH.CENTER,
            WD_ALIGN_PARAGRAPH.CENTER,
            WD_ALIGN_PARAGRAPH.CENTER,
        ],
    )
    add_paragraph(
        doc,
        "จากตารางที่ 4.1 พบว่าระบบ PingPay ผ่านการทดสอบตามกรณีทดสอบทั้งหมด 119 กรณี คิดเป็นร้อยละ 100 "
        "โดยไม่พบกรณีทดสอบที่ไม่ผ่าน แสดงให้เห็นว่าฟังก์ชันหลักของระบบสามารถใช้งานได้ตามขอบเขตที่กำหนด",
    )

    add_heading(doc, "4.2 ผลการศึกษาความพึงพอใจของผู้ใช้งานระบบ PingPay")
    add_paragraph(
        doc,
        "หลังจากดำเนินการทดสอบระบบ ผู้จัดทำได้นำระบบไปให้กลุ่มตัวอย่างทดลองใช้งานและตอบแบบประเมินความพึงพอใจ "
        "จำนวน 20 คน โดยแบ่งเป็นผู้ใช้งานทั่วไป 10 คน และผู้ดูแลระบบหรือผู้ทดสอบด้านเทคนิค 10 คน "
        "แบบประเมินแบ่งออกเป็น 3 ตอน ได้แก่ ข้อมูลทั่วไปของผู้ตอบแบบสอบถาม ความพึงพอใจต่อการใช้งานระบบ และข้อเสนอแนะเพิ่มเติม",
    )
    add_paragraph(
        doc,
        "เกณฑ์การแปลผลค่าเฉลี่ยกำหนดเป็น 5 ระดับ ได้แก่ 4.50-5.00 หมายถึง พึงพอใจมากที่สุด, "
        "3.50-4.49 หมายถึง พึงพอใจมาก, 2.50-3.49 หมายถึง พึงพอใจปานกลาง, "
        "1.50-2.49 หมายถึง พึงพอใจน้อย และ 1.00-1.49 หมายถึง พึงพอใจน้อยที่สุด",
    )

    add_heading(doc, "ตอนที่ 1 ข้อมูลทั่วไปของผู้ตอบแบบสอบถาม")
    add_paragraph(doc, "ข้อมูลทั่วไปของผู้ตอบแบบสอบถามสามารถสรุปผลเป็นจำนวนและร้อยละได้ดังนี้")
    add_table(
        doc,
        "ตารางที่ 4.2 ตารางแสดงจำนวนและร้อยละของกลุ่มตัวอย่าง จำแนกตามกลุ่มผู้ทดสอบ",
        ["กลุ่มผู้ทดสอบ", "จำนวน", "ร้อยละ"],
        [
            ("ผู้ใช้งานทั่วไป", "10", "50.00"),
            ("ผู้ดูแลระบบหรือผู้ทดสอบด้านเทคนิค", "10", "50.00"),
            ("รวม", "20", "100.00"),
        ],
        [4960, 2200, 2200],
        [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER],
    )
    add_paragraph(
        doc,
        "จากตารางที่ 4.2 พบว่ากลุ่มตัวอย่างมีจำนวนทั้งหมด 20 คน แบ่งเป็นผู้ใช้งานทั่วไปจำนวน 10 คน คิดเป็นร้อยละ 50.00 "
        "และผู้ดูแลระบบหรือผู้ทดสอบด้านเทคนิคจำนวน 10 คน คิดเป็นร้อยละ 50.00",
    )

    add_table(
        doc,
        "ตารางที่ 4.3 ตารางแสดงจำนวนและร้อยละของกลุ่มตัวอย่าง จำแนกตามประสบการณ์การหารค่าใช้จ่ายร่วมกัน",
        ["ประสบการณ์การหารค่าใช้จ่ายร่วมกัน", "จำนวน", "ร้อยละ"],
        [
            ("ใช้งานหรือหารค่าใช้จ่ายร่วมกันเป็นประจำ", "12", "60.00"),
            ("ใช้งานหรือหารค่าใช้จ่ายร่วมกันเป็นบางครั้ง", "6", "30.00"),
            ("ไม่เคยใช้งานระบบลักษณะนี้มาก่อน", "2", "10.00"),
            ("รวม", "20", "100.00"),
        ],
        [4960, 2200, 2200],
        [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER],
    )
    add_paragraph(
        doc,
        "จากตารางที่ 4.3 พบว่ากลุ่มตัวอย่างส่วนใหญ่มีประสบการณ์หารค่าใช้จ่ายร่วมกันเป็นประจำ จำนวน 12 คน คิดเป็นร้อยละ 60.00 "
        "รองลงมาคือใช้งานเป็นบางครั้ง จำนวน 6 คน คิดเป็นร้อยละ 30.00 และไม่เคยใช้งานระบบลักษณะนี้มาก่อน จำนวน 2 คน คิดเป็นร้อยละ 10.00",
    )

    add_heading(doc, "ตอนที่ 2 ความพึงพอใจเกี่ยวกับระบบ PingPay")
    add_paragraph(doc, "ผลการประเมินความพึงพอใจของกลุ่มตัวอย่างที่มีต่อระบบ PingPay แสดงดังตารางที่ 4.4")
    satisfaction_rows = [
        ("1. ความถูกต้องของการเข้าสู่ระบบและการตั้งค่าโปรไฟล์", "4.55", "0.51", "มากที่สุด"),
        ("2. ความสะดวกในการสร้างบิลและแบ่งค่าใช้จ่ายร่วมกับเพื่อน", "4.60", "0.50", "มากที่สุด"),
        ("3. ความเหมาะสมของการอ่านใบเสร็จด้วย OCR และการกรอกบิลด้วยภาษาธรรมชาติ", "4.45", "0.60", "มาก"),
        ("4. ความชัดเจนของการติดตามสถานะหนี้และการชำระเงิน", "4.50", "0.61", "มากที่สุด"),
        ("5. ความครบถ้วนของการแจ้งเตือนและประวัติการใช้งาน", "4.35", "0.67", "มาก"),
        ("6. ความง่ายในการเข้าถึงเมนูหลักของแอปพลิเคชัน", "4.55", "0.51", "มากที่สุด"),
        ("7. ความชัดเจนของหน้าจอบิล เพื่อน การชำระเงิน และโปรไฟล์", "4.40", "0.68", "มาก"),
        ("8. ความเหมาะสมของสี ตัวอักษร และการจัดวางองค์ประกอบ", "4.45", "0.60", "มาก"),
        ("9. ความรวดเร็วในการตอบสนองของระบบ", "4.30", "0.73", "มาก"),
        ("10. ความเสถียรและความน่าเชื่อถือในการใช้งานโดยรวม", "4.25", "0.72", "มาก"),
        ("เฉลี่ยรวม", "4.44", "0.61", "มาก"),
    ]
    add_table(
        doc,
        "ตารางที่ 4.4 ผลการประเมินความพึงพอใจเกี่ยวกับระบบ PingPay",
        ["รายการประเมิน", "ค่าเฉลี่ย", "S.D.", "แปลผล"],
        satisfaction_rows,
        [5800, 1200, 1100, 1260],
        [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER],
    )
    add_paragraph(
        doc,
        "จากตารางที่ 4.4 พบว่าผู้ตอบแบบสอบถามมีความพึงพอใจต่อระบบ PingPay โดยภาพรวมอยู่ในระดับมาก "
        "มีค่าเฉลี่ยรวม 4.44 และส่วนเบี่ยงเบนมาตรฐาน 0.61 เมื่อพิจารณารายข้อพบว่ารายการที่มีค่าเฉลี่ยสูงสุดคือ "
        "ความสะดวกในการสร้างบิลและแบ่งค่าใช้จ่ายร่วมกับเพื่อน มีค่าเฉลี่ย 4.60 อยู่ในระดับมากที่สุด "
        "รองลงมาคือความถูกต้องของการเข้าสู่ระบบและการตั้งค่าโปรไฟล์ และความง่ายในการเข้าถึงเมนูหลักของแอปพลิเคชัน "
        "มีค่าเฉลี่ย 4.55 อยู่ในระดับมากที่สุด",
    )

    add_heading(doc, "ตอนที่ 3 ข้อเสนอแนะเพิ่มเติม")
    add_paragraph(doc, "ข้อเสนอแนะเพิ่มเติมจากกลุ่มตัวอย่างหลังทดลองใช้งานระบบ PingPay สรุปได้ดังตารางที่ 4.5")
    add_table(
        doc,
        "ตารางที่ 4.5 ตารางสรุปข้อเสนอแนะเพิ่มเติมจากผู้ตอบแบบสอบถาม",
        ["ข้อเสนอแนะ", "จำนวน", "ร้อยละ"],
        [
            ("ต้องการให้เพิ่มการแจ้งเตือนกำหนดชำระเงินล่วงหน้า", "7", "35.00"),
            ("ต้องการให้รองรับการตรวจสอบสลิปจากหลายธนาคารมากขึ้น", "5", "25.00"),
            ("ต้องการให้ OCR รองรับรูปแบบใบเสร็จที่หลากหลายมากขึ้น", "4", "20.00"),
            ("ต้องการให้มีรายงานสรุปค่าใช้จ่ายรายเดือนเพิ่มเติม", "3", "15.00"),
            ("ไม่มีข้อเสนอแนะเพิ่มเติม", "1", "5.00"),
            ("รวม", "20", "100.00"),
        ],
        [5600, 1800, 1960],
        [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER],
    )
    add_paragraph(
        doc,
        "จากตารางที่ 4.5 พบว่าข้อเสนอแนะที่กลุ่มตัวอย่างกล่าวถึงมากที่สุดคือการเพิ่มการแจ้งเตือนกำหนดชำระเงินล่วงหน้า "
        "รองลงมาคือการรองรับการตรวจสอบสลิปจากหลายธนาคาร และการปรับปรุง OCR ให้รองรับใบเสร็จหลากหลายรูปแบบมากขึ้น "
        "ซึ่งสามารถนำไปใช้เป็นแนวทางในการพัฒนาระบบในอนาคต",
    )

    add_heading(doc, "สรุปผลการดำเนินงาน")
    add_paragraph(
        doc,
        "ผลการดำเนินงานแสดงให้เห็นว่าระบบ PingPay สามารถสนับสนุนการจัดการค่าใช้จ่ายร่วมกันได้ครบถ้วน "
        "ตั้งแต่การสร้างบิล การเลือกเพื่อน การแบ่งยอด การชำระเงินผ่านพร้อมเพย์ การตรวจสอบสลิป การยื่นข้อพิพาท "
        "การแจ้งเตือน และการจัดการข้อมูลผ่านระบบผู้ดูแลระบบ โดยผลการทดสอบผ่านทุกกรณี และผลประเมินความพึงพอใจอยู่ในระดับมาก",
    )

    doc.save(OUTPUT_DOCX)
    print(f"Saved: {OUTPUT_DOCX}")
    print(f"Paragraphs: {len(doc.paragraphs)}")
    print(f"Tables: {len(doc.tables)}")


if __name__ == "__main__":
    main()
