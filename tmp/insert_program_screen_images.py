from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


DOCX = Path("docs/บทที่ 3_แก้ไขแล้ว_ตารางครบ.docx")
OUTPUT_DOCX = Path("docs/บทที่ 3_แก้ไขแล้ว_รูปครบ.docx")
PREVIEW = Path("preview")

SCREEN_IMAGES = [
    ("68_seed_home_loaded.png", "ภาพที่ 3.5 หน้าหลักของระบบ"),
    ("79_seed_my_bills_with_data.png", "ภาพที่ 3.6 หน้ารายการบิล"),
    ("80_seed_bill_detail_with_data.png", "ภาพที่ 3.7 หน้ารายละเอียดบิล"),
    ("81_seed_create_bill_initial.png", "ภาพที่ 3.8 หน้าสร้างบิลใหม่"),
    ("70_seed_payments_debts_with_data.png", "ภาพที่ 3.9 หน้าหนี้ที่ต้องชำระ"),
    ("72_seed_payments_receivables_with_data.png", "ภาพที่ 3.10 หน้ารายการเงินที่ต้องรับ"),
    ("77_seed_monthly_analytics_with_data.png", "ภาพที่ 3.11 หน้าสรุปรายเดือน"),
    ("78_seed_yearly_analytics_with_data.png", "ภาพที่ 3.12 หน้าสรุปรายปี"),
    ("46_friends_main_empty.png", "ภาพที่ 3.13 หน้าจัดการเพื่อน"),
    ("35_profile_my_qr_sheet.png", "ภาพที่ 3.14 หน้ารหัส QR โปรไฟล์"),
    ("30_profile_main_top.png", "ภาพที่ 3.15 หน้าโปรไฟล์ผู้ใช้"),
    ("36_profile_promptpay_setup_sheet.png", "ภาพที่ 3.16 หน้าตั้งค่า PromptPay"),
]


def set_cell_width(cell, width):
    tc_pr = cell._tc.get_or_add_tcPr()
    tcw = tc_pr.find(qn("w:tcW"))
    if tcw is None:
        tcw = OxmlElement("w:tcW")
        tc_pr.append(tcw)
    tcw.set(qn("w:w"), str(width))
    tcw.set(qn("w:type"), "dxa")


def set_cell_margins(cell, top=80, start=80, bottom=80, end=80):
    tc_pr = cell._tc.get_or_add_tcPr()
    mar = tc_pr.find(qn("w:tcMar"))
    if mar is None:
        mar = OxmlElement("w:tcMar")
        tc_pr.append(mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.autofit = False
    table.allow_autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_layout = tbl_pr.find(qn("w:tblLayout"))
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    if grid is None:
        grid = OxmlElement("w:tblGrid")
        table._tbl.insert(0, grid)
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            set_cell_width(cell, width)
            set_cell_margins(cell)


def clear_cell(cell):
    for paragraph in cell.paragraphs:
        for run in list(paragraph.runs):
            run._element.getparent().remove(run._element)
        paragraph.text = ""


def format_paragraph(paragraph, size=14, bold=False):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.line_spacing = 1.0
    for run in paragraph.runs:
        run.font.name = "TH Sarabun New"
        run._element.rPr.rFonts.set(qn("w:ascii"), "TH Sarabun New")
        run._element.rPr.rFonts.set(qn("w:hAnsi"), "TH Sarabun New")
        run._element.rPr.rFonts.set(qn("w:cs"), "TH Sarabun New")
        run.font.size = Pt(size)
        run.bold = bold


def main():
    doc = Document(DOCX)
    screen_tables = doc.tables[23:29]
    if len(screen_tables) != 6:
        raise RuntimeError(f"Expected 6 screen tables, found {len(screen_tables)}")

    idx = 0
    for table in screen_tables:
        set_table_geometry(table, [4680, 4680])
        for cell in table.rows[0].cells:
            image_name, _caption = SCREEN_IMAGES[idx]
            image_path = PREVIEW / image_name
            if not image_path.exists():
                raise FileNotFoundError(image_path)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            clear_cell(cell)
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            run = paragraph.add_run()
            run.add_picture(str(image_path), width=Inches(2.35))
            idx += 1

        for cell in table.rows[1].cells:
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            for paragraph in cell.paragraphs:
                format_paragraph(paragraph, size=14)

    output_path = DOCX
    try:
        doc.save(output_path)
    except PermissionError:
        output_path = OUTPUT_DOCX
        doc.save(output_path)

    doc = Document(output_path)
    print(f"Saved: {output_path}")
    print(f"Inline shapes: {len(doc.inline_shapes)}")
    print(f"Screen images inserted: {len(SCREEN_IMAGES)}")
    captions = []
    for table in doc.tables[23:29]:
        for cell in table.rows[1].cells:
            captions.append(cell.text.strip())
    missing_captions = [caption for _, caption in SCREEN_IMAGES if not any(caption in c for c in captions)]
    print(f"Missing expected captions: {missing_captions}")
    if missing_captions:
        raise SystemExit(2)


if __name__ == "__main__":
    main()
