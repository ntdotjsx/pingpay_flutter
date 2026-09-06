from pathlib import Path

from docx import Document


DOCX = Path("docs/บทที่ 3_แก้ไขแล้ว.docx")


def main() -> None:
    doc = Document(DOCX)
    print("paragraphs", len(doc.paragraphs), "tables", len(doc.tables), "sections", len(doc.sections))
    for i, p in enumerate(doc.paragraphs):
        text = p.text.strip()
        if "ตารางที่ 3." in text or "3.3.3" in text:
            print(f"P{i:03d}: {text}")
    for i, table in enumerate(doc.tables, start=1):
        first_rows = []
        for row in table.rows[:3]:
            first_rows.append(" | ".join(cell.text.replace("\n", " / ").strip()[:120] for cell in row.cells))
        print(f"T{i:02d} rows={len(table.rows)} cols={len(table.columns)}")
        for row in first_rows:
            print("   ", row)


if __name__ == "__main__":
    main()
