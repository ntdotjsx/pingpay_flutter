from pathlib import Path
import re

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


DOCX = Path("docs/บทที่ 3_แก้ไขแล้ว.docx")
FALLBACK_DOCX = Path("docs/บทที่ 3_แก้ไขแล้ว_ตารางครบ.docx")


def field(name, desc, dtype, key="", ref="", sample=""):
    return {
        "name": name,
        "desc": desc,
        "dtype": dtype,
        "key": key,
        "ref": ref,
        "sample": sample,
    }


TABLES = [
    {
        "no": "3.2",
        "caption": "ตารางข้อมูลอัตลักษณ์การเข้าสู่ระบบภายนอก",
        "table": "auth_identities",
        "rows": [
            field("id", "รหัสอัตลักษณ์การเข้าสู่ระบบ", "uuid", "PK, Default: gen_random_uuid()", "", "550e8400-e29b-41d4-a716-446655440201"),
            field("user_id", "รหัสผู้ใช้งานที่เชื่อมโยงกับอัตลักษณ์นี้", "uuid", "FK, Not Null, Unique", "users.id", "550e8400-e29b-41d4-a716-446655440001"),
            field("provider", "ชื่อผู้ให้บริการยืนยันตัวตนภายนอก", "varchar(32)", "Not Null, Default: google", "", "google"),
            field("provider_user_id", "รหัสผู้ใช้จากผู้ให้บริการภายนอก", "varchar(128)", "Not Null, Unique ร่วมกับ provider", "", "109876543210987654321"),
            field("created_at", "วันและเวลาที่สร้างข้อมูลอัตลักษณ์", "timestamp", "Not Null, Default: now()", "", "2026-09-01 08:05:00"),
            field("updated_at", "วันและเวลาที่แก้ไขข้อมูลอัตลักษณ์ล่าสุด", "timestamp", "Not Null, Default: now()", "", "2026-09-01 08:05:00"),
        ],
    },
    {
        "no": "3.3",
        "caption": "ตารางข้อมูลรหัสผ่าน PIN และความปลอดภัย",
        "table": "user_credentials",
        "rows": [
            field("user_id", "รหัสผู้ใช้งานเจ้าของข้อมูลความปลอดภัย", "uuid", "PK, FK", "users.id", "550e8400-e29b-41d4-a716-446655440001"),
            field("pin_hash", "รหัสผ่านตัวเลขที่ผ่านการเข้ารหัสแล้ว", "text", "", "", "$argon2id$v=19$m=65536..."),
            field("failed_attempts", "จำนวนครั้งที่กรอกรหัสผ่านตัวเลขผิด", "integer", "Not Null, Default: 0", "", "0"),
            field("locked_until", "วันและเวลาสิ้นสุดการระงับชั่วคราว", "timestamp", "", "", "2026-09-01 08:20:00"),
            field("updated_at", "วันและเวลาที่แก้ไขข้อมูลความปลอดภัยล่าสุด", "timestamp", "Not Null, Default: now()", "", "2026-09-01 08:10:00"),
        ],
    },
    {
        "no": "3.4",
        "caption": "ตารางข้อมูลเซสชันการใช้งานอุปกรณ์เดียว",
        "table": "auth_sessions",
        "rows": [
            field("id", "รหัสเซสชันการใช้งาน", "uuid", "PK, Default: gen_random_uuid()", "", "550e8400-e29b-41d4-a716-446655440301"),
            field("user_id", "รหัสผู้ใช้งานเจ้าของเซสชัน", "uuid", "FK, Not Null, Index", "users.id", "550e8400-e29b-41d4-a716-446655440001"),
            field("refresh_token_hash", "โทเคนต่ออายุที่ผ่านการเข้ารหัสแล้ว", "text", "Not Null", "", "$2b$12$uQz..."),
            field("device_info", "รายละเอียดอุปกรณ์ที่ใช้เข้าสู่ระบบ", "text", "", "", "Samsung Galaxy A55"),
            field("ip_address", "หมายเลขเครือข่ายของอุปกรณ์ขณะเข้าสู่ระบบ", "varchar(64)", "", "", "203.0.113.15"),
            field("expires_at", "วันและเวลาหมดอายุของเซสชัน", "timestamp", "Not Null", "", "2026-10-01 08:10:00"),
            field("created_at", "วันและเวลาที่สร้างเซสชัน", "timestamp", "Not Null, Default: now()", "", "2026-09-01 08:10:00"),
        ],
    },
    {
        "no": "3.5",
        "caption": "ตารางข้อมูลอุปกรณ์และโทเคนแจ้งเตือน",
        "table": "device_tokens",
        "rows": [
            field("id", "รหัสข้อมูลอุปกรณ์", "uuid", "PK, Default: gen_random_uuid()", "", "550e8400-e29b-41d4-a716-446655440401"),
            field("user_id", "รหัสผู้ใช้งานเจ้าของอุปกรณ์", "uuid", "FK, Not Null, Index", "users.id", "550e8400-e29b-41d4-a716-446655440001"),
            field("token", "โทเคนสำหรับส่งข้อความแจ้งเตือนไปยังอุปกรณ์", "text", "Not Null, Unique", "", "fcm_token_abc123"),
            field("platform", "แพลตฟอร์มของอุปกรณ์", "varchar(32)", "Not Null, Default: android", "", "android"),
            field("device_name", "ชื่ออุปกรณ์ที่ผู้ใช้มองเห็น", "varchar(128)", "", "", "โทรศัพท์ของสมชาย"),
            field("device_model", "รุ่นของอุปกรณ์", "varchar(128)", "", "", "Galaxy A55"),
            field("device_brand", "ยี่ห้อของอุปกรณ์", "varchar(64)", "", "", "Samsung"),
            field("os_version", "เวอร์ชันระบบปฏิบัติการของอุปกรณ์", "varchar(64)", "", "", "Android 15"),
            field("app_version", "เวอร์ชันแอปพลิเคชันที่ติดตั้งในอุปกรณ์", "varchar(64)", "", "", "1.0.3"),
            field("created_at", "วันและเวลาที่บันทึกอุปกรณ์", "timestamp", "Not Null, Default: now()", "", "2026-09-01 08:12:00"),
            field("updated_at", "วันและเวลาที่แก้ไขข้อมูลอุปกรณ์ล่าสุด", "timestamp", "Not Null, Default: now()", "", "2026-09-01 08:12:00"),
        ],
    },
    {
        "no": "3.6",
        "caption": "ตารางประวัติการยินยอมนโยบาย PDPA",
        "table": "consent_records",
        "rows": [
            field("id", "รหัสรายการยินยอม", "uuid", "PK, Default: gen_random_uuid()", "", "550e8400-e29b-41d4-a716-446655440501"),
            field("user_id", "รหัสผู้ใช้งานที่ให้ความยินยอม", "uuid", "FK, Not Null, Index", "users.id", "550e8400-e29b-41d4-a716-446655440001"),
            field("policy_version", "หมายเลขเวอร์ชันของนโยบายที่ยินยอม", "varchar(32)", "Not Null", "", "2026.09"),
            field("accepted_at", "วันและเวลาที่ผู้ใช้กดยินยอม", "timestamp", "Not Null, Default: now()", "", "2026-09-01 08:13:00"),
            field("ip_address", "หมายเลขเครือข่ายขณะกดยินยอม", "varchar(64)", "", "", "203.0.113.15"),
        ],
    },
    {
        "no": "3.7",
        "caption": "ตารางข้อมูลรหัสผ่านใช้ครั้งเดียว OTP",
        "table": "otp_verifications",
        "rows": [
            field("id", "รหัสรายการยืนยันรหัสใช้ครั้งเดียว", "uuid", "PK, Default: gen_random_uuid()", "", "550e8400-e29b-41d4-a716-446655440601"),
            field("user_id", "รหัสผู้ใช้งานที่ขอยืนยันตัวตน", "uuid", "FK, Not Null, Index", "users.id", "550e8400-e29b-41d4-a716-446655440001"),
            field("email", "อีเมลที่รับรหัสยืนยัน", "varchar(255)", "Not Null, Index", "", "somchai@example.com"),
            field("otp_hash", "รหัสใช้ครั้งเดียวที่ผ่านการเข้ารหัสแล้ว", "text", "Not Null", "", "$2b$12$otp..."),
            field("purpose", "วัตถุประสงค์ของการขอรหัสยืนยัน", "varchar(32)", "Not Null, Default: pin_reset", "", "pin_reset"),
            field("attempts", "จำนวนครั้งที่กรอกรหัสยืนยันแล้ว", "integer", "Not Null, Default: 0", "", "1"),
            field("max_attempts", "จำนวนครั้งสูงสุดที่อนุญาตให้กรอกรหัสยืนยัน", "integer", "Not Null, Default: 5", "", "5"),
            field("expires_at", "วันและเวลาที่รหัสยืนยันหมดอายุ", "timestamp", "Not Null", "", "2026-09-01 08:28:00"),
            field("verified_at", "วันและเวลาที่ยืนยันสำเร็จ", "timestamp", "", "", "2026-09-01 08:18:00"),
            field("created_at", "วันและเวลาที่สร้างรายการยืนยัน", "timestamp", "Not Null, Default: now()", "", "2026-09-01 08:13:00"),
            field("updated_at", "วันและเวลาที่แก้ไขรายการยืนยันล่าสุด", "timestamp", "Not Null, Default: now()", "", "2026-09-01 08:18:00"),
        ],
    },
    {
        "no": "3.8",
        "caption": "ตารางบันทึกเหตุการณ์ความปลอดภัย",
        "table": "security_events",
        "rows": [
            field("id", "รหัสเหตุการณ์ความปลอดภัย", "uuid", "PK, Default: gen_random_uuid()", "", "550e8400-e29b-41d4-a716-446655440701"),
            field("user_id", "รหัสผู้ใช้งานที่เกี่ยวข้องกับเหตุการณ์", "uuid", "FK, Index", "users.id", "550e8400-e29b-41d4-a716-446655440001"),
            field("event", "ชื่อเหตุการณ์ความปลอดภัยที่เกิดขึ้น", "varchar(64)", "Not Null", "", "pin_brute_force"),
            field("ip_address", "หมายเลขเครือข่ายขณะเกิดเหตุการณ์", "varchar(64)", "", "", "203.0.113.15"),
            field("metadata", "ข้อมูลประกอบเหตุการณ์แบบกึ่งโครงสร้าง", "jsonb", "", "", "{\"attempts\":6}"),
            field("created_at", "วันและเวลาที่บันทึกเหตุการณ์", "timestamp", "Not Null, Default: now()", "", "2026-09-01 08:19:00"),
        ],
    },
    {
        "no": "3.9",
        "caption": "ตารางข้อมูลความสัมพันธ์เพื่อน",
        "table": "friendships",
        "rows": [
            field("id", "รหัสความสัมพันธ์เพื่อน", "uuid", "PK, Default: gen_random_uuid()", "", "550e8400-e29b-41d4-a716-446655440801"),
            field("requester_id", "รหัสผู้ใช้งานที่ส่งคำขอเป็นเพื่อน", "uuid", "FK, Not Null, Unique ร่วมกับ addressee_id", "users.id", "550e8400-e29b-41d4-a716-446655440001"),
            field("addressee_id", "รหัสผู้ใช้งานที่ได้รับคำขอเป็นเพื่อน", "uuid", "FK, Not Null, Unique ร่วมกับ requester_id", "users.id", "550e8400-e29b-41d4-a716-446655440002"),
            field("status", "สถานะความสัมพันธ์ระหว่างผู้ใช้สองคน", "friend_status", "Not Null, Default: pending", "", "accepted"),
            field("created_at", "วันและเวลาที่สร้างคำขอเป็นเพื่อน", "timestamp", "Not Null, Default: now()", "", "2026-09-01 08:20:00"),
            field("responded_at", "วันและเวลาที่ตอบรับหรือปฏิเสธคำขอ", "timestamp", "", "", "2026-09-01 08:25:00"),
            field("removed_at", "วันและเวลาที่ลบความสัมพันธ์เพื่อน", "timestamp", "", "", ""),
        ],
    },
    {
        "no": "3.10",
        "caption": "ตารางข้อมูลบิลค่าใช้จ่ายและผลสแกน OCR",
        "table": "bills",
        "rows": [
            field("id", "รหัสบิลค่าใช้จ่าย", "uuid", "PK, Default: gen_random_uuid()", "", "550e8400-e29b-41d4-a716-446655440901"),
            field("owner_id", "รหัสผู้ใช้งานเจ้าของบิลหรือผู้สำรองจ่าย", "uuid", "FK, Not Null, Index", "users.id", "550e8400-e29b-41d4-a716-446655440001"),
            field("title", "ชื่อบิลที่ผู้ใช้ตั้งไว้", "varchar(128)", "", "", "อาหารเย็น"),
            field("currency", "สกุลเงินของบิล", "varchar(3)", "Not Null, Default: THB", "", "THB"),
            field("total_amount", "ยอดเงินรวมปัจจุบันของบิล", "numeric(12,2)", "Not Null", "", "1250.00"),
            field("original_total_amount", "ยอดเงินรวมตั้งต้นตอนสร้างบิล", "numeric(12,2)", "", "", "1250.00"),
            field("receipt_image_url", "ที่อยู่ไฟล์ภาพใบเสร็จ", "text", "", "", "https://cdn.pingpay.app/receipts/r001.jpg"),
            field("ocr_raw_data", "ข้อมูลดิบจากการอ่านข้อความบนใบเสร็จ", "jsonb", "", "", "{\"merchant\":\"ร้านอาหารดี\"}"),
            field("items_breakdown", "รายการแยกย่อยของสินค้า ภาษี และค่าบริการ", "jsonb", "", "", "{\"service\":125.00}"),
            field("status", "สถานะการชำระเงินของบิล", "bill_status", "Not Null, Default: unpaid", "", "unpaid"),
            field("created_at", "วันและเวลาที่สร้างบิล", "timestamp", "Not Null, Default: now()", "", "2026-09-01 18:30:00"),
            field("updated_at", "วันและเวลาที่แก้ไขบิลล่าสุด", "timestamp", "Not Null, Default: now()", "", "2026-09-01 18:35:00"),
            field("cancelled_at", "วันและเวลาที่ยกเลิกบิล", "timestamp", "", "", ""),
        ],
    },
    {
        "no": "3.11",
        "caption": "ตารางข้อมูลรายการหนี้รายบุคคลประจำบิล",
        "table": "bill_items",
        "rows": [
            field("id", "รหัสรายการหนี้รายบุคคล", "uuid", "PK, Default: gen_random_uuid()", "", "550e8400-e29b-41d4-a716-446655441001"),
            field("bill_id", "รหัสบิลที่รายการหนี้นี้สังกัด", "uuid", "FK, Not Null, Index", "bills.id", "550e8400-e29b-41d4-a716-446655440901"),
            field("debtor_id", "รหัสผู้ใช้งานที่เป็นลูกหนี้ในบิล", "uuid", "FK, Not Null, Index", "users.id", "550e8400-e29b-41d4-a716-446655440002"),
            field("original_amount", "ยอดหนี้ตั้งต้นของลูกหนี้รายนี้", "numeric(12,2)", "Not Null", "", "625.00"),
            field("current_amount", "ยอดหนี้ปัจจุบันหลังการแก้ไขหรือชำระบางส่วน", "numeric(12,2)", "Not Null", "", "625.00"),
            field("amount_paid", "ยอดเงินที่ชำระแล้ว", "numeric(12,2)", "Not Null, Default: 0", "", "0.00"),
            field("amount_written_off", "ยอดเงินที่เจ้าหนี้ยกหนี้ให้", "numeric(12,2)", "Not Null, Default: 0", "", "0.00"),
            field("status", "สถานะรายการหนี้ของลูกหนี้รายนี้", "bill_item_status", "Not Null, Default: unpaid", "", "unpaid"),
            field("is_acknowledged", "สถานะว่าลูกหนี้รับทราบรายการหนี้แล้วหรือไม่", "boolean", "Not Null, Default: false", "", "true"),
            field("acknowledged_at", "วันและเวลาที่ลูกหนี้กดรับทราบ", "timestamp", "", "", "2026-09-01 18:40:00"),
            field("is_locked", "สถานะล็อกการแก้ไขรายการหนี้", "boolean", "Not Null, Default: false", "", "false"),
            field("created_at", "วันและเวลาที่สร้างรายการหนี้", "timestamp", "Not Null, Default: now()", "", "2026-09-01 18:35:00"),
            field("updated_at", "วันและเวลาที่แก้ไขรายการหนี้ล่าสุด", "timestamp", "Not Null, Default: now()", "", "2026-09-01 18:40:00"),
        ],
    },
    {
        "no": "3.12",
        "caption": "ตารางข้อมูลการชำระเงินและสลิปโอนเงิน",
        "table": "payments",
        "rows": [
            field("id", "รหัสรายการชำระเงิน", "uuid", "PK, Default: gen_random_uuid()", "", "550e8400-e29b-41d4-a716-446655441101"),
            field("bill_item_id", "รหัสรายการหนี้ที่ชำระ", "uuid", "FK, Not Null, Index", "bill_items.id", "550e8400-e29b-41d4-a716-446655441001"),
            field("payer_id", "รหัสผู้ใช้งานที่เป็นผู้ชำระเงิน", "uuid", "FK, Not Null, Index", "users.id", "550e8400-e29b-41d4-a716-446655440002"),
            field("method", "รูปแบบการชำระเงิน", "payment_method", "Not Null", "", "full"),
            field("channel", "ช่องทางการชำระเงิน", "payment_channel", "Not Null, Default: promptpay_qr", "", "promptpay_qr"),
            field("amount", "ยอดเงินที่ชำระ", "numeric(12,2)", "Not Null", "", "625.00"),
            field("installment_number", "ลำดับงวดสำหรับการผ่อนชำระ", "integer", "", "", "1"),
            field("prompt_pay_qr_payload", "ข้อมูลสำหรับสร้างคิวอาร์พร้อมเพย์", "text", "", "", "0002010102112937..."),
            field("prompt_pay_qr_image_url", "ที่อยู่ไฟล์ภาพคิวอาร์พร้อมเพย์", "text", "", "", "https://cdn.pingpay.app/qr/p001.png"),
            field("prompt_pay_qr_generated_at", "วันและเวลาที่สร้างคิวอาร์พร้อมเพย์", "timestamp", "", "", "2026-09-01 18:42:00"),
            field("slip_image_url", "ที่อยู่ไฟล์ภาพสลิปโอนเงิน", "text", "", "", "https://cdn.pingpay.app/slips/s001.jpg"),
            field("slip_hash", "รหัสตรวจซ้ำของภาพสลิป", "varchar(64)", "Index", "", "9f86d081884c7d659a2feaa0c55ad015"),
            field("slip_ok_reference_id", "หมายเลขอ้างอิงจากผลตรวจสลิป", "varchar(128)", "Index", "", "SLIPOK-20260901-001"),
            field("slip_ok_verified_at", "วันและเวลาที่ตรวจสลิปสำเร็จ", "timestamp", "", "", "2026-09-01 18:45:00"),
            field("slip_ok_raw_response", "ข้อมูลตอบกลับดิบจากการตรวจสลิป", "jsonb", "", "", "{\"amount\":625.00}"),
            field("status", "สถานะการตรวจสอบและยืนยันการชำระเงิน", "payment_status", "Not Null, Default: pending_verification", "", "verified"),
            field("confirmed_by_owner_at", "วันและเวลาที่เจ้าของบิลยืนยันรายการ", "timestamp", "", "", "2026-09-01 18:50:00"),
            field("confirmed_by_owner_id", "รหัสเจ้าของบิลที่ยืนยันรายการ", "uuid", "FK", "users.id", "550e8400-e29b-41d4-a716-446655440001"),
            field("rejected_at", "วันและเวลาที่ปฏิเสธรายการ", "timestamp", "", "", ""),
            field("rejected_by_id", "รหัสผู้ใช้งานที่ปฏิเสธรายการ", "uuid", "FK", "users.id", ""),
            field("rejected_reason", "เหตุผลที่ปฏิเสธรายการชำระเงิน", "text", "", "", ""),
            field("created_at", "วันและเวลาที่สร้างรายการชำระเงิน", "timestamp", "Not Null, Default: now()", "", "2026-09-01 18:42:00"),
            field("updated_at", "วันและเวลาที่แก้ไขรายการชำระเงินล่าสุด", "timestamp", "Not Null, Default: now()", "", "2026-09-01 18:50:00"),
        ],
    },
    {
        "no": "3.13",
        "caption": "ตารางข้อมูลประวัติการตรวจสอบสลิป",
        "table": "payment_verifications",
        "rows": [
            field("id", "รหัสประวัติการตรวจสอบสลิป", "uuid", "PK, Default: gen_random_uuid()", "", "550e8400-e29b-41d4-a716-446655441201"),
            field("payment_id", "รหัสรายการชำระเงินที่ถูกตรวจสอบ", "uuid", "FK, Not Null, Index", "payments.id", "550e8400-e29b-41d4-a716-446655441101"),
            field("provider", "ชื่อผู้ให้บริการตรวจสอบสลิป", "varchar(32)", "Not Null, Default: easyslip", "", "easyslip"),
            field("status", "สถานะผลการตรวจสอบสลิป", "varchar(32)", "Not Null", "", "success"),
            field("provider_reference", "หมายเลขอ้างอิงจากผู้ให้บริการตรวจสอบสลิป", "varchar(128)", "Index", "", "ES-20260901-0001"),
            field("verified_amount", "ยอดเงินที่ตรวจพบจากสลิป", "numeric(12,2)", "", "", "625.00"),
            field("sender_info", "ข้อมูลผู้โอนเงินที่อ่านได้จากสลิป", "jsonb", "", "", "{\"name\":\"สมชาย\"}"),
            field("receiver_info", "ข้อมูลผู้รับเงินที่อ่านได้จากสลิป", "jsonb", "", "", "{\"name\":\"มานะ\"}"),
            field("failure_code", "รหัสข้อผิดพลาดเมื่อการตรวจสอบไม่สำเร็จ", "varchar(64)", "", "", ""),
            field("failure_message", "ข้อความอธิบายข้อผิดพลาดเป็นภาษาไทย", "text", "", "", ""),
            field("raw_response", "ข้อมูลตอบกลับดิบจากผู้ให้บริการตรวจสอบสลิป", "jsonb", "", "", "{\"status\":\"success\"}"),
            field("created_at", "วันและเวลาที่บันทึกผลการตรวจสอบ", "timestamp", "Not Null, Default: now()", "", "2026-09-01 18:45:00"),
        ],
    },
    {
        "no": "3.14",
        "caption": "ตารางสมุดบัญชีแยกประเภทธุรกรรมทางการเงิน",
        "table": "financial_transactions",
        "rows": [
            field("id", "รหัสธุรกรรมทางการเงิน", "uuid", "PK, Default: gen_random_uuid()", "", "550e8400-e29b-41d4-a716-446655441301"),
            field("bill_id", "รหัสบิลที่เกี่ยวข้องกับธุรกรรม", "uuid", "FK, Not Null, Index", "bills.id", "550e8400-e29b-41d4-a716-446655440901"),
            field("bill_item_id", "รหัสรายการหนี้ที่เกี่ยวข้องกับธุรกรรม", "uuid", "FK, Not Null, Index", "bill_items.id", "550e8400-e29b-41d4-a716-446655441001"),
            field("type", "ประเภทความเคลื่อนไหวทางการเงิน", "transaction_type", "Not Null, Index", "", "payment"),
            field("amount", "จำนวนเงินของธุรกรรม", "numeric(12,2)", "Not Null", "", "625.00"),
            field("currency", "สกุลเงินของธุรกรรม", "varchar(3)", "Not Null, Default: THB", "", "THB"),
            field("reference_id", "รหัสอ้างอิงรายการต้นทาง", "varchar(128)", "", "", "PAY-20260901-001"),
            field("created_by_id", "รหัสผู้ใช้งานที่สร้างธุรกรรม", "uuid", "FK, Not Null", "users.id", "550e8400-e29b-41d4-a716-446655440002"),
            field("metadata", "ข้อมูลประกอบธุรกรรมแบบกึ่งโครงสร้าง", "jsonb", "", "", "{\"source\":\"payment\"}"),
            field("created_at", "วันและเวลาที่บันทึกธุรกรรม", "timestamp", "Not Null, Default: now()", "", "2026-09-01 18:50:00"),
        ],
    },
    {
        "no": "3.15",
        "caption": "ตารางประวัติการแก้ไขบิลและปรับปรุงยอดหนี้",
        "table": "edit_logs",
        "rows": [
            field("id", "รหัสประวัติการแก้ไข", "uuid", "PK, Default: gen_random_uuid()", "", "550e8400-e29b-41d4-a716-446655441401"),
            field("action", "ประเภทการแก้ไขที่เกิดขึ้น", "edit_action", "Not Null", "", "bill_item_edited"),
            field("bill_id", "รหัสบิลที่ถูกแก้ไข", "uuid", "FK, Index", "bills.id", "550e8400-e29b-41d4-a716-446655440901"),
            field("bill_item_id", "รหัสรายการหนี้ที่ถูกแก้ไข", "uuid", "FK", "bill_items.id", "550e8400-e29b-41d4-a716-446655441001"),
            field("performed_by_id", "รหัสผู้ใช้งานที่ดำเนินการแก้ไข", "uuid", "FK, Not Null, Index", "users.id", "550e8400-e29b-41d4-a716-446655440001"),
            field("affected_user_id", "รหัสผู้ใช้งานที่ได้รับผลจากการแก้ไข", "uuid", "FK", "users.id", "550e8400-e29b-41d4-a716-446655440002"),
            field("previous_value", "ข้อมูลเดิมก่อนการแก้ไข", "jsonb", "", "", "{\"current_amount\":650.00}"),
            field("new_value", "ข้อมูลใหม่หลังการแก้ไข", "jsonb", "", "", "{\"current_amount\":625.00}"),
            field("note", "หมายเหตุหรือเหตุผลประกอบการแก้ไข", "text", "", "", "ปรับยอดตามใบเสร็จจริง"),
            field("notified_at", "วันและเวลาที่แจ้งผู้เกี่ยวข้องแล้ว", "timestamp", "", "", "2026-09-01 18:55:00"),
            field("created_at", "วันและเวลาที่บันทึกประวัติการแก้ไข", "timestamp", "Not Null, Default: now()", "", "2026-09-01 18:55:00"),
        ],
    },
    {
        "no": "3.16",
        "caption": "ตารางข้อมูลข้อพิพาททางการเงินและหลักฐานสองฝ่าย",
        "table": "disputes",
        "rows": [
            field("id", "รหัสข้อพิพาท", "uuid", "PK, Default: gen_random_uuid()", "", "550e8400-e29b-41d4-a716-446655441501"),
            field("bill_item_id", "รหัสรายการหนี้ที่ถูกโต้แย้ง", "uuid", "FK, Not Null, Index", "bill_items.id", "550e8400-e29b-41d4-a716-446655441001"),
            field("raised_by_id", "รหัสผู้ใช้งานที่ยื่นข้อพิพาท", "uuid", "FK, Not Null", "users.id", "550e8400-e29b-41d4-a716-446655440002"),
            field("reason", "เหตุผลที่ผู้ใช้ยื่นข้อพิพาท", "text", "Not Null", "", "ยอดเงินไม่ตรงกับรายการอาหารที่สั่ง"),
            field("evidence_url", "ที่อยู่ไฟล์หลักฐานจากผู้ยื่นข้อพิพาท", "text", "", "", "https://cdn.pingpay.app/evidence/e001.jpg"),
            field("status", "สถานะการพิจารณาข้อพิพาท", "dispute_status", "Not Null, Default: open", "", "open"),
            field("creditor_evidence_note", "คำชี้แจงจากเจ้าหนี้", "text", "", "", "แนบใบเสร็จและรายการหารค่าใช้จ่าย"),
            field("creditor_evidence_url", "ที่อยู่ไฟล์หลักฐานจากเจ้าหนี้", "text", "", "", "https://cdn.pingpay.app/evidence/e002.jpg"),
            field("creditor_responded_at", "วันและเวลาที่เจ้าหนี้ตอบกลับ", "timestamp", "", "", "2026-09-02 09:00:00"),
            field("resolved_by_id", "รหัสผู้ดูแลระบบที่ตัดสินข้อพิพาท", "uuid", "FK", "users.id", "550e8400-e29b-41d4-a716-446655440099"),
            field("resolution_note", "ผลสรุปและเหตุผลการตัดสิน", "text", "", "", "ตรวจสอบหลักฐานแล้วให้ปรับยอด"),
            field("resolved_at", "วันและเวลาที่ตัดสินข้อพิพาทเสร็จสิ้น", "timestamp", "", "", "2026-09-02 10:00:00"),
            field("created_at", "วันและเวลาที่ยื่นข้อพิพาท", "timestamp", "Not Null, Default: now()", "", "2026-09-01 19:00:00"),
        ],
    },
    {
        "no": "3.17",
        "caption": "ตารางบันทึกประวัติการปฏิบัติงานของผู้ดูแลระบบ",
        "table": "admin_action_logs",
        "rows": [
            field("id", "รหัสบันทึกการปฏิบัติงานของผู้ดูแลระบบ", "uuid", "PK, Default: gen_random_uuid()", "", "550e8400-e29b-41d4-a716-446655441601"),
            field("admin_id", "รหัสผู้ดูแลระบบที่ดำเนินการ", "uuid", "FK, Not Null, Index", "users.id", "550e8400-e29b-41d4-a716-446655440099"),
            field("action_type", "ประเภทการปฏิบัติงานของผู้ดูแลระบบ", "varchar(64)", "Not Null", "", "resolve_dispute"),
            field("target_user_id", "รหัสผู้ใช้งานที่เป็นเป้าหมายของการดำเนินการ", "uuid", "FK, Index", "users.id", "550e8400-e29b-41d4-a716-446655440002"),
            field("reason", "เหตุผลประกอบการดำเนินการ", "text", "", "", "ตัดสินข้อพิพาทตามหลักฐาน"),
            field("metadata", "ข้อมูลประกอบการดำเนินการแบบกึ่งโครงสร้าง", "jsonb", "", "", "{\"dispute_id\":\"...\"}"),
            field("created_at", "วันและเวลาที่บันทึกการดำเนินการ", "timestamp", "Not Null, Default: now()", "", "2026-09-02 10:00:00"),
        ],
    },
    {
        "no": "3.18",
        "caption": "ตารางบันทึกพฤติกรรมน่าสงสัยและการทุจริต",
        "table": "suspicious_activity_logs",
        "rows": [
            field("id", "รหัสบันทึกพฤติกรรมน่าสงสัย", "uuid", "PK, Default: gen_random_uuid()", "", "550e8400-e29b-41d4-a716-446655441701"),
            field("user_id", "รหัสผู้ใช้งานที่เกี่ยวข้องกับพฤติกรรม", "uuid", "FK, Index", "users.id", "550e8400-e29b-41d4-a716-446655440002"),
            field("type", "ประเภทพฤติกรรมที่เข้าข่ายผิดปกติ", "varchar(64)", "Not Null, Index", "", "duplicate_slip"),
            field("description", "รายละเอียดพฤติกรรมที่ตรวจพบ", "text", "Not Null", "", "พบการใช้สลิปใบเดิมมากกว่าหนึ่งรายการ"),
            field("metadata", "ข้อมูลประกอบการตรวจจับแบบกึ่งโครงสร้าง", "jsonb", "", "", "{\"payment_id\":\"...\"}"),
            field("created_at", "วันและเวลาที่บันทึกพฤติกรรม", "timestamp", "Not Null, Default: now()", "", "2026-09-02 11:00:00"),
        ],
    },
    {
        "no": "3.19",
        "caption": "ตารางบันทึกกิจกรรมการใช้งานทั่วไป",
        "table": "activity_logs",
        "rows": [
            field("id", "รหัสบันทึกกิจกรรม", "uuid", "PK, Default: gen_random_uuid()", "", "550e8400-e29b-41d4-a716-446655441801"),
            field("user_id", "รหัสผู้ใช้งานที่ทำกิจกรรม", "uuid", "FK, Index", "users.id", "550e8400-e29b-41d4-a716-446655440001"),
            field("action", "ชื่อกิจกรรมที่ผู้ใช้ทำในระบบ", "varchar(64)", "Not Null", "", "bill_created"),
            field("metadata", "ข้อมูลประกอบกิจกรรมแบบกึ่งโครงสร้าง", "jsonb", "", "", "{\"bill_id\":\"...\"}"),
            field("created_at", "วันและเวลาที่เกิดกิจกรรม", "timestamp", "Not Null, Default: now()", "", "2026-09-01 18:30:00"),
        ],
    },
    {
        "no": "3.20",
        "caption": "ตารางคิวงานส่งข้อความแจ้งเตือนระดับระบบ",
        "table": "notification_outbox",
        "rows": [
            field("id", "รหัสงานแจ้งเตือนในคิว", "uuid", "PK, Default: gen_random_uuid()", "", "550e8400-e29b-41d4-a716-446655441901"),
            field("event_type", "ประเภทเหตุการณ์ที่ทำให้ต้องส่งแจ้งเตือน", "varchar(64)", "Not Null", "", "payment_verified"),
            field("recipient_user_id", "รหัสผู้ใช้งานผู้รับแจ้งเตือน", "uuid", "FK, Not Null, Index", "users.id", "550e8400-e29b-41d4-a716-446655440001"),
            field("channel", "ช่องทางการส่งแจ้งเตือน", "varchar(32)", "Not Null, Default: fcm", "", "fcm"),
            field("payload", "เนื้อหาข้อความแจ้งเตือนแบบกึ่งโครงสร้าง", "jsonb", "Not Null", "", "{\"title\":\"ได้รับเงินแล้ว\"}"),
            field("deduplication_key", "รหัสป้องกันการสร้างงานแจ้งเตือนซ้ำ", "varchar(128)", "Not Null, Unique", "", "payment_verified:550e8400"),
            field("status", "สถานะงานแจ้งเตือนในคิว", "varchar(32)", "Not Null, Default: PENDING", "Index ร่วมกับ available_at", "PENDING"),
            field("attempts", "จำนวนครั้งที่พยายามส่งแล้ว", "integer", "Not Null, Default: 0", "", "0"),
            field("max_attempts", "จำนวนครั้งสูงสุดที่อนุญาตให้พยายามส่ง", "integer", "Not Null, Default: 5", "", "5"),
            field("available_at", "วันและเวลาที่พร้อมให้ระบบเริ่มส่ง", "timestamp", "Not Null, Default: now()", "Index ร่วมกับ status", "2026-09-01 18:51:00"),
            field("locked_at", "วันและเวลาที่งานถูกล็อกเพื่อประมวลผล", "timestamp", "", "", "2026-09-01 18:52:00"),
            field("sent_at", "วันและเวลาที่ส่งแจ้งเตือนสำเร็จ", "timestamp", "", "", "2026-09-01 18:52:05"),
            field("failed_at", "วันและเวลาที่ส่งแจ้งเตือนไม่สำเร็จ", "timestamp", "", "", ""),
            field("last_error", "ข้อความข้อผิดพลาดล่าสุดจากการส่งแจ้งเตือน", "text", "", "", ""),
            field("created_at", "วันและเวลาที่สร้างงานแจ้งเตือน", "timestamp", "Not Null, Default: now()", "", "2026-09-01 18:51:00"),
            field("updated_at", "วันและเวลาที่แก้ไขงานแจ้งเตือนล่าสุด", "timestamp", "Not Null, Default: now()", "", "2026-09-01 18:52:05"),
        ],
    },
    {
        "no": "3.21",
        "caption": "ตารางบันทึกผลการส่งมอบการแจ้งเตือนไปยัง Firebase",
        "table": "notification_deliveries",
        "rows": [
            field("id", "รหัสผลการส่งมอบแจ้งเตือน", "uuid", "PK, Default: gen_random_uuid()", "", "550e8400-e29b-41d4-a716-446655442001"),
            field("notification_id", "รหัสงานแจ้งเตือนที่นำไปส่ง", "uuid", "FK, Not Null, Index", "notification_outbox.id", "550e8400-e29b-41d4-a716-446655441901"),
            field("provider", "ชื่อผู้ให้บริการส่งแจ้งเตือน", "varchar(32)", "Not Null, Default: fcm", "", "fcm"),
            field("recipient_target", "ปลายทางอุปกรณ์หรือผู้รับที่ใช้ส่งแจ้งเตือน", "varchar(256)", "", "", "device_token_abc123"),
            field("status", "สถานะผลการส่งมอบแจ้งเตือน", "varchar(32)", "Not Null", "", "SENT"),
            field("attempt_number", "ลำดับครั้งที่พยายามส่ง", "integer", "Not Null", "", "1"),
            field("response_payload", "ข้อมูลตอบกลับจากผู้ให้บริการส่งแจ้งเตือน", "jsonb", "", "", "{\"message_id\":\"0:123\"}"),
            field("error_message", "ข้อความข้อผิดพลาดจากการส่งแจ้งเตือน", "text", "", "", ""),
            field("created_at", "วันและเวลาที่บันทึกผลการส่งมอบ", "timestamp", "Not Null, Default: now()", "", "2026-09-01 18:52:05"),
        ],
    },
    {
        "no": "3.22",
        "caption": "ตารางแคตตาล็อกสินค้าของรางวัล",
        "table": "reward_items",
        "rows": [
            field("id", "รหัสสินค้าของรางวัล", "uuid", "PK, Default: gen_random_uuid()", "", "550e8400-e29b-41d4-a716-446655442101"),
            field("title", "ชื่อสินค้าของรางวัล", "varchar(128)", "Not Null", "", "แก้วน้ำพิงเพย์"),
            field("description", "รายละเอียดสินค้าของรางวัล", "text", "", "", "แก้วน้ำเก็บความเย็นพร้อมโลโก้"),
            field("points_cost", "จำนวนคะแนนที่ต้องใช้แลก", "integer", "Not Null", "", "250"),
            field("category", "หมวดหมู่ของสินค้าของรางวัล", "varchar(64)", "Not Null, Default: physical", "", "physical"),
            field("image_url", "ที่อยู่ไฟล์ภาพสินค้าของรางวัล", "text", "", "", "https://cdn.pingpay.app/rewards/cup.png"),
            field("in_stock", "จำนวนสินค้าคงเหลือ", "integer", "Not Null, Default: 100", "", "80"),
            field("is_active", "สถานะการเปิดให้แลกสินค้าของรางวัล", "boolean", "Not Null, Default: true", "", "true"),
            field("created_at", "วันและเวลาที่สร้างรายการของรางวัล", "timestamp", "Not Null, Default: now()", "", "2026-09-01 09:00:00"),
            field("updated_at", "วันและเวลาที่แก้ไขรายการของรางวัลล่าสุด", "timestamp", "Not Null, Default: now()", "", "2026-09-01 09:10:00"),
        ],
    },
    {
        "no": "3.23",
        "caption": "ตารางประวัติการแลกของรางวัลและการจัดส่ง",
        "table": "reward_redemptions",
        "rows": [
            field("id", "รหัสรายการแลกของรางวัล", "uuid", "PK, Default: gen_random_uuid()", "", "550e8400-e29b-41d4-a716-446655442201"),
            field("user_id", "รหัสผู้ใช้งานที่แลกของรางวัล", "uuid", "FK, Not Null, Index", "users.id", "550e8400-e29b-41d4-a716-446655440001"),
            field("reward_item_id", "รหัสสินค้าของรางวัลที่ถูกแลก", "uuid", "FK, Not Null", "reward_items.id", "550e8400-e29b-41d4-a716-446655442101"),
            field("points_spent", "จำนวนคะแนนที่ถูกหักจากผู้ใช้", "integer", "Not Null", "", "250"),
            field("status", "สถานะการจัดส่งของรางวัล", "reward_redemption_status", "Not Null, Default: pending_delivery", "", "pending_delivery"),
            field("recipient_name", "ชื่อผู้รับพัสดุ", "varchar(128)", "Not Null", "", "สมชาย ใจดี"),
            field("phone_number", "หมายเลขโทรศัพท์ผู้รับพัสดุ", "varchar(32)", "Not Null", "", "0812345678"),
            field("shipping_address", "ที่อยู่สำหรับจัดส่งพัสดุ", "text", "Not Null", "", "99/1 ถนนสุขุมวิท แขวงพระโขนง กรุงเทพฯ 10110"),
            field("tracking_number", "หมายเลขติดตามพัสดุ", "varchar(64)", "", "", "TH1234567890"),
            field("created_at", "วันและเวลาที่สร้างรายการแลกของรางวัล", "timestamp", "Not Null, Default: now()", "", "2026-09-01 09:15:00"),
            field("updated_at", "วันและเวลาที่แก้ไขรายการแลกของรางวัลล่าสุด", "timestamp", "Not Null, Default: now()", "", "2026-09-01 09:20:00"),
        ],
    },
]


def para_text(element):
    return "".join(t.text or "" for t in element.iter(qn("w:t"))).strip()


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, width):
    tc_pr = cell._tc.get_or_add_tcPr()
    tcw = tc_pr.find(qn("w:tcW"))
    if tcw is None:
        tcw = OxmlElement("w:tcW")
        tc_pr.append(tcw)
    tcw.set(qn("w:w"), str(width))
    tcw.set(qn("w:type"), "dxa")


def set_cell_margins(cell, top=60, start=80, bottom=60, end=80):
    tc_pr = cell._tc.get_or_add_tcPr()
    mar = tc_pr.find(qn("w:tcMar"))
    if mar is None:
        mar = OxmlElement("w:tcMar")
        tc_pr.append(mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_layout = tbl_pr.find(qn("w:tblLayout"))
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    if grid is None:
        grid = OxmlElement("w:tblGrid")
        table._tbl.insert(0, grid)
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            set_cell_width(cell, width)


def repeat_header_row(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = tr_pr.find(qn("w:tblHeader"))
    if tbl_header is None:
        tbl_header = OxmlElement("w:tblHeader")
        tr_pr.append(tbl_header)
    tbl_header.set(qn("w:val"), "true")


def format_cell(cell, *, bold=False, header=False, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    set_cell_margins(cell)
    if header:
        set_cell_shading(cell, "D9EAF7")
    for paragraph in cell.paragraphs:
        paragraph.alignment = align
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.0
        for run in paragraph.runs:
            run.font.name = "TH Sarabun New"
            run._element.rPr.rFonts.set(qn("w:ascii"), "TH Sarabun New")
            run._element.rPr.rFonts.set(qn("w:hAnsi"), "TH Sarabun New")
            run._element.rPr.rFonts.set(qn("w:cs"), "TH Sarabun New")
            run.font.size = Pt(12)
            run.bold = bold


def format_caption(paragraph, template=None):
    if template is not None:
        try:
            paragraph.style = template.style
        except Exception:
            pass
        paragraph.alignment = template.alignment
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.keep_with_next = True
    for run in paragraph.runs:
        run.font.name = "TH Sarabun New"
        run._element.rPr.rFonts.set(qn("w:ascii"), "TH Sarabun New")
        run._element.rPr.rFonts.set(qn("w:hAnsi"), "TH Sarabun New")
        run._element.rPr.rFonts.set(qn("w:cs"), "TH Sarabun New")
        run.font.size = Pt(16)


def build_table(doc, headers, spec, users_table_style):
    caption = doc.add_paragraph(f"ตารางที่ {spec['no']} {spec['caption']} ({spec['table']})")
    table = doc.add_table(rows=1, cols=len(headers))
    try:
        table.style = users_table_style
    except Exception:
        table.style = "Table Grid"
    table.autofit = False
    table.allow_autofit = False

    hdr_cells = table.rows[0].cells
    for cell, header in zip(hdr_cells, headers):
        cell.text = header
        format_cell(cell, bold=True, header=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    repeat_header_row(table.rows[0])

    for row in spec["rows"]:
        cells = table.add_row().cells
        values = [row["name"], row["desc"], row["dtype"], row["key"], row["ref"], row["sample"]]
        for idx, (cell, value) in enumerate(zip(cells, values)):
            cell.text = value
            align = WD_ALIGN_PARAGRAPH.LEFT
            if idx in (2, 3, 4):
                align = WD_ALIGN_PARAGRAPH.CENTER
            format_cell(cell, align=align)

    widths = [1350, 2500, 1250, 1700, 1300, 1500]
    set_table_geometry(table, widths)
    return caption, table


def remove_element(element):
    parent = element.getparent()
    if parent is not None:
        parent.remove(element)


def main():
    doc = Document(DOCX)
    body = doc._body._element

    users_table = None
    for table in doc.tables:
        headers = [cell.text.strip() for cell in table.rows[0].cells]
        if len(headers) == 6 and headers[0].replace("\n", " ").strip().startswith("Attribute"):
            users_table = table
            break
    if users_table is None:
        raise RuntimeError("Cannot find table 3.1 data dictionary table.")

    actual_caption = None
    for child in body:
        if child.tag == qn("w:p") and "ตารางที่ 3.1 ตารางข้อมูลผู้ใช้งานและโปรไฟล์ (users)" in para_text(child):
            actual_caption = child
            break

    remove_started = False
    old_caption_pat = re.compile(r"^ตารางที่ 3\.(?:1|[2-9]|1[0-9]|2[0-3])\b")
    for child in list(body):
        if child is actual_caption:
            remove_started = False
            continue
        if child.tag == qn("w:p"):
            text = para_text(child)
            if text == "ตารางที่ 3.1 ตารางข้อมูลผู้ใช้งานและโปรไฟล์":
                remove_started = True
            if remove_started and old_caption_pat.match(text):
                remove_element(child)
                continue
            if remove_started and text.startswith("ตารางที่ 3.23 "):
                remove_started = False

    # Remove any previously generated data dictionary tables after table 3.1.
    seen_users_table = False
    for child in list(body):
        if child is users_table._tbl:
            seen_users_table = True
            continue
        if not seen_users_table:
            continue
        if child.tag == qn("w:p") and para_text(child).startswith("3.3.4"):
            break
        remove_element(child)

    headers = [cell.text for cell in users_table.rows[0].cells]
    users_style = users_table.style
    caption_template = None
    if actual_caption is not None:
        from docx.text.paragraph import Paragraph
        caption_template = Paragraph(actual_caption, doc)

    new_elements = []
    for spec in TABLES:
        caption, table = build_table(doc, headers, spec, users_style)
        format_caption(caption, caption_template)
        spacer = doc.add_paragraph("")
        spacer.paragraph_format.space_after = Pt(3)
        new_elements.extend([caption._element, table._tbl, spacer._element])

    anchor = users_table._tbl
    for element in reversed(new_elements):
        remove_element(element)
        anchor.addnext(element)

    output_path = DOCX
    try:
        doc.save(output_path)
    except PermissionError:
        output_path = FALLBACK_DOCX
        doc.save(output_path)

    doc = Document(output_path)
    captions = [p.text.strip() for p in doc.paragraphs if p.text.strip().startswith("ตารางที่ 3.")]
    inserted = [c for c in captions if re.match(r"^ตารางที่ 3\.(?:[2-9]|1[0-9]|2[0-3])\b", c)]
    desc_violations = []
    for table in doc.tables[1:23]:
        if len(table.columns) < 2:
            continue
        for row_idx, row in enumerate(table.rows[1:], start=2):
            desc = row.cells[1].text.strip()
            if re.search(r"[A-Za-z]", desc):
                desc_violations.append((row_idx, desc))
    print(f"Saved: {output_path}")
    print(f"Total tables: {len(doc.tables)}")
    print(f"Inserted captions 3.2-3.23: {len(inserted)}")
    print(f"Description ASCII-letter violations: {len(desc_violations)}")
    if desc_violations:
        for item in desc_violations[:10]:
            print(item)
        raise SystemExit(2)


if __name__ == "__main__":
    main()
