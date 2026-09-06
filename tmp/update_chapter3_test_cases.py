from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


INPUT_DOCX = Path("docs/บทที่ 3_แก้ไขแล้ว_รูปครบ.docx")
OUTPUT_DOCX = Path("docs/บทที่ 3_แก้ไขแล้ว_TestCase.docx")


HEADERS = [
    "Pre-condition",
    "กรณีทดสอบ\n(Test Case)",
    "ขั้นตอนการทดสอบ\n(Test Step)",
    "Test Data",
    "Expect Results",
    "สถานะ",
    "หมายเหตุ",
]


TABLE_SPECS = [
    {
        "caption": "ตารางที่ 3.4 กรณีทดสอบเข้าสู่ระบบของผู้ดูแลระบบ",
        "pre": "เปิดใช้งานหน้าเข้าสู่ระบบของผู้ดูแลระบบ PingPay Developer Console",
        "rows": [
            (
                "เข้าสู่ระบบด้วยบัญชีผู้ดูแลระบบสำเร็จ",
                "1. เปิดหน้า /login\n2. กดปุ่มเข้าสู่ระบบด้วยบัญชี Google\n3. ระบบรับรหัสยืนยันจาก Google\n4. ระบบตรวจสอบสิทธิ์ผู้ใช้\n5. ระบบนำไปยัง Dashboard",
                "บัญชี Google ที่มี role = developer",
                "เข้าสู่ระบบสำเร็จ บันทึก admin_token และแสดงหน้า Dashboard",
                "ตรวจจากโค้ด",
                "อ้างอิง login/+page.svelte และ verifyGoogleToken",
            ),
            (
                "รับ token จาก callback สำเร็จ",
                "1. เปิด /login พร้อมพารามิเตอร์ token\n2. อ่าน token จาก URL\n3. บันทึก token\n4. เปลี่ยนเส้นทางไปหน้าหลัก",
                "/login?token=<JWT>",
                "ระบบบันทึก token และเข้าสู่หน้าหลักโดยไม่ต้องกรอกซ้ำ",
                "ตรวจจากโค้ด",
                "อ้างอิง onMount ในหน้า login",
            ),
            (
                "เข้าสู่ระบบด้วย token โดยตรงสำเร็จ",
                "1. เปิดเมนูขั้นสูง\n2. วาง access token\n3. กด Sign In with Token\n4. ระบบบันทึก token\n5. ไปยัง Dashboard",
                "JWT ของผู้ดูแลระบบ",
                "เข้าสู่ระบบสำเร็จและแสดงเมนูผู้ดูแลระบบ",
                "ตรวจจากโค้ด",
                "อ้างอิง handleManualTokenLogin",
            ),
            (
                "ไม่กรอก token ในช่องขั้นสูง",
                "1. เปิดเมนูขั้นสูง\n2. ปล่อยช่อง token ว่าง\n3. กด Sign In with Token",
                "ค่าว่าง",
                "ระบบแสดงข้อความแจ้งเตือนให้กรอก access token",
                "ตรวจจากโค้ด",
                "มี validation ก่อนบันทึก token",
            ),
            (
                "Google ไม่ส่ง credential กลับมา",
                "1. กดเข้าสู่ระบบด้วย Google\n2. ระบบได้รับ response ที่ไม่มี credential\n3. ตรวจสอบข้อความ error",
                "credential ว่าง",
                "ระบบแสดงข้อความแจ้งเตือนว่าไม่มีข้อมูล credential",
                "ตรวจจากโค้ด",
                "อ้างอิง handleCredentialResponse",
            ),
            (
                "บัญชีไม่ใช่ผู้ดูแลระบบ",
                "1. เข้าสู่ระบบด้วยบัญชีผู้ใช้ทั่วไป\n2. ระบบตรวจสอบ role\n3. ตรวจสอบผลลัพธ์",
                "role = user",
                "ระบบปฏิเสธการเข้าใช้งาน Developer Console",
                "ตรวจจากโค้ด",
                "verifyGoogleToken ตรวจ role ต้องเป็น developer",
            ),
            (
                "เปิดหน้าผู้ดูแลระบบโดยยังไม่เข้าสู่ระบบ",
                "1. เปิด URL ใดก็ได้ที่ไม่ใช่ /login\n2. ตรวจสอบ token ใน localStorage\n3. ระบบเปลี่ยนเส้นทาง",
                "/users หรือ /bills",
                "ระบบเปลี่ยนเส้นทางกลับไปหน้า /login",
                "ตรวจจากโค้ด",
                "อ้างอิง checkAuth ใน +layout.svelte",
            ),
            (
                "ออกจากระบบ",
                "1. เข้าสู่ระบบสำเร็จ\n2. กด Sign Out\n3. ตรวจสอบ token\n4. ตรวจสอบหน้าปัจจุบัน",
                "admin_token เดิม",
                "ระบบลบ token และกลับไปหน้า /login",
                "ตรวจจากโค้ด",
                "อ้างอิง logout และ clearToken",
            ),
        ],
    },
    {
        "caption": "ตารางที่ 3.4 กรณีทดสอบการลิงก์ไปหน้าต่าง ๆ ของผู้ดูแลระบบ (ต่อ)",
        "pre": "เข้าสู่ระบบด้วยบัญชีผู้ดูแลระบบแล้ว และเมนูด้านข้างของ Developer Console แสดงผลครบถ้วน",
        "rows": [
            ("ลิงก์ไปหน้า Dashboard", "1. กดเมนู Dashboard\n2. ตรวจสอบหน้า KPI และกราฟสถิติ", "/", "แสดงภาพรวมระบบและสถิติสำคัญ", "ตรวจจากโค้ด", "เรียก getDashboard และ getAnalytics"),
            ("ลิงก์ไปหน้า Bills", "1. กดเมนู Bills\n2. ตรวจสอบตารางรายการบิล\n3. กด Inspect เพื่อดูรายละเอียด", "/bills, /bills/{id}", "แสดงรายการบิลและหน้าแสดงรายละเอียดบิล", "ตรวจจากโค้ด", "รองรับ filter, search, CSV"),
            ("ลิงก์ไปหน้า Payments", "1. กดเมนู Payments\n2. ตรวจสอบรายการชำระเงิน\n3. เปิดรายละเอียดรายการ", "/payments, /payments/{id}", "แสดงรายการชำระเงิน สถานะ สลิป และรายละเอียด", "ตรวจจากโค้ด", "รองรับดูสลิปและ payload"),
            ("ลิงก์ไปหน้า Transactions", "1. กดเมนู Transactions\n2. เลือกประเภทธุรกรรม\n3. กด Refresh", "/transactions", "แสดงสมุดบัญชีแยกประเภทธุรกรรมการเงิน", "ตรวจจากโค้ด", "รองรับ filter และ CSV"),
            ("ลิงก์ไปหน้า Rewards Store", "1. กดเมนู Rewards Store\n2. สลับแท็บรายการของรางวัลและรายการแลกของรางวัล", "/rewards", "แสดงแคตตาล็อกและรายการจัดส่งของรางวัล", "ตรวจจากโค้ด", "มีสร้าง แก้ไข ลบ และอัปเดตสถานะส่งของ"),
            ("ลิงก์ไปหน้า Notifications", "1. กดเมนู Notifications\n2. ตรวจสอบคิวงานแจ้งเตือน\n3. เปิด modal ส่งข้อความ", "/notifications", "แสดงรายการแจ้งเตือนและฟอร์มส่ง Firebase", "ตรวจจากโค้ด", "รองรับส่งแบบทั้งหมด รายคน และ token"),
            ("ลิงก์ไปหน้า Security Events", "1. กดเมนู Security Events\n2. เลือกประเภทเหตุการณ์\n3. เปิดรายละเอียด metadata", "/security", "แสดงเหตุการณ์ความปลอดภัยและข้อมูลประกอบ", "ตรวจจากโค้ด", "ตรวจเหตุการณ์ PIN และ IP"),
            ("ลิงก์ไปหน้า Activity Logs", "1. กดเมนู Activity Logs\n2. ค้นหาหรือกรองตาม action\n3. เปิดรายละเอียด log", "/activity-logs", "แสดงบันทึกกิจกรรมทั่วไปของผู้ใช้", "ตรวจจากโค้ด", "รองรับ purge และลบรายการ"),
            ("ลิงก์ไปหน้า Suspicious", "1. กดเมนู Suspicious\n2. ค้นหาบันทึกพฤติกรรมน่าสงสัย\n3. เปิดฟอร์มเพิ่ม flag", "/suspicious", "แสดงบันทึกความเสี่ยงและเพิ่มรายการเฝ้าระวังได้", "ตรวจจากโค้ด", "รองรับค้นหาผู้ใช้"),
            ("ลิงก์ไปหน้า Users", "1. กดเมนู Users\n2. ค้นหาผู้ใช้\n3. เปิดหน้า Detail", "/users, /users/{id}", "แสดงรายชื่อผู้ใช้และหน้ารายละเอียด", "ตรวจจากโค้ด", "มี suspend, ban, restore"),
            ("ลิงก์ไปหน้า Disputes", "1. กดเมนู Disputes\n2. กรองสถานะข้อพิพาท\n3. กด Review", "/disputes, /disputes/{id}", "แสดงข้อพิพาทและหน้าพิจารณาหลักฐาน", "ตรวจจากโค้ด", "รองรับ mark review และ resolve"),
            ("ลิงก์ไปหน้า Audit Log", "1. กดเมนู Audit Log\n2. ค้นหาการกระทำของผู้ดูแลระบบ", "/audit-logs", "แสดงประวัติการกระทำของผู้ดูแลระบบ", "ตรวจจากโค้ด", "รองรับ export และ clear all"),
            ("ลิงก์ไปหน้า Maintenance", "1. กดเมนู Maintenance\n2. กด Refresh DB Counts\n3. ตรวจสอบปุ่มล้างข้อมูล", "/maintenance", "แสดงจำนวนข้อมูลในฐานข้อมูลและเครื่องมือบำรุงรักษา", "ตรวจจากโค้ด", "มี confirm ก่อนล้างข้อมูล"),
        ],
    },
    {
        "caption": "ตารางที่ 3.4 กรณีทดสอบฟังก์ชันจัดการข้อมูลของผู้ดูแลระบบ (ต่อ)",
        "pre": "เข้าสู่ระบบสำเร็จและ API ฝั่งผู้ดูแลระบบพร้อมให้บริการ",
        "rows": [
            ("โหลด Dashboard และกราฟวิเคราะห์", "1. เปิด Dashboard\n2. รอโหลดข้อมูล\n3. ตรวจสอบการ์ดสถิติและกราฟ", "ข้อมูลสถิติจาก /dashboard และ /analytics", "แสดงจำนวนผู้ใช้ บิล ธุรกรรม ข้อพิพาท การแจ้งเตือน และกราฟสรุป", "ตรวจจากโค้ด", "มี fallback เมื่อ analytics ไม่มีข้อมูล"),
            ("ค้นหาและกรองผู้ใช้", "1. เปิด Users\n2. กรอกคำค้น\n3. เลือกสถานะและบทบาท\n4. กด Refresh", "search, accountStatus, role", "ตารางแสดงผู้ใช้ตามเงื่อนไขและจัดเรียงได้", "ตรวจจากโค้ด", "ใช้ TableHandler และ SearchInput"),
            ("ระงับบัญชีผู้ใช้", "1. เปิด modal Suspend\n2. กรอกเหตุผล\n3. กำหนดจำนวนวัน\n4. กด Confirm", "reason, durationDays", "ระบบเรียก suspendUser และโหลดรายชื่อใหม่", "ตรวจจากโค้ด", "ปุ่มยืนยันถูกปิดถ้าไม่กรอกเหตุผล"),
            ("แบนและกู้คืนบัญชีผู้ใช้", "1. เลือก Ban หรือ Restore\n2. กรอกเหตุผล\n3. กด Confirm", "reason", "ระบบอัปเดตสถานะบัญชีและแสดงข้อความสำเร็จ", "ตรวจจากโค้ด", "เรียก banUser หรือ unsuspendUser"),
            ("ตรวจรายละเอียดบิล", "1. เปิด Bills\n2. กด Inspect\n3. เปิดดู OCR raw data หรือ breakdown", "bill id", "แสดงเจ้าของบิล ยอดเงิน ลูกหนี้ รูปใบเสร็จ และข้อมูล OCR", "ตรวจจากโค้ด", "หน้า /bills/{id} มี toggle raw data"),
            ("ตรวจรายการชำระเงินและสลิป", "1. เปิด Payments\n2. กรองสถานะ ช่องทาง หรือวิธีจ่าย\n3. เปิดรูปสลิป\n4. เปิดหน้ารายละเอียด", "status, channel, method, date", "แสดงยอดเงิน สถานะ ผู้จ่าย สลิป และข้อมูลตรวจสลิป", "ตรวจจากโค้ด", "หน้า detail รองรับ raw response และ QR payload"),
            ("ตรวจธุรกรรมทางการเงิน", "1. เปิด Transactions\n2. เลือก type\n3. กำหนดช่วงวันที่\n4. ส่งออก CSV", "type, dateFrom, dateTo", "แสดงรายการธุรกรรมพร้อมยอดเงิน ผู้สร้าง และวันเวลา", "ตรวจจากโค้ด", "ครอบคลุม debt, payment, refund, write off"),
            ("พิจารณาข้อพิพาท", "1. เปิด Dispute Detail\n2. กด Under Review\n3. ตรวจหลักฐานทั้งสองฝ่าย\n4. เลือกผลตัดสินและกรอก note\n5. กด Resolve", "resolved_paid, resolved_written_off, resolved_rejected", "ระบบอัปเดตสถานะข้อพิพาทและบันทึกผลการตัดสิน", "ตรวจจากโค้ด", "มี slip preview และ raw verification"),
            ("จัดการร้านค้าของรางวัล", "1. เปิด Rewards\n2. สร้างรายการของรางวัล\n3. แก้ไขรายการ\n4. ลบรายการ", "title, pointsCost, stock", "ระบบเพิ่ม แก้ไข หรือลบรายการของรางวัล", "ตรวจจากโค้ด", "ปุ่ม save ตรวจชื่อและแต้มมากกว่า 0"),
            ("อัปเดตสถานะการจัดส่งของรางวัล", "1. เปิดแท็บ Redemptions\n2. เลือก Update Delivery\n3. เลือกสถานะ\n4. กรอก tracking number\n5. บันทึก", "status, trackingNumber", "สถานะรายการแลกของรางวัลถูกอัปเดต", "ตรวจจากโค้ด", "เรียก updateRedemptionStatus"),
            ("ส่งข้อความแจ้งเตือน", "1. เปิด Notifications\n2. กด Send Firebase Push\n3. เลือกกลุ่มเป้าหมาย\n4. กรอกหัวข้อและข้อความ\n5. ส่ง", "all, user, token", "ระบบส่งแจ้งเตือนและแสดงจำนวนสำเร็จ/ล้มเหลว", "ตรวจจากโค้ด", "รองรับแนบรูปและ data payload"),
            (" retry งานแจ้งเตือนที่ล้มเหลว", "1. เปิด Notifications\n2. กรองสถานะ FAILED หรือ SKIPPED\n3. กด Retry", "notification id", "ระบบเรียก retry และโหลดรายการใหม่", "ตรวจจากโค้ด", "ปุ่มแสดงเฉพาะสถานะที่ retry ได้"),
            ("ดูเหตุการณ์ความปลอดภัย", "1. เปิด Security\n2. กรอง userId หรือ event\n3. เปิดรายละเอียด", "pin_brute_force, suspicious_login", "แสดงเวลา เหตุการณ์ ผู้ใช้ หมายเลขเครือข่าย และ metadata", "ตรวจจากโค้ด", "รองรับ export CSV"),
            ("จัดการบันทึกพฤติกรรมน่าสงสัย", "1. เปิด Suspicious\n2. เพิ่ม flag พร้อมคำอธิบาย\n3. เปิดรายละเอียด\n4. ลบรายการหรือ clear all", "type, description, user", "ระบบเพิ่ม แสดง ลบ และล้างบันทึกความเสี่ยงได้", "ตรวจจากโค้ด", "มี confirm ก่อนล้างทั้งหมด"),
            ("บำรุงรักษาข้อมูลระบบ", "1. เปิด Maintenance\n2. refresh จำนวนแถวฐานข้อมูล\n3. purge log เก่า\n4. clear log ตามประเภท", "purge_old, clear_activity, clear_suspicious, clear_audit", "ระบบแสดงผลสำเร็จหรือ error หลังดำเนินการ", "ตรวจจากโค้ด", "มี confirm ก่อนคำสั่งล้างข้อมูล"),
        ],
    },
    {
        "caption": "ตารางที่ 3.4 กรณีทดสอบข้อผิดพลาดและสถานะระบบของผู้ดูแลระบบ (ต่อ)",
        "pre": "เข้าสู่ระบบหรือเรียกใช้งานหน้าผู้ดูแลระบบในสถานการณ์ผิดพลาด",
        "rows": [
            ("API ตอบกลับ 401", "1. เปิดหน้าที่ต้องใช้สิทธิ์\n2. ใช้ token หมดอายุ\n3. ระบบเรียก API\n4. ตรวจสอบการเปลี่ยนเส้นทาง", "token หมดอายุ", "ระบบลบ token และกลับไปหน้า /login", "ตรวจจากโค้ด", "request() จัดการ 401 กลางระบบ"),
            ("API ตอบกลับ error", "1. เปิดหน้ารายการข้อมูล\n2. จำลอง API error\n3. ตรวจข้อความแจ้งเตือน", "server error", "ระบบแสดงกล่องข้อความ error สีแดง", "ตรวจจากโค้ด", "พบรูปแบบ error ในหลายหน้า"),
            ("ไม่พบข้อมูลตามตัวกรอง", "1. เปิดหน้ารายการ\n2. ใส่ตัวกรองที่ไม่มีข้อมูล\n3. ตรวจสอบตาราง", "search ไม่พบข้อมูล", "ระบบแสดงข้อความว่าไม่พบรายการที่ตรงกัน", "ตรวจจากโค้ด", "รองรับใน Users, Bills, Payments, Logs"),
            ("สถานะ loading", "1. เปิดหน้า list หรือ detail\n2. ระบบเริ่มโหลด API\n3. ตรวจสอบหน้าจอระหว่างรอ", "loading = true", "ระบบแสดง LoadingLottie ก่อนข้อมูลพร้อม", "ตรวจจากโค้ด", "ใช้ LoadingLottie หลายหน้า"),
            ("ยกเลิกคำสั่งล้างข้อมูล", "1. เปิด Maintenance หรือ Logs\n2. กดปุ่มล้างข้อมูล\n3. กด Cancel ใน confirm", "cancel confirm", "ระบบไม่เรียก API ล้างข้อมูล", "ตรวจจากโค้ด", "runCleanup และ handleClearAll ตรวจ confirm"),
            ("ส่งแจ้งเตือนโดยกรอกข้อมูลไม่ครบ", "1. เปิด modal ส่งแจ้งเตือน\n2. ไม่กรอก title หรือ body\n3. ตรวจปุ่มส่ง", "title/body ว่าง", "ปุ่มส่งถูกปิดจนกว่าจะกรอกข้อมูลจำเป็นครบ", "ตรวจจากโค้ด", "รวมกรณี target user ต้องเลือกผู้ใช้"),
            ("กรอก JSON payload ไม่ถูกต้อง", "1. เปิด modal ส่งแจ้งเตือน\n2. กรอก data payload ผิดรูปแบบ\n3. กดส่ง", "{ผิดรูปแบบ}", "ระบบแสดงข้อความแจ้งว่ารูปแบบ JSON ไม่ถูกต้อง", "ตรวจจากโค้ด", "อ้างอิง handleSendFcm"),
            ("เปิดหน้ารายละเอียดที่ไม่พบข้อมูล", "1. เปิดหน้า detail ด้วย id ที่ไม่มีอยู่\n2. ระบบเรียก API\n3. ตรวจผลลัพธ์", "id ไม่ถูกต้อง", "ระบบแสดง error หรือข้อความไม่พบข้อมูล", "ตรวจจากโค้ด", "มีใน bills, payments, users, disputes detail"),
            ("ผลการรัน test ฝั่ง Developer Console", "1. รัน npm test ในโฟลเดอร์ developer-console\n2. ตรวจผลลัพธ์จาก command", "npm test -- --run", "คำสั่งไม่สำเร็จเพราะ dependency ของ Vitest ใน node_modules ขาด", "รอติดตั้ง dependency", "ไม่ใช่ผล test case ล้มเหลวจากโค้ด"),
            ("ผลการรัน test ฝั่ง Flutter", "1. รัน flutter test\n2. รอผลลัพธ์\n3. หยุดเมื่อไม่คืนผลภายในเวลาที่กำหนด", "flutter test", "คำสั่งไม่คืน output ภายในเวลาที่รอ จึงยังสรุปผลผ่านจริงไม่ได้", "รอทดสอบซ้ำ", "ตารางนี้อ้างอิงโค้ดเป็นหลัก"),
        ],
    },
]


def text_of(element):
    return "".join(t.text or "" for t in element.iter(qn("w:t"))).strip()


def remove_element(element):
    parent = element.getparent()
    if parent is not None:
        parent.remove(element)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, width):
    tc_pr = cell._tc.get_or_add_tcPr()
    tcw = tc_pr.find(qn("w:tcW"))
    if tcw is None:
        tcw = OxmlElement("w:tcW")
        tc_pr.append(tcw)
    tcw.set(qn("w:w"), str(width))
    tcw.set(qn("w:type"), "dxa")


def set_cell_margins(cell, top=60, start=60, bottom=60, end=60):
    tc_pr = cell._tc.get_or_add_tcPr()
    mar = tc_pr.find(qn("w:tcMar"))
    if mar is None:
        mar = OxmlElement("w:tcMar")
        tc_pr.append(mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def keep_row_together(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = tr_pr.find(qn("w:cantSplit"))
    if cant_split is None:
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)
    cant_split.set(qn("w:val"), "true")


def repeat_header_row(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = tr_pr.find(qn("w:tblHeader"))
    if tbl_header is None:
        tbl_header = OxmlElement("w:tblHeader")
        tr_pr.append(tbl_header)
    tbl_header.set(qn("w:val"), "true")


def set_table_geometry(table, widths):
    table.autofit = False
    table.allow_autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_layout = tbl_pr.find(qn("w:tblLayout"))
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    if grid is None:
        grid = OxmlElement("w:tblGrid")
        table._tbl.insert(0, grid)
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        keep_row_together(row)
        for cell, width in zip(row.cells, widths):
            set_cell_width(cell, width)


def format_runs(paragraph, size=11, bold=False):
    for run in paragraph.runs:
        run.font.name = "TH Sarabun New"
        run._element.rPr.rFonts.set(qn("w:ascii"), "TH Sarabun New")
        run._element.rPr.rFonts.set(qn("w:hAnsi"), "TH Sarabun New")
        run._element.rPr.rFonts.set(qn("w:cs"), "TH Sarabun New")
        run.font.size = Pt(size)
        run.bold = bold


def format_paragraph(paragraph, size=11, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT):
    paragraph.alignment = align
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    format_runs(paragraph, size, bold)


def format_cell(cell, header=False, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    set_cell_margins(cell)
    if header:
        set_cell_shading(cell, "D9EAF7")
    for paragraph in cell.paragraphs:
        format_paragraph(paragraph, size=11, bold=header, align=align)


def format_caption(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_before = Pt(8)
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.keep_with_next = True
    format_runs(paragraph, size=14, bold=True)


def format_body(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.first_line_indent = Pt(28)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.0
    format_runs(paragraph, size=16)


def build_case_table(doc, spec):
    caption = doc.add_paragraph(spec["caption"])
    format_caption(caption)
    table = doc.add_table(rows=1, cols=len(HEADERS))
    table.style = "Table Grid"
    for cell, header in zip(table.rows[0].cells, HEADERS):
        cell.text = header
        format_cell(cell, header=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    repeat_header_row(table.rows[0])

    pre_start = None
    pre_end = None
    for row_idx, row_values in enumerate(spec["rows"], start=1):
        cells = table.add_row().cells
        values = [spec["pre"] if row_idx == 1 else "", *row_values]
        for col_idx, (cell, value) in enumerate(zip(cells, values)):
            cell.text = value
            align = WD_ALIGN_PARAGRAPH.CENTER if col_idx in (3, 5) else WD_ALIGN_PARAGRAPH.LEFT
            format_cell(cell, align=align)
        if row_idx == 1:
            pre_start = cells[0]
        pre_end = cells[0]

    if pre_start is not None and pre_end is not None and pre_start is not pre_end:
        merged = pre_start.merge(pre_end)
        merged.text = spec["pre"]
        format_cell(merged, align=WD_ALIGN_PARAGRAPH.LEFT)

    set_table_geometry(table, [1150, 1350, 2350, 1200, 1800, 720, 790])
    spacer = doc.add_paragraph("")
    spacer.paragraph_format.space_after = Pt(2)
    return [caption._element, table._tbl, spacer._element]


def main():
    doc = Document(INPUT_DOCX)
    body = doc._body._element

    heading = None
    after = None
    for child in body:
        if child.tag != qn("w:p"):
            continue
        text = text_of(child)
        if text.startswith("3.5 การทดสอบระบบโดยวิธี Test Case"):
            heading = child
            after = None
        elif heading is not None and text.startswith("3.6 "):
            after = child

    if heading is None or after is None:
        raise RuntimeError("Could not locate section 3.5 boundaries.")

    in_range = False
    for child in list(body):
        if child is heading:
            in_range = True
            continue
        if child is after:
            break
        if in_range:
            remove_element(child)

    new_elements = []
    intro = doc.add_paragraph(
        "การทดสอบระบบใช้วิธี Test Case เพื่อประเมินความถูกต้องของฟังก์ชันหลัก "
        "โดยกำหนดเงื่อนไขก่อนทดสอบ ขั้นตอน ข้อมูลทดสอบ ผลที่คาดหวัง "
        "และสถานะผลการทดสอบ ทั้งนี้สถานะในตารางเป็นผลจากการอ่านโค้ดและการตรวจคำสั่งทดสอบในสภาพแวดล้อมพัฒนา "
        "ยังไม่ใช่ผลการทดสอบใช้งานจริงกับผู้ใช้ปลายทาง"
    )
    format_body(intro)
    new_elements.append(intro._element)

    note = doc.add_paragraph(
        "จากการตรวจโค้ดพบว่า Developer Console มีระบบเข้าสู่ระบบด้วย Google และ token โดยตรง "
        "มีการตรวจสิทธิ์ role ผู้ดูแลระบบ การ redirect เมื่อไม่ได้เข้าสู่ระบบ และมีหน้าใช้งานหลักสำหรับ Dashboard, Bills, Payments, Transactions, Rewards, Notifications, Security, Activity Logs, Suspicious, Users, Disputes, Audit Logs และ Maintenance"
    )
    format_body(note)
    new_elements.append(note._element)

    for spec in TABLE_SPECS:
        if "ฟังก์ชันจัดการข้อมูล" in spec["caption"]:
            page_break = doc.add_paragraph()
            page_break.add_run().add_break(WD_BREAK.PAGE)
            new_elements.append(page_break._element)
        new_elements.extend(build_case_table(doc, spec))

    for element in reversed(new_elements):
        remove_element(element)
        heading.addnext(element)

    doc.save(OUTPUT_DOCX)

    doc = Document(OUTPUT_DOCX)
    captions = [p.text.strip() for p in doc.paragraphs if p.text.strip().startswith("ตารางที่ 3.4 กรณีทดสอบ")]
    print(f"Saved: {OUTPUT_DOCX}")
    print(f"Test case tables: {len(captions)}")
    print(f"Total document tables: {len(doc.tables)}")
    print(f"Inserted test case rows: {sum(len(t.rows)-1 for t in doc.tables[-4:])}")


if __name__ == "__main__":
    main()
