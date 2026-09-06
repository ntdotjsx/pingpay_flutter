from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


INPUT_DOCX = Path("docs/บทที่ 3_แก้ไขแล้ว_TestCase_APP_ตามตัวอย่าง.docx")
OUTPUT_DOCX = Path("docs/บทที่ 3_แก้ไขแล้ว_TestCase_APP_3.6.docx")


SECTION_36_PARAGRAPHS = [
    (
        "ผู้จัดทำได้ดำเนินการทดสอบและประเมินประสิทธิภาพของระบบตามกรณีทดสอบ (Test Case) "
        "โดยแบ่งการทดสอบออกเป็นส่วนผู้ดูแลระบบและส่วนผู้ใช้งาน เพื่อให้ครอบคลุมฟังก์ชันหลักของระบบ PingPay"
    ),
    (
        "3.6.1 การทดสอบส่วนการทำงานของผู้ดูแลระบบ (Admin) ดำเนินการทดสอบในวันที่ 25 สิงหาคม 2569"
    ),
    (
        "ครอบคลุมฟังก์ชันการเข้าสู่ระบบของผู้ดูแลระบบ การตรวจสอบสิทธิ์การใช้งาน "
        "การลิงก์ไปยังหน้าภาพรวม บิล การชำระเงิน ธุรกรรม ผู้ใช้งาน ข้อพิพาท ของรางวัล "
        "การแจ้งเตือน ความปลอดภัย บันทึกกิจกรรม บันทึกผู้ดูแลระบบ และการบำรุงรักษาระบบ "
        "รวมถึงการจัดการข้อมูลหลักที่เกี่ยวข้องกับสมาชิก บิล และธุรกรรมของระบบ"
    ),
    (
        "3.6.2 การทดสอบส่วนการทำงานของสมาชิก (Member) ดำเนินการทดสอบในวันที่ 26 สิงหาคม 2569"
    ),
    (
        "ครอบคลุมฟังก์ชันการเข้าสู่ระบบด้วย Google การยอมรับนโยบาย PDPA การตั้งค่าและตรวจสอบรหัส PIN "
        "การตั้งค่าโปรไฟล์และช่องทางรับเงิน การจัดการเพื่อน การสร้างบิล การอ่านใบเสร็จด้วย OCR "
        "การกรอกบิลด้วยภาษาธรรมชาติ การแบ่งยอดค่าใช้จ่าย การชำระเงินผ่าน PromptPay และสลิปโอนเงิน "
        "การยื่นข้อพิพาท การรับแจ้งเตือน การแลกของรางวัล และการส่งข้อเสนอแนะ"
    ),
    "3.6.3 การเก็บรวบรวมข้อมูลจากแบบประเมินหลังจากดำเนินการทดสอบระบบ",
    (
        "ภายหลังการทดสอบระบบ ผู้จัดทำได้ให้กลุ่มตัวอย่างทดลองใช้งานระบบ PingPay "
        "และตอบแบบประเมินความพึงพอใจผ่านแบบฟอร์มออนไลน์ (Google Forms) "
        "เพื่อรวบรวมความคิดเห็นและข้อเสนอแนะที่มีต่อระบบ โดยแบบประเมินครอบคลุม 4 ด้านหลัก ได้แก่"
    ),
    "1) ด้านความถูกต้องและความครบถ้วนของฟังก์ชันการทำงาน (Functionality)",
    "2) ด้านการออกแบบส่วนติดต่อผู้ใช้และความง่ายต่อการใช้งาน (Usability & UI Design)",
    "3) ด้านประสิทธิภาพ ความรวดเร็ว และความเสถียรของระบบ (Performance)",
    "4) ด้านความพึงพอใจโดยรวมที่มีต่อระบบ",
    (
        "จากนั้น ผู้จัดทำได้นำข้อมูลที่ได้จากแบบประเมินมาตรวจสอบความสมบูรณ์ "
        "และนำไปวิเคราะห์ผลทางสถิติ ได้แก่ ค่าความถี่ ร้อยละ ค่าเฉลี่ย (X̄) "
        "และส่วนเบี่ยงเบนมาตรฐาน (S.D.) เพื่อสรุประดับความพึงพอใจของกลุ่มตัวอย่างที่มีต่อระบบ PingPay "
        "และใช้ประกอบการสรุปผลการดำเนินโครงงานต่อไป"
    ),
]


def text_of(element):
    return "".join(node.text or "" for node in element.iter() if node.tag == qn("w:t")).strip()


def remove_element(element):
    parent = element.getparent()
    if parent is not None:
        parent.remove(element)


def set_page_break_before(paragraph_element):
    p_pr = paragraph_element.find(qn("w:pPr"))
    if p_pr is None:
        p_pr = OxmlElement("w:pPr")
        paragraph_element.insert(0, p_pr)
    page_break = p_pr.find(qn("w:pageBreakBefore"))
    if page_break is None:
        page_break = OxmlElement("w:pageBreakBefore")
        p_pr.append(page_break)


def format_runs(paragraph, size=16, bold=False):
    for run in paragraph.runs:
        run.font.name = "TH Sarabun New"
        run._element.rPr.rFonts.set(qn("w:ascii"), "TH Sarabun New")
        run._element.rPr.rFonts.set(qn("w:hAnsi"), "TH Sarabun New")
        run._element.rPr.rFonts.set(qn("w:cs"), "TH Sarabun New")
        run.font.size = Pt(size)
        run.bold = bold


def format_body(paragraph, first_indent=True, left_indent=0):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.first_line_indent = Pt(28) if first_indent else Pt(0)
    paragraph.paragraph_format.left_indent = Pt(left_indent)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.line_spacing = 1.0
    format_runs(paragraph, size=15)


def main():
    doc = Document(INPUT_DOCX)
    body = doc._body._element

    heading = None
    after = None
    for child in body:
        text = text_of(child)
        if text.startswith("3.6 ศึกษาวิเคราะห์และการเก็บรวบรวมข้อมูล"):
            heading = child
            after = None
            continue
        if heading is not None and text.startswith("3.7 "):
            after = child

    if heading is None or after is None:
        raise RuntimeError("Could not locate real 3.6 section boundaries.")

    in_range = False
    removed = 0
    for child in list(body):
        if child is heading:
            in_range = True
            continue
        if child is after:
            break
        if in_range:
            remove_element(child)
            removed += 1

    new_elements = []
    for text in SECTION_36_PARAGRAPHS:
        paragraph = doc.add_paragraph(text)
        if text.startswith(("1)", "2)", "3)", "4)")):
            format_body(paragraph, first_indent=False, left_indent=28)
        else:
            format_body(paragraph, first_indent=True)
        new_elements.append(paragraph._element)

    for element in reversed(new_elements):
        remove_element(element)
        heading.addnext(element)

    set_page_break_before(after)
    doc.save(OUTPUT_DOCX)

    doc = Document(OUTPUT_DOCX)
    section_lines = []
    capture = False
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text.startswith("3.6 ศึกษาวิเคราะห์และการเก็บรวบรวมข้อมูล"):
            capture = True
        if capture:
            section_lines.append(text)
        if capture and text.startswith("3.7 "):
            break
    print(f"Saved: {OUTPUT_DOCX}")
    print(f"Removed elements: {removed}")
    print(f"Section 3.6 lines incl heading/3.7 marker: {len([x for x in section_lines if x])}")
    for line in section_lines[:4]:
        print(line)


if __name__ == "__main__":
    main()
