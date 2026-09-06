from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


INPUT_DOCX = Path("docs/บทที่ 3_แก้ไขแล้ว_TestCase_แนวนอน.docx")
OUTPUT_DOCX = Path("docs/บทที่ 3_แก้ไขแล้ว_TestCase_APP_แนวนอน.docx")

HEADERS = [
    "Pre-condition",
    "กรณีทดสอบ\n(Test Case)",
    "ขั้นตอนการทดสอบ\n(Test Step)",
    "Test Data",
    "Expect Results",
    "สถานะ",
    "หมายเหตุ",
]

LANDSCAPE_WIDTHS = [1450, 1700, 3100, 1550, 2400, 900, 1140]


APP_TABLE_SPECS = [
    {
        "caption": "ตารางที่ 3.4 กรณีทดสอบการเริ่มใช้งานแอปพลิเคชันของผู้ใช้งาน (ต่อ)",
        "pre": "ติดตั้งและเปิดใช้งานแอปพลิเคชัน PingPay บนอุปกรณ์ของผู้ใช้งาน",
        "rows": [
            (
                "เปิดแอปครั้งแรก",
                "1. เปิดแอป PingPay\n2. ระบบเริ่มตรวจสอบสถานะผู้ใช้\n3. รอให้โหลดสถานะสำเร็จ",
                "/splash",
                "แสดงหน้าโหลดระหว่างตรวจสอบ session แล้วพาไปหน้าที่ตรงกับสถานะผู้ใช้",
                "ตรวจจากโค้ด",
                "อ้างอิง app_router.dart",
            ),
            (
                "ยังไม่ได้เข้าสู่ระบบ",
                "1. เปิดแอปโดยไม่มี token\n2. ระบบตรวจ session\n3. ระบบ redirect",
                "ไม่มี accessToken",
                "ระบบพาไปหน้าเข้าสู่ระบบ",
                "ตรวจจากโค้ด",
                "อ้างอิง AuthNotifier.checkSession",
            ),
            (
                "เข้าสู่ระบบด้วย Google สำเร็จ",
                "1. กดเข้าสู่ระบบด้วย Google\n2. รับ idToken หรือ accessToken\n3. ส่ง token ไปตรวจสอบ\n4. บันทึก token ที่ได้รับ",
                "บัญชี Google ของผู้ใช้",
                "เข้าสู่ระบบสำเร็จและโหลดข้อมูลผู้ใช้ล่าสุด",
                "มี test บางส่วน",
                "components_and_login_test.dart",
            ),
            (
                "เข้าสู่ระบบไม่สำเร็จ",
                "1. กดเข้าสู่ระบบด้วย Google\n2. ระบบตรวจสอบ token ไม่ผ่าน\n3. จับ error",
                "token ไม่ถูกต้อง",
                "แสดงข้อความผิดพลาดและคงอยู่หน้าเข้าสู่ระบบ",
                "ตรวจจากโค้ด",
                "อ้างอิง LoginScreen และ AuthNotifier",
            ),
            (
                "ผู้ใช้ต้องยอมรับ PDPA",
                "1. เข้าสู่ระบบสำเร็จ\n2. ระบบอ่าน onboardingState\n3. พบสถานะต้องยอมรับ PDPA",
                "pdpaRequired",
                "ระบบพาไปหน้า PDPA",
                "ตรวจจากโค้ด",
                "อ้างอิง router redirect",
            ),
            (
                "ผู้ใช้ต้องตั้งรหัส PIN",
                "1. เข้าสู่ระบบสำเร็จ\n2. ระบบอ่าน onboardingState\n3. พบสถานะต้องตั้ง PIN",
                "pinRequired",
                "ระบบพาไปหน้าตั้งรหัส PIN",
                "ตรวจจากโค้ด",
                "อ้างอิง PinSetupScreen",
            ),
            (
                "ผู้ใช้ต้องกรอกโปรไฟล์",
                "1. เข้าสู่ระบบสำเร็จ\n2. ระบบอ่าน onboardingState\n3. พบสถานะต้องกรอกข้อมูลโปรไฟล์",
                "profileRequired",
                "ระบบพาไปหน้าตั้งค่าชื่อผู้ใช้และข้อมูลส่วนตัว",
                "ตรวจจากโค้ด",
                "อ้างอิง UsernameSetupScreen",
            ),
            (
                "ผู้ใช้ผ่านขั้นตอนเริ่มต้นครบแล้ว",
                "1. เข้าสู่ระบบสำเร็จ\n2. ระบบอ่านสถานะ completed\n3. ตรวจสอบการล็อกแอป",
                "completed",
                "ระบบพาไปหน้า Home หรือหน้า PIN Lock ตามสถานะล็อก",
                "ตรวจจากโค้ด",
                "อ้างอิง router redirect",
            ),
            (
                "session หมดอายุหรือถูกย้ายอุปกรณ์",
                "1. เรียก API ที่ต้องใช้ token\n2. Backend ตอบกลับ 401 หรือแจ้ง session ถูกยุติ\n3. Interceptor ล้าง token",
                "401, SESSION_TERMINATED",
                "ระบบนำผู้ใช้ออกจากระบบและแสดงเหตุผลด้านความปลอดภัย",
                "ตรวจจากโค้ด",
                "อ้างอิง auth_interceptor.dart",
            ),
        ],
    },
    {
        "caption": "ตารางที่ 3.4 กรณีทดสอบการลิงก์ไปหน้าต่าง ๆ ของแอปพลิเคชันผู้ใช้งาน (ต่อ)",
        "pre": "ผู้ใช้งานเข้าสู่ระบบและผ่านขั้นตอนเริ่มต้นครบแล้ว",
        "rows": [
            (
                "ไปหน้า Home",
                "1. เปิดแอปหลังเข้าสู่ระบบ\n2. ระบบเลือกแท็บหลัก Home\n3. โหลดข้อมูลภาพรวม",
                "/home",
                "แสดงหน้าหลักและข้อมูลกิจกรรมของผู้ใช้",
                "ตรวจจากโค้ด",
                "อ้างอิง StatefulShellRoute",
            ),
            (
                "ไปหน้าชำระเงินและหนี้",
                "1. กดแท็บ Payments\n2. ระบบเปิดหน้ารายการหนี้\n3. เลือกแท็บลูกหนี้หรือเจ้าหนี้",
                "/payments, /payments?tab=receivables",
                "แสดงยอดค้างชำระและยอดที่ต้องรับจากเพื่อน",
                "มี test บางส่วน",
                "payment_screen_test.dart",
            ),
            (
                "ไปหน้าบิลของฉัน",
                "1. กดแท็บ My Bills\n2. ระบบโหลดบิลที่ผู้ใช้สร้าง\n3. แสดงตัวกรองสถานะบิล",
                "/bills/my",
                "แสดงรายการบิลของผู้ใช้และสรุปยอดรวม",
                "ตรวจจากโค้ด",
                "อ้างอิง MyBillsScreen",
            ),
            (
                "ไปหน้าร้านของรางวัล",
                "1. กดแท็บ Rewards\n2. ระบบโหลดแต้มและรายการของรางวัล\n3. แสดงประวัติการแลก",
                "/rewards",
                "แสดงของรางวัล แต้มปัจจุบัน และประวัติการแลก",
                "ตรวจจากโค้ด",
                "อ้างอิง RewardRepository",
            ),
            (
                "ไปหน้าโปรไฟล์",
                "1. กดแท็บ Profile\n2. ระบบโหลดข้อมูลผู้ใช้\n3. แสดงการตั้งค่าบัญชีและช่องทางรับเงิน",
                "/profile",
                "แสดงโปรไฟล์ผู้ใช้ การตั้งค่า PromptPay และเมนูออกจากระบบ",
                "ตรวจจากโค้ด",
                "อ้างอิง ProfileScreen",
            ),
            (
                "ไปหน้าสร้างบิล",
                "1. กดปุ่มสร้างบิล\n2. ระบบเปิดหน้าสร้างบิลแบบ modal\n3. โหลดรายชื่อเพื่อน",
                "/bills/create",
                "แสดงฟอร์มสร้างบิลและตัวเลือกเพื่อน",
                "มี test บางส่วน",
                "bill_flow_test.dart",
            ),
            (
                "ไปหน้ารายละเอียดบิล",
                "1. เลือกรายการบิล\n2. ส่ง billId ผ่าน route\n3. ระบบโหลดรายละเอียดบิล",
                "/bills/{id}",
                "แสดงรายละเอียดบิล รายการหนี้ และประวัติการชำระเงิน",
                "ตรวจจากโค้ด",
                "อ้างอิง BillDetailScreen",
            ),
            (
                "ไปหน้าเพื่อน",
                "1. เปิดเมนูเพื่อน\n2. ระบบโหลดรายชื่อเพื่อนและคำขอ\n3. แสดงแท็บเพื่อนและคำขอ",
                "/friends",
                "แสดงรายชื่อเพื่อน คำขอเข้า และคำขอที่ส่งออก",
                "มี test บางส่วน",
                "friends_flow_test.dart",
            ),
            (
                "ไปหน้าเพิ่มเพื่อน",
                "1. กดเพิ่มเพื่อน\n2. ระบบเปิดหน้าค้นหารหัสผู้ใช้\n3. ผู้ใช้กรอกรหัสเพื่อน",
                "/friends/add",
                "แสดงช่องค้นหาและปุ่มส่งคำขอเป็นเพื่อน",
                "มี test บางส่วน",
                "friends_flow_test.dart",
            ),
            (
                "ไปหน้าสแกน QR เพื่อน",
                "1. กดสแกน QR\n2. ระบบเปิดกล้องหรือหน้าสแกน\n3. อ่านรหัสผู้ใช้จาก QR",
                "/friends/scan",
                "แสดงหน้าสแกน QR และปุ่มนำทางกลับ",
                "มี test บางส่วน",
                "qr_scan_test.dart",
            ),
            (
                "ไปหน้ารายละเอียดเพื่อน",
                "1. เลือกเพื่อนจากรายการ\n2. ส่ง friendshipId ผ่าน route\n3. ระบบโหลดรายละเอียด",
                "/friends/{id}",
                "แสดงข้อมูลเพื่อนและสถานะหนี้ที่เกี่ยวข้อง",
                "ตรวจจากโค้ด",
                "อ้างอิง FriendDetailScreen",
            ),
            (
                "ไปหน้าสรุปรายเดือน",
                "1. เปิดหน้าสรุปค่าใช้จ่ายรายเดือน\n2. ระบบรวมข้อมูลบิลและหนี้\n3. แสดงกราฟและตัวเลขสรุป",
                "/analytics/monthly",
                "แสดงกราฟรายเดือนและสรุปยอดค่าใช้จ่าย",
                "มี test บางส่วน",
                "monthly_summary_test.dart",
            ),
            (
                "ไปหน้าส่งข้อเสนอแนะ",
                "1. เปิดเมนู Feedback\n2. เลือกประเภทเรื่อง\n3. กรอกข้อความและส่งข้อมูล",
                "/feedback",
                "แสดงแบบฟอร์มแจ้งปัญหาและส่งข้อเสนอแนะได้",
                "มี test บางส่วน",
                "feedback_sheet_test.dart",
            ),
        ],
    },
    {
        "caption": "ตารางที่ 3.4 กรณีทดสอบฟังก์ชันเพื่อนของผู้ใช้งาน (ต่อ)",
        "pre": "ผู้ใช้งานเข้าสู่ระบบและอยู่ในเมนูเพื่อนของแอปพลิเคชัน",
        "rows": [
            (
                "แสดงรายชื่อเพื่อน",
                "1. เปิดหน้าเพื่อน\n2. ระบบเรียกข้อมูลเพื่อน\n3. แสดงรายการและชื่อเล่น",
                "limit = 50",
                "ผู้ใช้เห็นรายชื่อเพื่อนและสถานะความสัมพันธ์",
                "มี test บางส่วน",
                "friends_flow_test.dart",
            ),
            (
                "ไม่มีเพื่อนในระบบ",
                "1. เปิดหน้าเพื่อน\n2. ระบบได้รับรายการว่าง\n3. แสดง empty state",
                "items = []",
                "แสดงข้อความแนะนำให้เพิ่มเพื่อน",
                "มี test บางส่วน",
                "friends_flow_test.dart",
            ),
            (
                "ค้นหาผู้ใช้จากรหัส",
                "1. เปิดหน้าเพิ่มเพื่อน\n2. กรอกรหัสผู้ใช้\n3. ระบบเรียก API ค้นหา",
                "userCode = PP123456",
                "แสดงข้อมูลผู้ใช้ที่ค้นพบ",
                "ตรวจจากโค้ด",
                "อ้างอิง searchUser",
            ),
            (
                "ค้นหารหัสผู้ใช้ไม่พบ",
                "1. กรอกรหัสที่ไม่มีในระบบ\n2. ระบบรับ USER_NOT_FOUND\n3. แสดงผลว่าไม่พบผู้ใช้",
                "userCode ไม่ถูกต้อง",
                "ไม่สร้างคำขอเป็นเพื่อนและแจ้งผู้ใช้ว่าไม่พบข้อมูล",
                "ตรวจจากโค้ด",
                "อ้างอิง FriendsRepository",
            ),
            (
                "ส่งคำขอเป็นเพื่อนสำเร็จ",
                "1. ค้นหาผู้ใช้พบ\n2. กดส่งคำขอ\n3. ระบบส่ง userCode ไป backend",
                "userCode ของเพื่อน",
                "บันทึกคำขอเป็นเพื่อนและแสดงสถานะรอการตอบรับ",
                "ตรวจจากโค้ด",
                "อ้างอิง sendFriendRequest",
            ),
            (
                "ส่งคำขอให้ตัวเอง",
                "1. กรอกรหัสของตนเอง\n2. กดส่งคำขอ\n3. ระบบรับข้อผิดพลาด",
                "CANNOT_ADD_SELF",
                "แสดงข้อความว่าไม่สามารถส่งคำขอเป็นเพื่อนให้ตัวเองได้",
                "ตรวจจากโค้ด",
                "ข้อความเป็นภาษาไทยใน repository",
            ),
            (
                "ส่งคำขอซ้ำหรือเป็นเพื่อนอยู่แล้ว",
                "1. กรอกรหัสผู้ใช้ที่เคยส่งคำขอแล้วหรือเป็นเพื่อนแล้ว\n2. กดส่งคำขอ\n3. ระบบรับข้อผิดพลาด",
                "ALREADY_FRIENDS หรือ FRIEND_REQUEST_ALREADY_SENT",
                "แสดงข้อความแจ้งสถานะซ้ำให้ผู้ใช้ทราบ",
                "ตรวจจากโค้ด",
                "อ้างอิง FriendsRepository",
            ),
            (
                "ยอมรับคำขอเป็นเพื่อน",
                "1. เปิดแท็บคำขอเข้า\n2. เลือกคำขอ\n3. กดยอมรับ",
                "requestId",
                "สถานะความสัมพันธ์เปลี่ยนเป็นเพื่อน",
                "ตรวจจากโค้ด",
                "อ้างอิง acceptRequest",
            ),
            (
                "ปฏิเสธหรือยกเลิกคำขอ",
                "1. เปิดรายการคำขอ\n2. เลือกคำขอที่ต้องการ\n3. กดปฏิเสธหรือยกเลิก",
                "requestId",
                "คำขอถูกนำออกจากรายการที่เกี่ยวข้อง",
                "ตรวจจากโค้ด",
                "อ้างอิง rejectRequest และ cancelRequest",
            ),
            (
                "ลบเพื่อนที่ไม่มีหนี้ค้าง",
                "1. เปิดรายละเอียดเพื่อน\n2. ตรวจสอบหนี้ค้าง\n3. กดยืนยันลบเพื่อน",
                "friendshipId",
                "ลบความสัมพันธ์เพื่อนสำเร็จ",
                "ตรวจจากโค้ด",
                "อ้างอิง removeFriend",
            ),
            (
                "ลบเพื่อนที่มีหนี้ค้าง",
                "1. เปิดรายละเอียดเพื่อน\n2. ระบบตรวจพบหนี้ค้าง\n3. ผู้ใช้ยืนยันเงื่อนไขเพิ่มเติม",
                "confirmOutstandingDebt = true",
                "ระบบลบเพื่อนตามเงื่อนไขที่ผู้ใช้ยืนยัน",
                "ตรวจจากโค้ด",
                "อ้างอิง checkRemoval",
            ),
        ],
    },
    {
        "caption": "ตารางที่ 3.4 กรณีทดสอบการสร้างบิลและการแบ่งยอดของผู้ใช้งาน (ต่อ)",
        "pre": "ผู้ใช้งานเข้าสู่ระบบและมีรายชื่อเพื่อนสำหรับนำมาแบ่งค่าใช้จ่าย",
        "rows": [
            (
                "เปิดหน้าสร้างบิลเมื่อยังไม่มีเพื่อน",
                "1. เปิดหน้าสร้างบิล\n2. ระบบโหลดรายชื่อเพื่อนเป็นรายการว่าง\n3. แสดงสถานะไม่มีเพื่อน",
                "friends = []",
                "แสดงปุ่มเพิ่มเพื่อนและไม่แสดงปุ่มสร้างบิล",
                "มี test บางส่วน",
                "bill_flow_test.dart",
            ),
            (
                "เลือกเพื่อนเข้าบิล",
                "1. เปิดหน้าสร้างบิล\n2. กดปุ่มเพิ่มเพื่อน\n3. เลือกเพื่อนจาก bottom sheet\n4. กดยืนยัน",
                "รายชื่อเพื่อน 2 คน",
                "ระบบเพิ่มเพื่อนเข้า participant list",
                "มี test บางส่วน",
                "bill_flow_test.dart",
            ),
            (
                "ค้นหาเพื่อนในหน้าสร้างบิล",
                "1. เปิดตัวเลือกเพื่อน\n2. พิมพ์ชื่อหรือชื่อเล่น\n3. ระบบกรองรายการ",
                "คำค้นชื่อเพื่อน",
                "แสดงเฉพาะเพื่อนที่ตรงกับคำค้น",
                "ตรวจจากโค้ด",
                "อ้างอิง FriendSelectionBottomSheet",
            ),
            (
                "เพิ่มรายการสินค้าในบิล",
                "1. เปิดรายการสินค้า\n2. เพิ่มชื่อรายการและราคา\n3. กำหนดผู้ร่วมจ่าย\n4. บันทึก",
                "รายการอาหารและราคา",
                "ระบบแสดงรายการสินค้าและรวมยอดได้ถูกต้อง",
                "มี test บางส่วน",
                "bill_flow_test.dart",
            ),
            (
                "ลบรายการสินค้าหลายรายการ",
                "1. เปิดรายการสินค้า\n2. เลือกหลายรายการ\n3. กดลบแบบกลุ่ม\n4. ยืนยันการลบ",
                "เลือกรายการมากกว่า 1 รายการ",
                "รายการที่เลือกถูกลบออกจากบิล",
                "มี test บางส่วน",
                "bill_flow_test.dart",
            ),
            (
                "สแกนใบเสร็จด้วย OCR",
                "1. เลือกรูปใบเสร็จ\n2. อัปโหลดไฟล์ไป API OCR\n3. รอผลวิเคราะห์\n4. นำรายการมาแสดง",
                "receipt image",
                "ระบบดึงชื่อรายการ ราคา และยอดรวมจากใบเสร็จ",
                "ตรวจจากโค้ด",
                "อ้างอิง scanReceiptOcr",
            ),
            (
                "OCR ใช้เวลานาน",
                "1. อัปโหลดใบเสร็จ\n2. รอผลจาก backend\n3. ตรวจ timeout",
                "ใบเสร็จขนาดใหญ่",
                "ระบบรอผลได้ตามเวลาที่กำหนดและแจ้ง error เมื่อเกินเวลา",
                "ตรวจจากโค้ด",
                "send/receive timeout 90 วินาที",
            ),
            (
                "กรอกบิลด้วยภาษาธรรมชาติ",
                "1. เปิดแผ่นกรอกข้อความ\n2. พิมพ์ประโยคค่าใช้จ่าย\n3. ระบบ parse รายการ\n4. กดนำไปใช้",
                "กินชาบู 1200 หารกับ บาส",
                "ระบบแปลงข้อความเป็นชื่อบิล ยอดรวม รายการ และผู้ร่วมจ่าย",
                "มี test บางส่วน",
                "nli_bill_parser_test.dart",
            ),
            (
                "ระบุยอดจ่ายรายคนเอง",
                "1. กรอกข้อความที่มีชื่อและยอดรายคน\n2. ระบบ parse ยอดเฉพาะบุคคล\n3. ตรวจยอดรวม",
                "ค่ากาแฟ 240 บาส 80 เอ็มมี่ 160",
                "ระบบจับคู่ยอดกับเพื่อนและรวมยอดตรงกับบิล",
                "มี test บางส่วน",
                "nli_bill_parser_test.dart",
            ),
            (
                "แบ่งยอดเท่ากัน",
                "1. ใส่ยอดรวม\n2. เลือกจำนวนผู้ร่วมจ่าย\n3. กดคำนวณ\n4. ตรวจยอดแต่ละคน",
                "498 บาท, 2 คน",
                "ผู้ร่วมจ่ายแต่ละคนได้รับยอดแบ่งเท่ากันตามหลัก satang",
                "มี test บางส่วน",
                "bill_split_calculator_test.dart",
            ),
            (
                "แบ่งยอดแล้วไม่รวมเจ้าของบิล",
                "1. เลือกเพื่อนในบิล\n2. ตั้งค่าไม่รวมฉัน\n3. กดคำนวณ\n4. ตรวจยอดผู้ร่วมจ่าย",
                "ค่าน้ำมัน 800 ไม่รวมฉัน",
                "ยอดทั้งหมดถูกแบ่งให้เพื่อนตามเงื่อนไข",
                "มี test บางส่วน",
                "nli_bill_parser_test.dart",
            ),
            (
                "ยืนยันสร้างบิล",
                "1. กรอกข้อมูลครบ\n2. เปิดหน้าสรุปยืนยัน\n3. กดยืนยันสร้างบิล\n4. ระบบเรียก API",
                "title, totalAmount, participants",
                "สร้างบิลสำเร็จและได้รับข้อมูลบิลจาก backend",
                "ตรวจจากโค้ด",
                "อ้างอิง createBill",
            ),
            (
                "ป้องกันกดสร้างบิลซ้ำ",
                "1. เปิดหน้าสรุปยืนยัน\n2. กดสร้างบิล\n3. ระหว่างส่งข้อมูลให้กดซ้ำอีกครั้ง",
                "กำลัง submit",
                "ระบบไม่ส่งคำขอซ้ำระหว่างกำลังสร้างบิล",
                "มี test บางส่วน",
                "bill_flow_test.dart",
            ),
        ],
    },
    {
        "caption": "ตารางที่ 3.4 กรณีทดสอบการชำระเงินและข้อพิพาทของผู้ใช้งาน (ต่อ)",
        "pre": "ผู้ใช้งานมีบิลหรือหนี้ค้างชำระในระบบ PingPay",
        "rows": [
            (
                "แสดงรายการหนี้ที่ต้องจ่าย",
                "1. เปิดหน้า Payments\n2. ระบบเรียกข้อมูล debts\n3. แสดงสรุปยอดและจำนวนรายการ",
                "/api/v1/bills/debts",
                "ผู้ใช้เห็นยอดหนี้ทั้งหมดและรายการที่ต้องชำระ",
                "มี test บางส่วน",
                "payment_screen_test.dart",
            ),
            (
                "รับสภาพหนี้",
                "1. เปิดรายละเอียดหนี้\n2. ผู้ใช้ยืนยันรับสภาพหนี้\n3. ระบบส่ง billItemId",
                "billItemId",
                "สถานะหนี้เปลี่ยนเป็นรับทราบแล้ว",
                "ตรวจจากโค้ด",
                "อ้างอิง acknowledgeDebt",
            ),
            (
                "ส่งชำระเงินพร้อมสลิป",
                "1. เปิดหนี้ที่ต้องชำระ\n2. แนบสลิปโอนเงิน\n3. ใส่ยอดชำระ\n4. กดส่งหลักฐาน",
                "slipFile, amount",
                "ระบบสร้างรายการชำระเงินและรอเจ้าของบิลยืนยัน",
                "ตรวจจากโค้ด",
                "อ้างอิง submitPaymentWithSlip",
            ),
            (
                "ส่งชำระเงินด้วยข้อมูล QR",
                "1. เปิดหน้าชำระเงิน\n2. ระบบส่งข้อมูล QR พร้อมยอด\n3. บันทึก payment",
                "qrData, method, channel",
                "ระบบบันทึกช่องทางชำระเงินและข้อมูล QR ที่เกี่ยวข้อง",
                "ตรวจจากโค้ด",
                "อ้างอิง PaymentRepository",
            ),
            (
                "เจ้าของบิลยืนยันการชำระเงิน",
                "1. เปิดรายการรอยืนยัน\n2. ตรวจข้อมูลสลิป\n3. กดยืนยันรับเงิน",
                "paymentId",
                "สถานะ payment สำเร็จและยอดค้างลดลง",
                "ตรวจจากโค้ด",
                "อ้างอิง confirmPayment",
            ),
            (
                "เจ้าของบิลปฏิเสธสลิป",
                "1. เปิดรายการรอยืนยัน\n2. ตรวจพบสลิปไม่ถูกต้อง\n3. กรอกเหตุผล\n4. กดปฏิเสธ",
                "paymentId, reason",
                "สถานะ payment ถูกปฏิเสธและแจ้งเหตุผลให้ผู้จ่าย",
                "ตรวจจากโค้ด",
                "อ้างอิง rejectPayment",
            ),
            (
                "แสดงประวัติชำระเงินของบิล",
                "1. เปิดรายละเอียดบิล\n2. ระบบเรียก payment history\n3. แสดงรายการตามเวลา",
                "billId",
                "แสดงประวัติการชำระเงินของบิลครบถ้วน",
                "ตรวจจากโค้ด",
                "อ้างอิง getBillPaymentHistory",
            ),
            (
                "แก้ไขข้อมูลบิล",
                "1. เปิดรายละเอียดบิลที่ตนเป็นเจ้าของ\n2. แก้ชื่อ รายละเอียด หรือยอดรวม\n3. บันทึก",
                "billId, title, totalAmount",
                "ข้อมูลบิลถูกปรับปรุงและแสดงยอดล่าสุด",
                "ตรวจจากโค้ด",
                "อ้างอิง editBill",
            ),
            (
                "แก้ยอดผู้ร่วมจ่าย",
                "1. เปิดรายละเอียดบิล\n2. เลือกผู้ร่วมจ่าย\n3. แก้ยอดรายคน\n4. บันทึก",
                "participantId, newAmount",
                "ยอดของผู้ร่วมจ่ายถูกปรับและระบบคำนวณยอดใหม่",
                "ตรวจจากโค้ด",
                "อ้างอิง editParticipantAmount",
            ),
            (
                "ยกหนี้หรือหักลบหนี้",
                "1. เปิดรายละเอียดบิล\n2. เลือกรายการหนี้\n3. ระบุเหตุผล\n4. กดยืนยัน",
                "participants, reason",
                "ระบบบันทึก write-off และปรับยอดค้างชำระ",
                "ตรวจจากโค้ด",
                "อ้างอิง writeOffDebt",
            ),
            (
                "ยกเลิกบิล",
                "1. เปิดรายละเอียดบิล\n2. กดยกเลิกบิล\n3. ระบุเหตุผล\n4. ยืนยัน",
                "billId, reason",
                "สถานะบิลเปลี่ยนเป็นยกเลิกและไม่แสดงเป็นยอดรอชำระ",
                "มี test บางส่วน",
                "line_calendar_test.dart",
            ),
            (
                "ยื่นข้อพิพาท",
                "1. เปิดรายการหนี้\n2. กดแจ้งข้อพิพาท\n3. ระบุเหตุผลและหลักฐาน\n4. ส่งข้อมูล",
                "billItemId, reason, evidenceUrl",
                "ระบบสร้างข้อพิพาทและแสดงสถานะรอตรวจสอบ",
                "มี test บางส่วน",
                "dispute_model_test.dart",
            ),
            (
                "เจ้าของบิลส่งหลักฐานโต้แย้ง",
                "1. เปิดข้อพิพาทที่เกี่ยวข้อง\n2. กรอกคำชี้แจง\n3. แนบหลักฐาน\n4. ส่งข้อมูล",
                "disputeId, note, evidenceUrl",
                "ระบบบันทึกหลักฐานฝั่งเจ้าของบิล",
                "ตรวจจากโค้ด",
                "อ้างอิง submitCreditorEvidence",
            ),
        ],
    },
    {
        "caption": "ตารางที่ 3.4 กรณีทดสอบการแจ้งเตือน รางวัล และโปรไฟล์ของผู้ใช้งาน (ต่อ)",
        "pre": "ผู้ใช้งานเข้าสู่ระบบและมีข้อมูลบัญชีในระบบ PingPay",
        "rows": [
            (
                "ลงทะเบียนอุปกรณ์เพื่อรับแจ้งเตือน",
                "1. เปิดแอปหลังเข้าสู่ระบบ\n2. ระบบขอหรืออ่าน FCM token\n3. ส่ง token และข้อมูลอุปกรณ์ไป backend",
                "fcmToken, platform",
                "อุปกรณ์ถูกผูกกับบัญชีผู้ใช้สำหรับรับแจ้งเตือน",
                "ตรวจจากโค้ด",
                "อ้างอิง registerDeviceToken",
            ),
            (
                "แสดงศูนย์แจ้งเตือน",
                "1. เปิด notification center\n2. ระบบเรียกแจ้งเตือนของผู้ใช้\n3. แปลง event type เป็นข้อความ",
                "userId, limit = 50",
                "แสดงรายการแจ้งเตือนเป็นภาษาไทยตามชนิดเหตุการณ์",
                "มี test บางส่วน",
                "notification_filter_test.dart",
            ),
            (
                "กรองแจ้งเตือนที่อ่านแล้ว",
                "1. เปิดศูนย์แจ้งเตือน\n2. ทำเครื่องหมายว่าอ่านแล้ว\n3. โหลดรายการใหม่",
                "isRead = true",
                "แจ้งเตือนที่อ่านแล้วถูกซ่อนจาก feed ตามเงื่อนไข",
                "มี test บางส่วน",
                "notification_filter_test.dart",
            ),
            (
                "ล้างสถานะอ่านเก่ากว่า 30 วัน",
                "1. เปิดแอป\n2. ระบบอ่านข้อมูลการอ่านแจ้งเตือนในเครื่อง\n3. ลบข้อมูลเก่า",
                "read record เกิน 30 วัน",
                "ข้อมูลสถานะอ่านที่หมดอายุถูกลบออก",
                "มี test บางส่วน",
                "notification_filter_test.dart",
            ),
            (
                "โหลดรายการของรางวัล",
                "1. เปิดหน้า Rewards\n2. ระบบเรียกรายการสินค้า active\n3. แสดงรายการพร้อมแต้มที่ใช้",
                "/api/v1/rewards/items",
                "ผู้ใช้เห็นรายการของรางวัลที่แลกได้",
                "ตรวจจากโค้ด",
                "อ้างอิง getRewardItems",
            ),
            (
                "ตรวจแต้มผู้ใช้",
                "1. เปิดหน้า Rewards\n2. ระบบเรียกข้อมูลแต้ม\n3. แสดงแต้มและที่อยู่จัดส่งที่บันทึกไว้",
                "/api/v1/rewards/points",
                "แสดงแต้มปัจจุบันและข้อมูลจัดส่งของผู้ใช้",
                "ตรวจจากโค้ด",
                "อ้างอิง getUserPointsInfo",
            ),
            (
                "แลกของรางวัลสำเร็จ",
                "1. เลือกของรางวัล\n2. กรอกชื่อ เบอร์โทร และที่อยู่จัดส่ง\n3. กดยืนยันแลก",
                "rewardItemId, recipientName, phoneNumber, shippingAddress",
                "ระบบบันทึกคำขอแลกของรางวัลและตัดแต้มตามเงื่อนไข",
                "ตรวจจากโค้ด",
                "อ้างอิง redeemReward",
            ),
            (
                "แสดงประวัติแลกของรางวัล",
                "1. เปิดหน้า Rewards\n2. ระบบเรียกประวัติการแลก\n3. แสดงสถานะจัดส่ง",
                "/api/v1/rewards/history",
                "แสดงรายการของรางวัลที่เคยแลกและสถานะล่าสุด",
                "ตรวจจากโค้ด",
                "อ้างอิง getRedemptionHistory",
            ),
            (
                "อัปเดตช่องทางรับเงิน",
                "1. เปิดหน้าโปรไฟล์\n2. แก้ PromptPay หรือบัญชีธนาคาร\n3. บันทึกข้อมูล",
                "promptPayId, bankAccountNumber",
                "ข้อมูลช่องทางรับเงินถูกบันทึกในโปรไฟล์",
                "ตรวจจากโค้ด",
                "อ้างอิง updateProfile",
            ),
            (
                "อัปเดตที่อยู่จัดส่ง",
                "1. เปิดหน้าโปรไฟล์หรือ Rewards\n2. กรอกชื่อผู้รับ เบอร์โทร ที่อยู่\n3. บันทึก",
                "recipientName, phone, address",
                "ข้อมูลจัดส่งถูกอัปเดตและนำไปใช้ในการแลกของรางวัล",
                "ตรวจจากโค้ด",
                "อ้างอิง updateShippingAddress",
            ),
            (
                "เปลี่ยนรหัส PIN",
                "1. เปิดหน้าเปลี่ยน PIN\n2. กรอกรหัสเดิมและรหัสใหม่\n3. ยืนยัน",
                "currentPin, newPin",
                "ระบบบันทึก PIN ใหม่และอัปเดตสถานะผู้ใช้",
                "ตรวจจากโค้ด",
                "อ้างอิง changePin",
            ),
            (
                "ลืมรหัส PIN",
                "1. เปิดขั้นตอนลืม PIN\n2. ขอ OTP\n3. กรอกรหัส OTP\n4. ตั้ง PIN ใหม่",
                "email, otp, resetToken",
                "ระบบยืนยัน OTP และตั้ง PIN ใหม่สำเร็จ",
                "ตรวจจากโค้ด",
                "อ้างอิง requestPinResetOtp",
            ),
            (
                "ส่งข้อเสนอแนะหรือแจ้งปัญหา",
                "1. เปิดหน้า Feedback\n2. เลือกประเภทและความรุนแรง\n3. กรอกรายละเอียด\n4. ส่งข้อมูล",
                "category, message",
                "ระบบตรวจข้อมูลที่จำเป็นและส่งคำขอ feedback",
                "มี test บางส่วน",
                "feedback_sheet_test.dart",
            ),
            (
                "ออกจากระบบ",
                "1. เปิดหน้าโปรไฟล์\n2. กดออกจากระบบ\n3. ระบบเรียก logout และล้าง token",
                "/api/v1/auth/me/logout",
                "ระบบล้างข้อมูล session และกลับไปหน้าเข้าสู่ระบบ",
                "ตรวจจากโค้ด",
                "อ้างอิง AuthRepository.logout",
            ),
        ],
    },
]


def text_of(element):
    return "".join(node.text or "" for node in element.iter() if node.tag == qn("w:t")).strip()


def remove_element(element):
    parent = element.getparent()
    if parent is not None:
        parent.remove(element)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, width):
    tc_pr = cell._tc.get_or_add_tcPr()
    tcw = tc_pr.find(qn("w:tcW"))
    if tcw is None:
        tcw = OxmlElement("w:tcW")
        tc_pr.append(tcw)
    tcw.set(qn("w:w"), str(width))
    tcw.set(qn("w:type"), "dxa")


def set_cell_margins(cell, top=60, start=60, bottom=60, end=60):
    tc_pr = cell._tc.get_or_add_tcPr()
    mar = tc_pr.find(qn("w:tcMar"))
    if mar is None:
        mar = OxmlElement("w:tcMar")
        tc_pr.append(mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def repeat_header_row(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = tr_pr.find(qn("w:tblHeader"))
    if tbl_header is None:
        tbl_header = OxmlElement("w:tblHeader")
        tr_pr.append(tbl_header)
    tbl_header.set(qn("w:val"), "true")


def set_table_geometry(table, widths):
    table.autofit = False
    table.allow_autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_layout = tbl_pr.find(qn("w:tblLayout"))
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    if grid is None:
        grid = OxmlElement("w:tblGrid")
        table._tbl.insert(0, grid)
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            set_cell_width(cell, width)


def format_runs(paragraph, size=11, bold=False):
    for run in paragraph.runs:
        run.font.name = "TH Sarabun New"
        run._element.rPr.rFonts.set(qn("w:ascii"), "TH Sarabun New")
        run._element.rPr.rFonts.set(qn("w:hAnsi"), "TH Sarabun New")
        run._element.rPr.rFonts.set(qn("w:cs"), "TH Sarabun New")
        run.font.size = Pt(size)
        run.bold = bold


def format_paragraph(paragraph, size=11, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT):
    paragraph.alignment = align
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    format_runs(paragraph, size=size, bold=bold)


def format_cell(cell, header=False, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
    set_cell_margins(cell)
    if header:
        set_cell_shading(cell, "D9EAF7")
    for paragraph in cell.paragraphs:
        format_paragraph(paragraph, size=11, bold=header, align=align)


def format_caption(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_before = Pt(8)
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.keep_with_next = True
    format_runs(paragraph, size=14, bold=True)


def build_case_table(doc, spec):
    caption = doc.add_paragraph(spec["caption"])
    format_caption(caption)
    table = doc.add_table(rows=1, cols=len(HEADERS))
    table.style = "Table Grid"
    for cell, header in zip(table.rows[0].cells, HEADERS):
        cell.text = header
        format_cell(cell, header=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    repeat_header_row(table.rows[0])

    pre_start = None
    pre_end = None
    for row_idx, row_values in enumerate(spec["rows"], start=1):
        cells = table.add_row().cells
        values = [spec["pre"] if row_idx == 1 else "", *row_values]
        for col_idx, (cell, value) in enumerate(zip(cells, values)):
            cell.text = value
            align = WD_ALIGN_PARAGRAPH.CENTER if col_idx in (3, 5) else WD_ALIGN_PARAGRAPH.LEFT
            format_cell(cell, align=align)
        if row_idx == 1:
            pre_start = cells[0]
        pre_end = cells[0]

    if pre_start is not None and pre_end is not None and pre_start is not pre_end:
        merged = pre_start.merge(pre_end)
        merged.text = spec["pre"]
        format_cell(merged, align=WD_ALIGN_PARAGRAPH.LEFT)

    set_table_geometry(table, LANDSCAPE_WIDTHS)
    spacer = doc.add_paragraph("")
    spacer.paragraph_format.space_after = Pt(2)
    return [caption._element, table._tbl, spacer._element]


def main():
    doc = Document(INPUT_DOCX)
    body = doc._body._element

    insert_before = None
    seen_real_35 = False
    for child in body:
        text = text_of(child)
        if text.startswith("3.5 การทดสอบระบบโดยวิธี Test Case"):
            seen_real_35 = True
        if seen_real_35 and child.tag == qn("w:p") and child.find(".//" + qn("w:sectPr")) is not None:
            # The section break immediately before 3.6 returns the document to portrait.
            insert_before = child

    if insert_before is None:
        raise RuntimeError("Could not locate the portrait-return section break before 3.6.")

    new_elements = []
    note = doc.add_paragraph(
        "กรณีทดสอบส่วนแอปพลิเคชันผู้ใช้งานต่อไปนี้จัดทำจากการตรวจโค้ด Flutter, repository, router และไฟล์ทดสอบที่มีอยู่ในโครงการ โดยครอบคลุมการเริ่มใช้งาน การลิงก์หน้า เพื่อน การสร้างบิล การชำระเงิน ข้อพิพาท การแจ้งเตือน รางวัล และโปรไฟล์ผู้ใช้งาน"
    )
    note.alignment = WD_ALIGN_PARAGRAPH.LEFT
    note.paragraph_format.first_line_indent = Pt(28)
    note.paragraph_format.space_before = Pt(6)
    note.paragraph_format.space_after = Pt(4)
    note.paragraph_format.line_spacing = 1.0
    format_runs(note, size=16)
    new_elements.append(note._element)

    for spec in APP_TABLE_SPECS:
        new_elements.extend(build_case_table(doc, spec))

    for element in new_elements:
        remove_element(element)
        insert_before.addprevious(element)

    doc.save(OUTPUT_DOCX)
    doc = Document(OUTPUT_DOCX)
    captions = [p.text.strip() for p in doc.paragraphs if p.text.strip().startswith("ตารางที่ 3.4 กรณีทดสอบ")]
    print(f"Saved: {OUTPUT_DOCX}")
    print(f"Test case captions: {len(captions)}")
    print(f"Total document tables: {len(doc.tables)}")
    print(f"Inserted app/user rows: {sum(len(spec['rows']) for spec in APP_TABLE_SPECS)}")


if __name__ == "__main__":
    main()
