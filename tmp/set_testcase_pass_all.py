from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt


INPUT_DOCX = Path("docs/บทที่ 3_แก้ไขแล้ว_TestCase_APP_แนวนอน.docx")
OUTPUT_DOCX = Path("docs/บทที่ 3_แก้ไขแล้ว_TestCase_APP_ผ่านทั้งหมด.docx")


def format_runs(paragraph, size=11):
    for run in paragraph.runs:
        run.font.name = "TH Sarabun New"
        run._element.rPr.rFonts.set(qn("w:ascii"), "TH Sarabun New")
        run._element.rPr.rFonts.set(qn("w:hAnsi"), "TH Sarabun New")
        run._element.rPr.rFonts.set(qn("w:cs"), "TH Sarabun New")
        run.font.size = Pt(size)
        run.bold = False


def set_cell_pass(cell):
    cell.text = "ผ่าน"
    for paragraph in cell.paragraphs:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.0
        format_runs(paragraph, size=11)


def is_test_case_table(table):
    if len(table.columns) != 7 or len(table.rows) < 2:
        return False
    return table.cell(0, 0).text.strip() == "Pre-condition"


def main():
    doc = Document(INPUT_DOCX)
    changed_rows = 0
    changed_cells = 0
    test_tables = 0

    for table in doc.tables:
        if not is_test_case_table(table):
            continue
        test_tables += 1
        for row in table.rows[1:]:
            cells = row.cells
            for col_idx in (4, 5, 6):
                set_cell_pass(cells[col_idx])
                changed_cells += 1
            changed_rows += 1

    doc.save(OUTPUT_DOCX)

    doc = Document(OUTPUT_DOCX)
    remaining = []
    for table_idx, table in enumerate(doc.tables):
        if not is_test_case_table(table):
            continue
        for row_idx, row in enumerate(table.rows[1:], start=1):
            values = [row.cells[i].text.strip() for i in (4, 5, 6)]
            if values != ["ผ่าน", "ผ่าน", "ผ่าน"]:
                remaining.append((table_idx, row_idx, values))

    print(f"Saved: {OUTPUT_DOCX}")
    print(f"Test case tables: {test_tables}")
    print(f"Changed rows: {changed_rows}")
    print(f"Changed cells: {changed_cells}")
    print(f"Rows not pass: {len(remaining)}")


if __name__ == "__main__":
    main()
