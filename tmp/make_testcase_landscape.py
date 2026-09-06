from copy import deepcopy
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


INPUT_DOCX = Path("docs/บทที่ 3_แก้ไขแล้ว_TestCase.docx")
OUTPUT_DOCX = Path("docs/บทที่ 3_แก้ไขแล้ว_TestCase_แนวนอน.docx")
TMP_DOCX = Path("tmp/chapter3_testcase_landscape_work.docx")

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
NS = {"w": W}


def set_cell_width(cell, width):
    tc_pr = cell._tc.get_or_add_tcPr()
    tcw = tc_pr.find(qn("w:tcW"))
    if tcw is None:
        tcw = OxmlElement("w:tcW")
        tc_pr.append(tcw)
    tcw.set(qn("w:w"), str(width))
    tcw.set(qn("w:type"), "dxa")


def set_table_width(table, widths):
    table.autofit = False
    table.allow_autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_layout = tbl_pr.find(qn("w:tblLayout"))
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    if grid is None:
        grid = OxmlElement("w:tblGrid")
        table._tbl.insert(0, grid)
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            set_cell_width(cell, width)


def element_text(element):
    return "".join(t.text or "" for t in element.findall(".//w:t", NS)).strip()


def ensure_section_type(sect_pr, section_type="nextPage"):
    section_type_el = sect_pr.find("w:type", NS)
    if section_type_el is None:
        section_type_el = OxmlElement("w:type")
        sect_pr.insert(0, section_type_el)
    section_type_el.set(qn("w:val"), section_type)


def make_section_break_paragraph(sect_pr):
    p = OxmlElement("w:p")
    p_pr = OxmlElement("w:pPr")
    p.append(p_pr)
    p_pr.append(sect_pr)
    return p


def make_landscape_sect_pr(base_sect_pr):
    sect_pr = deepcopy(base_sect_pr)
    ensure_section_type(sect_pr, "nextPage")
    pg_sz = sect_pr.find("w:pgSz", NS)
    if pg_sz is None:
        pg_sz = OxmlElement("w:pgSz")
        sect_pr.append(pg_sz)
    pg_sz.set(qn("w:w"), "15840")
    pg_sz.set(qn("w:h"), "12240")
    pg_sz.set(qn("w:orient"), "landscape")
    return sect_pr


def make_portrait_sect_pr(base_sect_pr):
    sect_pr = deepcopy(base_sect_pr)
    ensure_section_type(sect_pr, "nextPage")
    pg_sz = sect_pr.find("w:pgSz", NS)
    if pg_sz is None:
        pg_sz = OxmlElement("w:pgSz")
        sect_pr.append(pg_sz)
    pg_sz.set(qn("w:w"), "12240")
    pg_sz.set(qn("w:h"), "15840")
    if qn("w:orient") in pg_sz.attrib:
        del pg_sz.attrib[qn("w:orient")]
    return sect_pr


def widen_test_case_tables():
    doc = Document(INPUT_DOCX)
    # Letter landscape with existing margins L=1.5", R=1.0" gives 8.5" usable width.
    landscape_widths = [1450, 1700, 3100, 1550, 2400, 900, 1140]
    test_case_tables = [t for t in doc.tables if len(t.columns) == 7 and t.cell(0, 0).text.strip() == "Pre-condition"]
    if len(test_case_tables) < 4:
        raise RuntimeError("Could not locate the four test case tables.")
    for table in test_case_tables[-4:]:
        set_table_width(table, landscape_widths)
    doc.save(TMP_DOCX)


def add_landscape_section():
    with ZipFile(TMP_DOCX, "r") as zin:
        files = {name: zin.read(name) for name in zin.namelist()}

    from lxml import etree

    root = etree.fromstring(files["word/document.xml"])
    body = root.find("w:body", NS)
    base_sect_pr = body.find("w:sectPr", NS)
    if base_sect_pr is None:
        raise RuntimeError("Could not locate body section properties.")

    children = list(body)
    start_idx = None
    end_idx = None
    for idx, child in enumerate(children):
        text = element_text(child)
        if text.startswith("3.5 การทดสอบระบบโดยวิธี Test Case"):
            start_idx = idx
            end_idx = None
            continue
        if start_idx is not None and text.startswith("3.6 "):
            end_idx = idx

    if start_idx is None or end_idx is None:
        raise RuntimeError("Could not locate test case table section boundaries.")

    portrait_break = make_section_break_paragraph(make_portrait_sect_pr(base_sect_pr))
    landscape_break = make_section_break_paragraph(make_landscape_sect_pr(base_sect_pr))

    body.insert(start_idx, portrait_break)
    body.insert(end_idx + 1, landscape_break)

    with ZipFile(OUTPUT_DOCX, "w", ZIP_DEFLATED) as zout:
        for name, data in files.items():
            if name == "word/document.xml":
                data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            zout.writestr(name, data)


def main():
    widen_test_case_tables()
    add_landscape_section()
    doc = Document(OUTPUT_DOCX)
    print(f"Saved: {OUTPUT_DOCX}")
    print(f"Sections: {len(doc.sections)}")
    for i, section in enumerate(doc.sections, 1):
        print(i, section.orientation, section.page_width, section.page_height)


if __name__ == "__main__":
    main()
